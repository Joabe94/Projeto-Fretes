#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Simulacao de uso real: aplica mutacoes no arquivo, recalcula e confere o efeito.

Cobre o item 44 do projeto (teste de erros) e o item 45 (duplicidade). Cada
cenario roda sobre uma copia limpa do arquivo original.
"""
from __future__ import annotations

import datetime as dt
import os
import shutil
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "gerador"))

import openpyxl

from comum import LINHA_DADOS as DAT, gs
import dados as D
from recalc_lo import recalc

ORIG = "SISTEMA_FINANCEIRO_PESSOAL_V1.xlsx"
TMP = "/tmp/mut.xlsx"
RES: list[tuple[str, bool, str]] = []

# indices de coluna usados nas mutacoes
LC = dict(id=1, data=2, desc=3, tipo=4, status=5, valor=6, o_tipo=7, o_id=8, d_tipo=9,
          d_id=10, cat=11, sub=12, forma=13, cmp=14, parc=15, meta=16, rec=17, rel=18,
          obs=19, classif=23, vrec=30, vdesp=31, valid=36, vinc=37)
PC = dict(id=1, cmp=2, num=3, venc=4, valor=5, status=15)


def idx_lanc(pred):
    """Indice (0-based) do primeiro lancamento do dataset que satisfaz pred."""
    for i, l in enumerate(D.LANCAMENTOS):
        if pred(l):
            return i
    raise LookupError("lancamento nao encontrado")


def idx_parc(pred):
    for i, p in enumerate(D.PARCELAS):
        if pred(p):
            return i
    raise LookupError("parcela nao encontrada")


def ler(caminho):
    return openpyxl.load_workbook(caminho, data_only=True)


def cenario(nome, mutar, conferir, espera_erro_validacao=False):
    shutil.copy(ORIG, TMP)
    wb = openpyxl.load_workbook(TMP)
    mutar(wb)
    wb.save(TMP)
    rc = recalc(TMP, 600)
    if rc.get("status") != "success":
        RES.append((nome, False, f"formulas quebraram: {rc}"))
        print(f"  [FALHA] {nome}: {rc}")
        return
    v = ler(TMP)
    try:
        msg = conferir(v)
        RES.append((nome, True, msg))
        print(f"  [OK]    {nome}: {msg}")
    except AssertionError as e:
        RES.append((nome, False, str(e)))
        print(f"  [FALHA] {nome}: {e}")


def kpi(v):
    """Coleta os principais numeros do arquivo recalculado."""
    c = v["CAD_Contas"]
    k = v["CAD_Cartoes"]
    inv = v["CAD_Investimentos"]
    m = v["CAD_Metas"]
    cm = v["CAD_Compromissos"]
    t = v["TESTES"]
    falhas = None
    for r in range(1, 12):
        if t.cell(r, 1).value == "TESTES COM FALHA":
            falhas = t.cell(r, 2).value
    return dict(
        saldos=[c.cell(DAT + i, 10).value for i in range(3)],
        disp=[c.cell(DAT + i, 12).value for i in range(3)],
        reserv=[c.cell(DAT + i, 11).value for i in range(3)],
        dividas=[k.cell(DAT + i, 12).value for i in range(2)],
        inv_saldo=[inv.cell(DAT + i, 15).value for i in range(2)],
        metas=[m.cell(DAT + i, 8).value for i in range(4)],
        meta_sit=[m.cell(DAT + i, 15).value for i in range(4)],
        cmp_atual=[cm.cell(DAT + i, 23).value for i in range(3)],
        cmp_devedor=[cm.cell(DAT + i, 22).value for i in range(3)],
        falhas_testes=falhas,
    )


BASE = None


def main():
    global BASE
    print("=" * 78)
    print("SIMULACAO DE USO REAL - TESTE DE ERROS E DUPLICIDADE")
    print("=" * 78)
    BASE = kpi(ler(ORIG))
    print(f"Situacao inicial: saldos {[gs(x) for x in BASE['saldos']]}")
    print(f"                  parcela atual {BASE['cmp_atual']}  "
          f"falhas na bateria interna: {BASE['falhas_testes']}\n")

    # ---------------------------------------------------------------- 1 EDICAO
    i_alu = idx_lanc(lambda l: l["desc"] == "Aluguel" and l["data"] == dt.date(2026, 3, 5))

    def m1(wb):
        wb["LANCAMENTOS"].cell(DAT + i_alu, LC["valor"], 2_500_000)

    def c1(v):
        n = kpi(v)
        d = BASE["saldos"][0] - n["saldos"][0]
        assert abs(d - 500_000) < 1, f"saldo deveria cair 500.000, caiu {gs(d)}"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return f"aluguel de mar/2026 2.000.000 -> 2.500.000; saldo caiu {gs(d)}"
    cenario("EDICAO de um lancamento realizado", m1, c1)

    # ------------------------------------------------------------ 2 CANCELAMENTO
    i_sal = idx_lanc(lambda l: l["desc"] == "Salario mensal" and l["data"] == dt.date(2026, 4, 30))

    def m2(wb):
        ws = wb["LANCAMENTOS"]
        ws.cell(DAT + i_sal, LC["status"], "CANCELADO")
        ws.cell(DAT + i_sal, LC["obs"], "CANCELADO: lancado por engano")

    def c2(v):
        n = kpi(v)
        d = BASE["saldos"][0] - n["saldos"][0]
        assert abs(d - 12_000_000) < 1, f"saldo deveria cair 12.000.000, caiu {gs(d)}"
        wl = v["LANCAMENTOS"]
        assert (wl.cell(DAT + i_sal, LC["vrec"]).value or 0) == 12_000_000, \
            "coluna Vlr_Receita nao deveria mudar (a classificacao independe do status)"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return f"salario de abr/2026 cancelado; saldo caiu {gs(d)} e o historico permaneceu"
    cenario("CANCELAMENTO / ESTORNO por status", m2, c2)

    # ------------------------------------------------------------- 3 ESTORNO PAR
    n_lanc = len(D.LANCAMENTOS)

    def m3(wb):
        ws = wb["LANCAMENTOS"]
        r = DAT + n_lanc
        vals = {"data": dt.date(2026, 6, 20), "desc": "Estorno de compra indevida",
                "tipo": "RECEITA", "status": "REALIZADO", "valor": 300_000,
                "o_tipo": "EXTERNO", "o_id": "", "d_tipo": "CONTA", "d_id": "CON-000001",
                "cat": "CAT-000010", "sub": "SUB-000041", "forma": "Transferencia",
                "rel": "LAN-000200", "obs": "Estorno vinculado ao lancamento LAN-000200"}
        for k_, val in vals.items():
            ws.cell(r, LC[k_], val)
        ws.cell(r, LC["data"]).number_format = "DD/MM/YYYY"

    def c3(v):
        n = kpi(v)
        d = n["saldos"][0] - BASE["saldos"][0]
        assert abs(d - 300_000) < 1, f"saldo deveria subir 300.000, subiu {gs(d)}"
        wl = v["LANCAMENTOS"]
        r = DAT + n_lanc
        assert wl.cell(r, LC["id"]).value == f"LAN-{n_lanc + 1:06d}", "ID automatico incorreto"
        assert wl.cell(r, LC["valid"]).value == "OK", \
            f"validacao: {wl.cell(r, LC['valid']).value}"
        assert "LAN-000200" in str(wl.cell(r, LC["vinc"]).value), "aviso de vinculo ausente"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return (f"nova linha recebeu {wl.cell(r, LC['id']).value} automaticamente, "
                f"validacao OK e aviso de vinculo exibido")
    cenario("ESTORNO com lancamento de contrapartida", m3, c3)

    # ------------------------------------------------------------- 4 EXCLUSAO
    i_ult = n_lanc - 1

    def m4(wb):
        ws = wb["LANCAMENTOS"]
        for col in range(2, 20):
            ws.cell(DAT + i_ult, col).value = None

    def c4(v):
        n = kpi(v)
        wl = v["LANCAMENTOS"]
        assert wl.cell(DAT + i_ult, LC["id"]).value in (None, ""), "ID deveria ficar vazio"
        assert wl.cell(DAT + i_ult, LC["valid"]).value in (None, ""), \
            "validacao deveria ficar vazia"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return "ultima linha apagada; nenhuma formula quebrou e a bateria continua em zero falhas"
    cenario("EXCLUSAO da ultima linha de LANCAMENTOS", m4, c4)

    # --------------------------------------------------------- 5 PAGAR PARCELA 44
    i_p44 = idx_parc(lambda p: p["cmp"] == "CMP-000001" and p["num"] == 44)
    i_l44 = idx_lanc(lambda l: l["parcela"] == D.PARCELAS[i_p44]["id"])

    def m5(wb):
        wb["LANCAMENTOS"].cell(DAT + i_l44, LC["status"], "REALIZADO")

    def c5(v):
        n = kpi(v)
        wp = v["PARCELAS"]
        assert wp.cell(DAT + i_p44, PC["status"]).value == "PAGA", \
            f"parcela 44 deveria ficar PAGA, esta {wp.cell(DAT + i_p44, PC['status']).value}"
        assert n["cmp_atual"][0] == "45/60", f"parcela atual deveria ser 45/60: {n['cmp_atual'][0]}"
        dd = BASE["cmp_devedor"][0] - n["cmp_devedor"][0]
        assert abs(dd - 1_800_000) < 1, f"saldo devedor deveria cair 1.800.000, caiu {gs(dd)}"
        ds = BASE["saldos"][0] - n["saldos"][0]
        assert abs(ds - 1_800_000) < 1, f"conta deveria cair 1.800.000, caiu {gs(ds)}"
        assert wp.cell(DAT + i_p44, 13).value == 1, "deveria haver exatamente 1 pagamento"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return ("parcela 44/60 paga: status PAGA, parcela atual 45/60, devedor -1.800.000, "
                "conta -1.800.000 e nenhum lancamento duplicado")
    cenario("PAGAMENTO da parcela 44/60 (projetado -> realizado)", m5, c5)

    # -------------------------------------------------- 6 DUPLICIDADE DE PARCELA
    def m6(wb):
        ws = wb["LANCAMENTOS"]
        ws.cell(DAT + i_l44, LC["status"], "REALIZADO")
        r = DAT + n_lanc
        src = DAT + i_l44
        for col in range(2, 20):
            ws.cell(r, col, ws.cell(src, col).value)
        ws.cell(r, LC["data"]).number_format = "DD/MM/YYYY"
        ws.cell(r, LC["desc"], "Terreno Santa Rita - parcela 44/60 (DIGITADA DUAS VEZES)")

    def c6(v):
        wp = v["PARCELAS"]
        assert wp.cell(DAT + i_p44, 13).value == 2, "deveria detectar 2 pagamentos"
        wa = v["ALERTAS"]
        achou = any(wa.cell(r, 1).value == "Parcelas com mais de um pagamento realizado"
                    and wa.cell(r, 4).value == "ALERTA" for r in range(1, wa.max_row + 1))
        assert achou, "o alerta de parcela duplicada nao disparou"
        n = kpi(v)
        assert n["falhas_testes"] and n["falhas_testes"] > 0, \
            "a bateria interna deveria acusar falha na duplicidade"
        return (f"duplicidade detectada: alerta ativo e {n['falhas_testes']} teste(s) da "
                "bateria interna em FALHA, como esperado")
    cenario("DUPLICIDADE: mesma parcela paga duas vezes (deve ser detectada)", m6, c6)

    # ------------------------------------------------------- 7 PARCELA ATRASADA
    def m7(wb):
        ws = wb["CFG_Sistema"]
        for r in range(1, 40):
            if ws.cell(r, 2).value == "Data de referencia (hoje)":
                ws.cell(r, 3, dt.date(2026, 10, 20))

    def c7(v):
        wp = v["PARCELAS"]
        atr = sum(1 for i in range(len(D.PARCELAS))
                  if wp.cell(DAT + i, PC["status"]).value == "ATRASADA")
        assert atr > 1, f"deveria haver mais parcelas atrasadas, ha {atr}"
        n = kpi(v)
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return f"data de referencia movida para 20/10/2026: {atr} parcelas ficaram ATRASADAS"
    cenario("PARCELA ATRASADA ao avancar a data de referencia", m7, c7)

    # --------------------------------------------------- 8 RESGATE TOTAL (AHORRO)
    i_res = idx_lanc(lambda l: l["desc"] == "Resgate total Ahorro a Plazo")
    i_ren = idx_lanc(lambda l: l["desc"] == "Rendimento Ahorro a Plazo (vencimento)")

    def m8(wb):
        ws = wb["LANCAMENTOS"]
        ws.cell(DAT + i_res, LC["status"], "REALIZADO")
        ws.cell(DAT + i_ren, LC["status"], "REALIZADO")

    def c8(v):
        n = kpi(v)
        assert abs((n["inv_saldo"][1] or 0)) <= 1, \
            f"o Ahorro a Plazo deveria zerar, esta em {gs(n['inv_saldo'][1])}"
        d = n["saldos"][1] - BASE["saldos"][1]
        assert abs(d - 11_100_000) < 1, f"a conta deveria subir 11.100.000, subiu {gs(d)}"
        wi = v["CAD_Investimentos"]
        assert (wi.cell(DAT + 1, 14).value or 0) == 1_100_000, \
            "o rendimento realizado deveria ser 1.100.000"
        wl = v["LANCAMENTOS"]
        assert (wl.cell(DAT + i_res, LC["vrec"]).value or 0) == 0, \
            "o principal do resgate nao pode virar receita"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return ("resgate total: investimento zerado, conta +11.100.000, rendimento "
                "1.100.000 como receita financeira e principal fora da receita")
    cenario("RESGATE TOTAL de investimento", m8, c8)

    # ------------------------------------------------------- 9 CANCELAR UMA META
    n_mov = len(D.MOV_METAS)

    def m9(wb):
        wm = wb["MOV_METAS"]
        r = DAT + n_mov
        for col, val in [(2, dt.date(2026, 9, 5)), (3, "MET-000001"),
                         (4, "LIBERACAO_CANCELAMENTO"), (5, 20_000_000),
                         (6, "Saldo disponivel"), (8, "Meta cancelada; valor liberado")]:
            wm.cell(r, col, val)
        wm.cell(r, 2).number_format = "DD/MM/YYYY"
        wb["CAD_Metas"].cell(DAT, 13, "Cancelada")
        wb["CAD_Metas"].cell(DAT, 14, dt.date(2026, 9, 5))

    def c9(v):
        n = kpi(v)
        assert abs(n["metas"][0] or 0) < 1, f"a meta deveria zerar: {gs(n['metas'][0])}"
        assert n["meta_sit"][0] == "CANCELADA", f"situacao: {n['meta_sit'][0]}"
        dr = BASE["reserv"][1] - n["reserv"][1]
        assert abs(dr - 20_000_000) < 1, f"o reservado deveria cair 20.000.000, caiu {gs(dr)}"
        dd = n["disp"][1] - BASE["disp"][1]
        assert abs(dd - 20_000_000) < 1, f"o disponivel deveria subir 20.000.000: {gs(dd)}"
        assert n["saldos"][1] == BASE["saldos"][1], "o saldo fisico da conta nao pode mudar"
        wm = v["MOV_METAS"]
        assert wm.cell(DAT + n_mov, 1).value == f"MOV-{n_mov + 1:06d}", "ID automatico incorreto"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return ("meta cancelada com saldo liberado: reservado -20.000.000, disponivel "
                "+20.000.000, saldo bancario inalterado e historico preservado")
    cenario("CANCELAMENTO de meta com saldo reservado", m9, c9)

    # ---------------------------------------------------- 10 CONCLUSAO DE META
    def m10(wb):
        wm = wb["MOV_METAS"]
        r = DAT + n_mov
        for col, val in [(2, dt.date(2026, 9, 5)), (3, "MET-000002"), (4, "RESERVA"),
                         (5, 23_600_000), (8, "Aporte final para concluir a meta")]:
            wm.cell(r, col, val)
        wm.cell(r, 2).number_format = "DD/MM/YYYY"

    def c10(v):
        n = kpi(v)
        assert abs((n["metas"][1] or 0) - 30_000_000) <= 1, \
            f"a meta deveria chegar a 30.000.000: {gs(n['metas'][1])}"
        assert n["meta_sit"][1] == "OBJETIVO ATINGIDO", f"situacao: {n['meta_sit'][1]}"
        wm = v["CAD_Metas"]
        assert abs((wm.cell(DAT + 1, 9).value or 0) - 1.0) < 1e-9, "percentual deveria ser 100%"
        assert abs(wm.cell(DAT + 1, 10).value or 0) <= 1, "faltante deveria ser zero"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return "meta atingiu 100%: situacao OBJETIVO ATINGIDO e valor faltante zerado"
    cenario("CONCLUSAO de meta ao atingir o objetivo", m10, c10)

    # ------------------------------------------- 11 NOVA SUBCATEGORIA + TROCA CAT
    n_sub = len(D.SUBCATEGORIAS)
    i_sup = idx_lanc(lambda l: l["desc"] == "Supermercado (compra 1)"
                     and l["data"].month == 5 and l["data"].year == 2026)

    def m11(wb):
        ws = wb["CAD_Subcategorias"]
        r = DAT + n_sub
        ws.cell(r, 2, "Alimentacao")
        ws.cell(r, 4, "Padaria")
        ws.cell(r, 8, "Ativo")
        wb["LANCAMENTOS"].cell(DAT + i_sup, LC["sub"], f"SUB-{n_sub + 1:06d}")

    def c11(v):
        ws = v["CAD_Subcategorias"]
        r = DAT + n_sub
        assert ws.cell(r, 1).value == f"SUB-{n_sub + 1:06d}", "ID automatico incorreto"
        assert ws.cell(r, 3).value == "CAT-000001", f"categoria: {ws.cell(r, 3).value}"
        wc = v["CAD_Categorias"]
        assert wc.cell(DAT, 4).value == 6, \
            f"Alimentacao deveria ter 6 subcategorias, tem {wc.cell(DAT, 4).value}"
        aux = v["AUX"]
        nomes = [aux.cell(DAT + i, 11).value for i in range(25)]
        assert "Padaria" in nomes, f"a lista dependente nao trouxe Padaria: {nomes[:8]}"
        wl = v["LANCAMENTOS"]
        assert wl.cell(DAT + i_sup, LC["valid"]).value == "OK", \
            f"validacao: {wl.cell(DAT + i_sup, LC['valid']).value}"
        assert wl.cell(DAT + i_sup, 25).value == "Padaria", "o nome resolvido nao atualizou"
        n = kpi(v)
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return ("subcategoria Padaria criada, entrou na lista dependente de Alimentacao "
                "e o lancamento reclassificado permaneceu valido")
    cenario("NOVA SUBCATEGORIA e reclassificacao de um lancamento", m11, c11)

    # ----------------------------------------------- 12 SUBCATEGORIA INCOMPATIVEL
    def m12(wb):
        wb["LANCAMENTOS"].cell(DAT + i_sup, LC["sub"], "SUB-000006")   # Combustivel

    def c12(v):
        wl = v["LANCAMENTOS"]
        val = str(wl.cell(DAT + i_sup, LC["valid"]).value)
        assert val.startswith("ERRO"), f"deveria acusar erro, retornou {val!r}"
        n = kpi(v)
        assert n["falhas_testes"] and n["falhas_testes"] > 0, \
            "a bateria interna deveria acusar o lancamento invalido"
        return f"erro detectado corretamente: {val}"
    cenario("LANCAMENTO ERRADO: subcategoria de outra categoria", m12, c12)

    # ------------------------------------------------------ 13 ORIGEM INEXISTENTE
    def m13(wb):
        wb["LANCAMENTOS"].cell(DAT + i_sup, LC["o_id"], "CON-000099")
        wb["LANCAMENTOS"].cell(DAT + i_sup, LC["o_tipo"], "CONTA")

    def c13(v):
        wl = v["LANCAMENTOS"]
        val = str(wl.cell(DAT + i_sup, LC["valid"]).value)
        assert val.startswith("ERRO"), f"deveria acusar erro, retornou {val!r}"
        return f"erro detectado corretamente: {val}"
    cenario("LANCAMENTO ERRADO: conta de origem inexistente", m13, c13)

    # ------------------------------------------------- 14 NOVA CONTA + TRANSFEREN
    def m14(wb):
        wc = wb["CAD_Contas"]
        r = DAT + 3
        wc.cell(r, 2, "Cuenta Corriente Vision")
        wc.cell(r, 3, "Ueno Bank")
        wc.cell(r, 5, "Conta corrente")
        wc.cell(r, 6, 0)
        wc.cell(r, 7, dt.date(2025, 12, 31))
        wc.cell(r, 7).number_format = "DD/MM/YYYY"
        wc.cell(r, 16, "Ativo")
        ws = wb["LANCAMENTOS"]
        rr = DAT + n_lanc
        for col, val in [(2, dt.date(2026, 8, 10)), (3, "Transferencia para a nova conta"),
                         (4, "TRANSFERENCIA"), (5, "REALIZADO"), (6, 1_000_000),
                         (7, "CONTA"), (8, "CON-000001"), (9, "CONTA"), (10, "CON-000004"),
                         (11, "CAT-000013"), (12, "SUB-000053"), (13, "Transferencia")]:
            ws.cell(rr, col, val)
        ws.cell(rr, 2).number_format = "DD/MM/YYYY"

    def c14(v):
        wc = v["CAD_Contas"]
        assert wc.cell(DAT + 3, 1).value == "CON-000004", "ID automatico incorreto"
        assert abs((wc.cell(DAT + 3, 10).value or 0) - 1_000_000) <= 1, \
            f"a nova conta deveria ter 1.000.000: {gs(wc.cell(DAT + 3, 10).value)}"
        n = kpi(v)
        assert abs((BASE["saldos"][0] - n["saldos"][0]) - 1_000_000) <= 1, \
            "a conta de origem deveria cair 1.000.000"
        soma_antes = sum(BASE["saldos"])
        soma_depois = sum(n["saldos"]) + (wc.cell(DAT + 3, 10).value or 0)
        assert abs(soma_antes - soma_depois) <= 1, \
            f"a transferencia nao pode alterar o total: {gs(soma_antes)} x {gs(soma_depois)}"
        wl = v["LANCAMENTOS"]
        rr = DAT + n_lanc
        assert (wl.cell(rr, LC["vrec"]).value or 0) == 0, "transferencia nao pode virar receita"
        assert (wl.cell(rr, LC["vdesp"]).value or 0) == 0, "transferencia nao pode virar despesa"
        assert (wl.cell(rr, 28).value or 0) == 0 and (wl.cell(rr, 29).value or 0) == 0, \
            "transferencia entre contas nao pode entrar no fluxo de caixa"
        assert wl.cell(rr, LC["valid"]).value == "OK", "validacao deveria passar"
        assert n["falhas_testes"] == 0, "bateria interna acusou falha"
        return ("conta CON-000004 criada e usada na hora; a transferencia nao virou receita "
                "nem despesa, nao entrou no fluxo e o total das contas ficou igual")
    cenario("NOVA CONTA e TRANSFERENCIA para ela", m14, c14)

    print("\n" + "=" * 78)
    ok = sum(1 for _n, o, _m in RES if o)
    print(f"RESULTADO: {ok}/{len(RES)} cenarios com o comportamento esperado")
    print("=" * 78)
    for nome, o, msg in RES:
        if not o:
            print(f"  FALHOU -> {nome}: {msg}")
    os.path.exists(TMP) and os.remove(TMP)
    return len(RES) - ok


if __name__ == "__main__":
    sys.exit(1 if main() else 0)
