# -*- coding: utf-8 -*-
"""Constantes, estilos e utilitarios compartilhados do Sistema Financeiro Pessoal V1."""
from __future__ import annotations

import datetime as dt

from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side

# ----------------------------------------------------------------------------
# PARAMETROS GLOBAIS
# ----------------------------------------------------------------------------
MOEDA_SIGLA = "Gs."
HOJE = dt.date(2026, 9, 5)            # data de referencia do sistema (editavel na aba CFG_Sistema)
DATA_INICIO = dt.date(2025, 12, 31)   # data dos saldos iniciais
ANO_BASE = 2026
HORIZONTE_MESES = 16                  # meses de projecao a partir de HOJE

ARQUIVO_SAIDA = "SISTEMA_FINANCEIRO_PESSOAL_V1.xlsx"

# Limites das areas de dados (linhas). Fora desses limites as formulas nao enxergam os dados.
MAX_LANC = 3000          # linhas do range nomeado de LANCAMENTOS
FORM_LANC = 900          # linhas com formulas pre-preenchidas em LANCAMENTOS
MAX_PARC = 1200
FORM_PARC = 700
MAX_MOV = 600
FORM_MOV = 300
LINHA_HDR = 4            # linha do cabecalho nas abas de dados
LINHA_DADOS = 5          # primeira linha de dados
MAX_CAD = 40             # linhas de cada cadastro
FORM_CAD = 40

# ----------------------------------------------------------------------------
# FORMATOS NUMERICOS
# ----------------------------------------------------------------------------
FMT_GS = '"Gs. "#,##0;-"Gs. "#,##0;"Gs. "0'
FMT_GS_SIMPLES = '"Gs. "#,##0'
FMT_NUM = "#,##0"
FMT_PCT = "0.0%"
FMT_PCT2 = "0.00%"
FMT_DATA = "DD/MM/YYYY"
FMT_MESANO = "MM/YYYY"
FMT_TXT = "@"

# ----------------------------------------------------------------------------
# PALETA
# ----------------------------------------------------------------------------
C_HEADER = "1F3864"      # azul escuro  - cabecalhos de tabela
C_TITULO = "2E5C8A"      # azul medio   - titulos de secao
C_SUB = "D9E2F3"         # azul claro   - subtotais / faixas
C_INPUT = "FFF2CC"       # amarelo      - celulas de digitacao
C_CALC = "F2F2F2"        # cinza        - celulas calculadas
C_OK = "C6EFCE"
C_ERRO = "FFC7CE"
C_ALERTA = "FFEB9C"
C_BRANCO = "FFFFFF"

FONTE = "Arial"

F_TITULO = Font(name=FONTE, size=14, bold=True, color="FFFFFF")
F_SECAO = Font(name=FONTE, size=11, bold=True, color="FFFFFF")
F_HEADER = Font(name=FONTE, size=9, bold=True, color="FFFFFF")
F_NORMAL = Font(name=FONTE, size=10)
F_BOLD = Font(name=FONTE, size=10, bold=True)
F_INPUT = Font(name=FONTE, size=10, color="0000FF")
F_CALC = Font(name=FONTE, size=10, color="000000")
F_LINK = Font(name=FONTE, size=10, color="008000")
F_PEQ = Font(name=FONTE, size=8, italic=True, color="595959")
F_KPI = Font(name=FONTE, size=16, bold=True, color="1F3864")

FILL_HEADER = PatternFill("solid", fgColor=C_HEADER)
FILL_TITULO = PatternFill("solid", fgColor=C_TITULO)
FILL_SUB = PatternFill("solid", fgColor=C_SUB)
FILL_INPUT = PatternFill("solid", fgColor=C_INPUT)
FILL_CALC = PatternFill("solid", fgColor=C_CALC)
FILL_OK = PatternFill("solid", fgColor=C_OK)
FILL_ERRO = PatternFill("solid", fgColor=C_ERRO)
FILL_ALERTA = PatternFill("solid", fgColor=C_ALERTA)

_thin = Side(style="thin", color="BFBFBF")
BORDA = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center")
AL_R = Alignment(horizontal="right", vertical="center")
AL_LW = Alignment(horizontal="left", vertical="top", wrap_text=True)

# ----------------------------------------------------------------------------
# DOMINIOS
# ----------------------------------------------------------------------------
TIPOS_OPERACAO = [
    "RECEITA", "DESPESA", "TRANSFERENCIA", "COMPRA_CARTAO", "PAGAMENTO_CARTAO",
    "APORTE", "RESGATE", "RENDIMENTO", "PAGAMENTO_PARCELA", "AJUSTE_SALDO",
]
# Classificacao economica de cada tipo de operacao (regra de integridade do item 42)
CLASSIFICACAO = {
    "RECEITA": "RECEITA",
    "RENDIMENTO": "RECEITA_FINANCEIRA",
    "DESPESA": "DESPESA",
    "COMPRA_CARTAO": "DESPESA",
    "PAGAMENTO_PARCELA": "DESPESA",
    "TRANSFERENCIA": "NEUTRO",
    "PAGAMENTO_CARTAO": "NEUTRO",
    "APORTE": "NEUTRO",
    "RESGATE": "NEUTRO",
    "AJUSTE_SALDO": "NEUTRO",
}
STATUS_LANC = ["REALIZADO", "PROJETADO", "CANCELADO"]
TIPOS_ENTIDADE = ["CONTA", "CARTAO", "INVESTIMENTO", "EXTERNO"]
TIPOS_CONTA = ["Conta corrente", "Conta poupanca", "Carteira/dinheiro", "Outra"]
FORMAS_PAGAMENTO = ["Dinheiro", "Debito", "Credito", "Transferencia", "Cheque",
                    "Debito automatico", "SIPAP/Giros", "Outro"]
PERIODICIDADES = ["Diario", "Semanal", "Quinzenal", "Mensal", "Bimestral",
                  "Trimestral", "Semestral", "Anual", "Personalizado"]
STATUS_CADASTRO = ["Ativo", "Inativo", "Encerrado"]
TIPOS_INVESTIMENTO = ["Fondo Mutuo", "Ahorro a Plazo", "CDA", "Acoes", "Outro"]
BASES_RENDIMENTO = ["Taxa anual", "Taxa mensal", "Valor fixo", "Manual"]
STATUS_META = ["Ativa", "Concluida", "Cancelada", "Pausada"]
STATUS_PARCELA = ["PAGA", "ABERTA", "ATRASADA", "PROJETADA", "CANCELADA"]
TIPOS_MOV_META = ["RESERVA", "RETIRADA", "TRANSF_ENTRADA", "TRANSF_SAIDA",
                  "LIBERACAO_CONCLUSAO", "LIBERACAO_CANCELAMENTO"]
SINAL_MOV_META = {"RESERVA": 1, "TRANSF_ENTRADA": 1, "RETIRADA": -1,
                  "TRANSF_SAIDA": -1, "LIBERACAO_CONCLUSAO": -1,
                  "LIBERACAO_CANCELAMENTO": -1}
TIPOS_INSTITUICAO = ["Banco", "Financeira", "Cooperativa", "Casa de cambio",
                     "Billetera/Fintech", "Corretora", "Nao aplica", "Outra"]
TIPOS_CATEGORIA = ["RECEITA", "DESPESA", "NEUTRO"]
TIPOS_BEM = ["Imovel", "Veiculo", "Equipamento", "Participacao societaria",
             "Semovente", "Outro"]
SIM_NAO = ["SIM", "NAO"]
DESTINOS_RETIRADA = ["Saldo disponivel", "Outra meta", "Investimento", "Despesa",
                     "Outra conta", "Outro"]

# ----------------------------------------------------------------------------
# UTILITARIOS
# ----------------------------------------------------------------------------

def id_fmt(prefixo: str, n: int) -> str:
    return f"{prefixo}-{n:06d}"


def add_meses(d: dt.date, n: int) -> dt.date:
    """Soma n meses preservando o dia sempre que possivel (fim de mes e limitado)."""
    ano = d.year + (d.month - 1 + n) // 12
    mes = (d.month - 1 + n) % 12 + 1
    dia = min(d.day, dias_no_mes(ano, mes))
    return dt.date(ano, mes, dia)


def dias_no_mes(ano: int, mes: int) -> int:
    if mes == 12:
        return 31
    return (dt.date(ano + (mes == 12), mes % 12 + 1, 1) - dt.timedelta(days=1)).day


def fim_do_mes(d: dt.date) -> dt.date:
    return dt.date(d.year, d.month, dias_no_mes(d.year, d.month))


def dia_do_mes(ano: int, mes: int, dia: int) -> dt.date:
    return dt.date(ano, mes, min(dia, dias_no_mes(ano, mes)))


def anomes(d: dt.date) -> int:
    return d.year * 100 + d.month


def col_letra(idx: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(idx)


def gs(v) -> str:
    """Formata um numero como Guarani para os relatorios de texto."""
    try:
        return f"Gs. {v:,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return str(v)
