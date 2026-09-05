# -*- coding: utf-8 -*-
"""Construcao do arquivo Excel do Sistema Financeiro Pessoal V1."""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from comum import *          # noqa: F401,F403
import comum as K
import dados as D

# ---------------------------------------------------------------------------
# LAYOUT DAS TABELAS PRINCIPAIS
# ---------------------------------------------------------------------------
# (cabecalho, largura, chave_do_dict | None quando for coluna calculada)
COLS_LANC = [
    ("ID_Lancamento", 13, "id"), ("Data", 11, "data"), ("Descricao", 42, "desc"),
    ("Tipo_Operacao", 19, "tipo"), ("Status", 12, "status"), ("Valor", 16, "valor"),
    ("Origem_Tipo", 14, "o_tipo"), ("Origem_ID", 13, "o_id"),
    ("Destino_Tipo", 14, "d_tipo"), ("Destino_ID", 13, "d_id"),
    ("Categoria_ID", 13, "cat"), ("Subcategoria_ID", 15, "sub"),
    ("Forma_Pagamento", 17, "forma"), ("ID_Compromisso", 15, "cmp"),
    ("ID_Parcela", 12, "parcela"), ("ID_Meta", 11, "meta"),
    ("ID_Recorrencia", 15, "rec"), ("ID_Relacionado", 15, "rel"),
    ("Observacao", 48, "obs"),
    ("Ano", 7, None), ("Mes", 6, None), ("AnoMes", 9, None),
    ("Classificacao", 20, None), ("Categoria", 24, None), ("Subcategoria", 24, None),
    ("Origem_Nome", 28, None), ("Destino_Nome", 28, None),
    ("Entrada_Caixa", 15, None), ("Saida_Caixa", 15, None),
    ("Vlr_Receita", 15, None), ("Vlr_Despesa", 15, None), ("Vlr_Rend_Fin", 15, None),
    ("Efeito_Conta", 15, None), ("Efeito_Cartao", 15, None), ("Efeito_Investim", 15, None),
    ("Validacao", 34, None), ("Vinculos", 30, None),
]
# indices (1-based) das colunas de LANCAMENTOS
LC = {h: i + 1 for i, (h, _w, _k) in enumerate(COLS_LANC)}

COLS_PARC = [
    ("ID_Parcela", 12, "id"), ("ID_Compromisso", 15, "cmp"), ("Num_Parcela", 12, "num"),
    ("Data_Vencimento", 15, "venc"), ("Valor_Parcela", 16, "valor"), ("Observacao", 26, "obs"),
    ("Compromisso", 32, None), ("Total_Parcelas", 14, None), ("Rotulo", 10, None),
    ("Pago_Antes_Sistema", 17, None), ("Pagto_Realizado", 16, None),
    ("Data_Pagamento", 15, None), ("Qtd_Lanc_Realizados", 17, None),
    ("Qtd_Lanc_Projetados", 17, None), ("Status_Parcela", 15, None),
    ("Dias_Atraso", 12, None), ("Diferenca_Valor", 15, None),
]
PC = {h: i + 1 for i, (h, _w, _k) in enumerate(COLS_PARC)}

COLS_MOV = [
    ("ID_Movimento", 13, "id"), ("Data", 11, "data"), ("ID_Meta", 11, "meta"),
    ("Tipo_Movimento", 22, "tipo"), ("Valor", 16, "valor"),
    ("Destino_Retirada", 20, "destino"), ("ID_Meta_Destino", 16, "meta_dest"),
    ("Observacao", 52, "obs"),
    ("Meta_Nome", 26, None), ("Sinal", 8, None), ("Valor_Com_Sinal", 17, None),
    ("Conta_Vinculada", 15, None), ("Validacao", 30, None),
]
MC = {h: i + 1 for i, (h, _w, _k) in enumerate(COLS_MOV)}


# ---------------------------------------------------------------------------
# HELPERS DE ESCRITA
# ---------------------------------------------------------------------------
class Livro:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.nomes: dict[str, str] = {}

    def aba(self, nome, cor=None, zoom=90):
        ws = self.wb.create_sheet(nome)
        if cor:
            ws.sheet_properties.tabColor = cor
        ws.sheet_view.zoomScale = zoom
        ws.sheet_view.showGridLines = False
        return ws

    def nome(self, nome, ref):
        self.nomes[nome] = ref
        self.wb.defined_names.add(DefinedName(nome, attr_text=ref))

    def salvar(self, caminho):
        self.wb.save(caminho)


def titulo(ws, linha, texto, ncols=12, sub=None):
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols)
    c = ws.cell(linha, 1, texto)
    c.font = K.F_TITULO
    c.fill = K.FILL_TITULO
    c.alignment = K.AL_L
    ws.row_dimensions[linha].height = 26
    if sub:
        ws.merge_cells(start_row=linha + 1, start_column=1, end_row=linha + 1, end_column=ncols)
        c2 = ws.cell(linha + 1, 1, sub)
        c2.font = K.F_PEQ
        c2.alignment = K.AL_L
        return linha + 2
    return linha + 1


def secao(ws, linha, texto, ncols=12):
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols)
    c = ws.cell(linha, 1, texto)
    c.font = K.F_SECAO
    c.fill = PatternFill("solid", fgColor=K.C_TITULO)
    c.alignment = K.AL_L
    ws.row_dimensions[linha].height = 20
    return linha + 1


def cabecalho(ws, linha, colunas, col_ini=1):
    """colunas = lista de (texto, largura). Retorna proxima linha."""
    for j, (txt, larg) in enumerate(colunas):
        c = ws.cell(linha, col_ini + j, txt)
        c.font = K.F_HEADER
        c.fill = K.FILL_HEADER
        c.alignment = K.AL_C
        c.border = K.BORDA
        ws.column_dimensions[CL(col_ini + j)].width = larg
    ws.row_dimensions[linha].height = 30
    return linha + 1


def rotulo(ws, linha, col, texto, largura=None, bold=True):
    c = ws.cell(linha, col, texto)
    c.font = K.F_BOLD if bold else K.F_NORMAL
    c.alignment = K.AL_L
    if largura:
        ws.column_dimensions[CL(col)].width = largura
    return c


def valor(ws, linha, col, v, fmt=None, entrada=False, bold=False, align=None):
    c = ws.cell(linha, col, v)
    c.font = K.F_INPUT if entrada else (K.F_BOLD if bold else K.F_NORMAL)
    if entrada:
        c.fill = K.FILL_INPUT
    if fmt:
        c.number_format = fmt
    c.border = K.BORDA
    c.alignment = align or (K.AL_L if isinstance(v, str) and not str(v).startswith("=") else K.AL_R)
    return c


def nota(ws, linha, texto, ncols=12):
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols)
    c = ws.cell(linha, 1, texto)
    c.font = K.F_PEQ
    c.alignment = K.AL_LW
    return linha + 1


def dv_lista(ws, nome_range, ref_celulas, titulo_msg="Selecione um valor"):
    dv = DataValidation(type="list", formula1=f"={nome_range}", allow_blank=True,
                        showDropDown=False, errorStyle="warning")
    dv.error = "Valor fora da lista cadastrada. Cadastre-o antes ou escolha um da lista."
    dv.errorTitle = "Valor nao cadastrado"
    dv.prompt = titulo_msg
    ws.add_data_validation(dv)
    dv.add(ref_celulas)
    return dv


HDR = K.LINHA_HDR       # 4
DAT = K.LINHA_DADOS     # 5

# Quando MODO_LIMPO e True o arquivo sai sem nenhum registro de exemplo:
# apenas as categorias e subcategorias padrao, que servem de ponto de partida.
MODO_LIMPO = False


def regs(lista, manter_no_limpo=False):
    """Devolve os registros de exemplo, ou nada quando o arquivo e o limpo."""
    return lista if (manter_no_limpo or not MODO_LIMPO) else []
AVISO = ("PARA APAGAR UM REGISTRO: selecione as celulas AMARELAS da linha e pressione DELETE "
         "(limpar conteudo). NUNCA use 'Excluir linha' nem classifique (sort) esta aba: os IDs "
         "sao gerados pela posicao da linha e os vinculos seriam quebrados em silencio. "
         "Para anular um lancamento sem perder o historico use o campo Status "
         "(CANCELADO / Inativo). Para pesquisar use os filtros do cabecalho.")


class Col:
    """Definicao de uma coluna de uma aba de dados."""

    def __init__(self, h, w, kind="in", key=None, formula=None, fmt=None, dv=None,
                 conv=None):
        self.h, self.w, self.kind = h, w, kind
        self.key, self.formula, self.fmt, self.dv, self.conv = key, formula, fmt, dv, conv


def esqueleto(L, nome, tit, cols, cor, nlin, aviso=AVISO):
    ws = L.aba(nome, cor)
    n = len(cols)
    titulo(ws, 1, tit, n)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    c = ws.cell(2, 1, aviso)
    c.font = Font(name=K.FONTE, size=8, italic=True, color="C00000")
    c.alignment = K.AL_LW
    ws.row_dimensions[2].height = 24
    cabecalho(ws, HDR, [(x.h, x.w) for x in cols])
    ws.freeze_panes = f"A{DAT}"
    ws.auto_filter.ref = f"A{HDR}:{CL(n)}{HDR + nlin}"
    return ws


def preencher(ws, cols, registros, nlin, prefixo_id=None):
    """Escreve valores (registros) e formulas ate nlin linhas."""
    for i in range(nlin):
        r = DAT + i
        rec = registros[i] if i < len(registros) else None
        for j, col in enumerate(cols, start=1):
            cel = ws.cell(r, j)
            cel.border = K.BORDA
            cel.font = K.F_NORMAL
            if col.kind == "auto":
                cel.value = f'=IF($B{r}="","","{prefixo_id}-"&TEXT(ROW()-{HDR},"000000"))'
                cel.font = K.F_BOLD
                cel.fill = K.FILL_CALC
                cel.alignment = K.AL_L
            elif col.kind == "calc":
                cel.value = col.formula.format(r=r)
                cel.fill = K.FILL_CALC
                cel.alignment = K.AL_R
            else:
                if rec is not None:
                    v = col.conv(rec) if col.conv else rec.get(col.key) if isinstance(rec, dict) else None
                    if v not in (None, ""):
                        cel.value = v
                cel.font = K.F_INPUT
                cel.fill = K.FILL_INPUT
                cel.alignment = K.AL_L if isinstance(cel.value, str) else K.AL_R
            if col.fmt:
                cel.number_format = col.fmt
    for j, col in enumerate(cols, start=1):
        if col.dv:
            dv_lista(ws, col.dv, f"{CL(j)}{DAT}:{CL(j)}{DAT + nlin - 1}")


# ---------------------------------------------------------------------------
# CFG_SISTEMA
# ---------------------------------------------------------------------------
def aba_cfg_sistema(L):
    ws = L.aba("CFG_Sistema", "1F3864")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 76
    r = titulo(ws, 1, "CFG_SISTEMA - PARAMETROS GERAIS DO SISTEMA", 4,
               "Somente as celulas amarelas devem ser alteradas. Todo o restante do sistema "
               "le estes parametros.")
    r += 1

    def par(linha, texto, val, nome, fmt=None, ajuda="", dv=None):
        rotulo(ws, linha, 2, texto)
        c = valor(ws, linha, 3, val, fmt, entrada=True)
        c.alignment = K.AL_C
        ws.cell(linha, 4, ajuda).font = K.F_PEQ
        ws.cell(linha, 4).alignment = K.AL_LW
        L.nome(nome, f"CFG_Sistema!$C${linha}")
        if dv:
            dv_lista(ws, dv, f"C{linha}")
        return linha + 1

    r = secao(ws, r, "1. IDENTIFICACAO E MOEDA", 4)
    r = par(r, "Nome do sistema", "Sistema Financeiro Pessoal V1", "CFG_NOME",
            ajuda="Titulo exibido no dashboard.")
    r = par(r, "Moeda padrao (simbolo)", "Gs.", "CFG_MOEDA",
            ajuda="Guarani paraguaio. O formato de celula usado em todo o sistema e "
                  '"Gs. "#.##0 (sem centavos).')
    r += 1
    r = secao(ws, r, "2. DATAS DE REFERENCIA", 4)
    r = par(r, "Data de referencia (hoje)", K.HOJE, "CFG_HOJE", K.FMT_DATA,
            "Define o que e passado e o que e futuro. Para acompanhar o dia corrente "
            "substitua por =HOJE(). Mantida fixa para que os testes sejam reproduziveis.")
    r = par(r, "Data de inicio do sistema (saldos iniciais)", K.DATA_INICIO, "CFG_INICIO",
            K.FMT_DATA, "Data a que se referem os saldos iniciais de contas e cartoes.")
    r = par(r, "Ano de analise (relatorios)", K.ANO_BASE, "CFG_ANO", "0",
            "Usado por REL_Mensal, REL_Anual e INDICADORES.")
    r = par(r, "Mes de analise (relatorios)", K.HOJE.month, "CFG_MES", "0",
            "Numero de 1 a 12. Usado por REL_Mensal.")
    r = par(r, "Horizonte de projecao (meses)", K.HORIZONTE_MESES, "CFG_HORIZONTE", "0",
            "Quantos meses a aba PROJECAO olha para frente a partir da data de referencia.")
    r += 1
    r = secao(ws, r, "3. REGRAS DE NEGOCIO", 4)
    r = par(r, "Considerar reservas de metas no fluxo projetado?", "SIM", "CFG_RESERVA_FLUXO",
            None, "Item 28 do projeto. SIM = o saldo projetado disponivel desconta as reservas "
                  "das metas marcadas com Considerar_No_Fluxo = SIM. NAO = as reservas ficam "
                  "registradas mas nao alteram a projecao.", dv="LST_SIMNAO")
    r += 1
    r = secao(ws, r, "4. PARAMETROS DE ALERTA", 4)
    r = par(r, "Saldo minimo desejado por conta", 4_000_000, "CFG_SALDO_MIN", K.FMT_GS,
            "Abaixo disso a aba ALERTAS acusa saldo baixo.")
    r = par(r, "Uso maximo do limite do cartao", 0.80, "CFG_PCT_CARTAO", K.FMT_PCT,
            "Percentual do limite a partir do qual o cartao entra em alerta.")
    r = par(r, "Dias de antecedencia para alerta de vencimento", 7, "CFG_DIAS_ALERTA", "0",
            "Faturas e parcelas que vencem dentro desse prazo entram em alerta.")
    r = par(r, "Dias de antecedencia para vencimento de investimento", 60, "CFG_DIAS_INV", "0",
            "Investimentos que vencem dentro desse prazo entram em alerta.")
    r = par(r, "Comprometimento maximo da renda mensal", 0.40, "CFG_PCT_RENDA", K.FMT_PCT,
            "Parcelas + faturas de cartao divididas pela receita media mensal.")
    r = par(r, "Variacao de despesa que gera alerta", 0.25, "CFG_PCT_DESP", K.FMT_PCT,
            "Aumento percentual da despesa do mes sobre a media dos ultimos meses.")
    r += 1
    r = secao(ws, r, "5. LIMITES TECNICOS DA VERSAO EXCEL", 4)
    for txt, val in [
        ("Linhas cobertas pelas formulas em LANCAMENTOS", K.MAX_LANC),
        ("Linhas com formulas ja preenchidas em LANCAMENTOS", K.FORM_LANC),
        ("Linhas cobertas pelas formulas em PARCELAS", K.MAX_PARC),
        ("Linhas cobertas pelas formulas em MOV_METAS", K.MAX_MOV),
        ("Linhas por aba de cadastro", K.MAX_CAD),
    ]:
        rotulo(ws, r, 2, txt, bold=False)
        c = ws.cell(r, 3, val); c.number_format = "#,##0"; c.alignment = K.AL_C
        c.font = K.F_NORMAL; c.fill = K.FILL_CALC; c.border = K.BORDA
        r += 1
    r = nota(ws, r + 1,
             "Ao ultrapassar as linhas com formulas preenchidas, copie a ultima linha da aba e "
             "cole para baixo: as formulas se ajustam sozinhas. Ao ultrapassar as linhas "
             "cobertas pelos intervalos nomeados, use Formulas > Gerenciador de Nomes e amplie "
             "o intervalo (por exemplo de $3004 para $6004).", 4)
    return ws


# ---------------------------------------------------------------------------
# CFG_LISTAS
# ---------------------------------------------------------------------------
def aba_cfg_listas(L):
    ws = L.aba("CFG_Listas", "1F3864")
    titulo(ws, 1, "CFG_LISTAS - DOMINIOS FIXOS DO SISTEMA", 17,
           "Listas de dominio usadas nas validacoes. Alterar um valor aqui exige revisar as "
           "formulas que dependem dele (principalmente a coluna Classificacao de LANCAMENTOS).")
    blocos = [
        ("TIPO_OPERACAO", K.TIPOS_OPERACAO, "MAP_TIPO", 20),
        ("CLASSIFICACAO", [K.CLASSIFICACAO[t] for t in K.TIPOS_OPERACAO], "MAP_CLASS", 21),
        ("STATUS_LANCAMENTO", K.STATUS_LANC, "LST_STATUS", 19),
        ("TIPO_ENTIDADE", K.TIPOS_ENTIDADE, "LST_ENT_TIPO", 16),
        ("TIPO_CONTA", K.TIPOS_CONTA, "LST_TIPOCONTA", 19),
        ("FORMA_PAGAMENTO", K.FORMAS_PAGAMENTO, "LST_FORMA", 19),
        ("PERIODICIDADE", K.PERIODICIDADES, "LST_PERIOD", 16),
        ("STATUS_CADASTRO", K.STATUS_CADASTRO, "LST_STCAD", 17),
        ("TIPO_INVESTIMENTO", K.TIPOS_INVESTIMENTO, "LST_TIPOINV", 19),
        ("BASE_RENDIMENTO", K.BASES_RENDIMENTO, "LST_BASEREND", 17),
        ("STATUS_META", K.STATUS_META, "LST_STMETA", 14),
        ("STATUS_PARCELA", K.STATUS_PARCELA, "LST_STPARC", 16),
        ("TIPO_MOV_META", K.TIPOS_MOV_META, "LST_TIPOMOV", 23),
        ("SINAL_MOV_META", [K.SINAL_MOV_META[t] for t in K.TIPOS_MOV_META], "MAP_SINALMOV", 16),
        ("SIM_NAO", K.SIM_NAO, "LST_SIMNAO", 10),
        ("DESTINO_RETIRADA_META", K.DESTINOS_RETIRADA, "LST_DESTRET", 22),
        ("MES_NOME", ["Janeiro", "Fevereiro", "Marco", "Abril", "Maio", "Junho", "Julho",
                      "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"],
         "LST_MESES", 13),
    ]
    cabecalho(ws, 3, [(b[0], b[3]) for b in blocos])
    for j, (h, itens, nome, _w) in enumerate(blocos, start=1):
        for i, it in enumerate(itens):
            c = ws.cell(4 + i, j, it)
            c.font = K.F_NORMAL
            c.border = K.BORDA
            c.alignment = K.AL_C if not isinstance(it, str) else K.AL_L
        L.nome(nome, f"CFG_Listas!${CL(j)}$4:${CL(j)}${3 + len(itens)}")
    ws.freeze_panes = "A4"
    return ws


# ---------------------------------------------------------------------------
# INTERVALOS NOMEADOS
# ---------------------------------------------------------------------------
def _bloco(L, aba, n_linhas, mapa):
    fim = DAT + n_linhas - 1
    for nome, letra in mapa.items():
        L.nome(nome, f"{aba}!${letra}${DAT}:${letra}${fim}")


def definir_nomes(L):
    _bloco(L, "LANCAMENTOS", K.MAX_LANC, {
        "L_ID": "A", "L_Data": "B", "L_Desc": "C", "L_Tipo": "D", "L_Status": "E",
        "L_Valor": "F", "L_OrigTipo": "G", "L_OrigID": "H", "L_DestTipo": "I",
        "L_DestID": "J", "L_CatID": "K", "L_SubID": "L", "L_Forma": "M", "L_CmpID": "N",
        "L_ParcID": "O", "L_MetaID": "P", "L_RecID": "Q", "L_RelID": "R",
        "L_Ano": "T", "L_Mes": "U", "L_AnoMes": "V", "L_Class": "W", "L_CatNome": "X",
        "L_SubNome": "Y", "L_OrigNome": "Z", "L_DestNome": "AA", "L_Ent": "AB",
        "L_Sai": "AC", "L_VRec": "AD", "L_VDesp": "AE", "L_VRend": "AF",
        "L_EfConta": "AG", "L_EfCartao": "AH", "L_EfInv": "AI", "L_Valid": "AJ"})
    _bloco(L, "PARCELAS", K.MAX_PARC, {
        "P_ID": "A", "P_CMP": "B", "P_NUM": "C", "P_VENC": "D", "P_VALOR": "E",
        "P_OBS": "F", "P_NOME": "G", "P_TOTAL": "H", "P_ROT": "I", "P_ANTES": "J",
        "P_PAGO": "K", "P_DTPAGO": "L", "P_QTDREAL": "M", "P_QTDPROJ": "N",
        "P_STATUS": "O", "P_ATRASO": "P", "P_DIF": "Q"})
    _bloco(L, "MOV_METAS", K.MAX_MOV, {
        "MM_ID": "A", "MM_DATA": "B", "MM_META": "C", "MM_TIPO": "D", "MM_VALOR": "E",
        "MM_DEST": "F", "MM_METADEST": "G", "MM_NOME": "I", "MM_SINAL": "J",
        "MM_VALSINAL": "K", "MM_CONTA": "L", "MM_VALID": "M"})
    _bloco(L, "CAD_Contas", K.MAX_CAD, {
        "C_ID": "A", "C_NOME": "B", "C_TIPO": "E", "C_SALDOINI": "F", "C_DATAINI": "G",
        "C_ENTR": "H", "C_SAID": "I", "C_SALDO": "J", "C_RESERV": "K", "C_DISP": "L",
        "C_ENTP": "M", "C_SAIP": "N", "C_PROJ": "O", "C_STATUS": "P"})
    _bloco(L, "CAD_Cartoes", K.MAX_CAD, {
        "K_ID": "A", "K_NOME": "B", "K_LIMITE": "E", "K_DIVINI": "F", "K_FECHDIA": "H",
        "K_VENCDIA": "I", "K_COMPRAS": "J", "K_PAGTOS": "K", "K_DIVIDA": "L",
        "K_PCT": "M", "K_DISP": "N", "K_FECHA": "O", "K_CICLOINI": "P", "K_VENC": "Q",
        "K_FATURA": "R", "K_FUTURAS": "S", "K_PARCFUT": "T", "K_COMPR": "U",
        "K_STATUS": "V"})
    _bloco(L, "CAD_Investimentos", K.MAX_CAD, {
        "I_ID": "A", "I_NOME": "B", "I_TIPO": "E", "I_VALINI": "F", "I_DATAINI": "G",
        "I_BASE": "H", "I_TAXA": "I", "I_PRAZO": "J", "I_VENC": "K", "I_APORTES": "L",
        "I_RESG": "M", "I_REND": "N", "I_SALDO": "O", "I_PRINC": "P", "I_RENDSALDO": "Q",
        "I_MESES": "R", "I_RENDPROJ": "S", "I_RENDEST": "T", "I_ESTVENC": "U",
        "I_APORTPROJ": "V", "I_RESGPROJ": "W", "I_STATUS": "X"})
    _bloco(L, "CAD_Categorias", K.MAX_CAD, {
        "CAT_ID": "A", "CAT_NOME": "B", "CAT_TIPO": "C", "CAT_QTDSUB": "D",
        "CAT_REAL": "E", "CAT_PROJ": "F", "CAT_STATUS": "G"})
    _bloco(L, "CAD_Subcategorias", K.MAX_CAD * 3, {
        "SUB_ID": "A", "SUB_CATNOME": "B", "SUB_CATID": "C", "SUB_NOME": "D",
        "SUB_REAL": "E", "SUB_PROJ": "F", "SUB_ORDEM": "G", "SUB_STATUS": "H"})
    _bloco(L, "CAD_Instituicoes", K.MAX_CAD, {"INS_ID": "A", "INS_NOME": "B"})
    _bloco(L, "CAD_Compromissos", K.MAX_CAD, {
        "CM_ID": "A", "CM_NOME": "B", "CM_CATID": "D", "CM_SUBID": "F", "CM_QTD": "G",
        "CM_VLRPARC": "H", "CM_TOTAL": "I", "CM_PAGASANTES": "M", "CM_ENTTIPO": "N", "CM_ENTID": "O",
        "CM_PAGAS": "P", "CM_ABERTAS": "Q", "CM_ATRAS": "R", "CM_PROJ": "S",
        "CM_REST": "T", "CM_PAGO": "U", "CM_DEVEDOR": "V", "CM_ATUAL": "W",
        "CM_PROXVENC": "X", "CM_STATUS": "Y"})
    _bloco(L, "CAD_Metas", K.MAX_CAD, {
        "M_ID": "A", "M_NOME": "B", "M_OBJ": "C", "M_PRAZO": "D", "M_CONTA": "E",
        "M_INV": "F", "M_FLUXO": "G", "M_RESERV": "H", "M_PCT": "I", "M_FALTA": "J",
        "M_MESES": "K", "M_APMENSAL": "L", "M_STATUS": "M", "M_CONCL": "N",
        "M_SITUACAO": "O", "M_CONTANOME": "P"})
    _bloco(L, "CAD_Recorrencias", K.MAX_CAD, {
        "R_ID": "A", "R_DESC": "B", "R_TIPO": "C", "R_VALOR": "D", "R_PERIOD": "E",
        "R_STATUS": "Q", "R_GERADOS": "R", "R_PROXIMA": "U"})
    _bloco(L, "CAD_Bens", K.MAX_CAD, {
        "B_ID": "A", "B_NOME": "B", "B_VALOR": "D", "B_CMP": "F", "B_DEVEDOR": "G",
        "B_LIQ": "H", "B_STATUS": "I"})
    n_ent = 3 * K.MAX_CAD
    L.nome("ENT_ID", f"AUX!$A${DAT}:$A${DAT + n_ent - 1}")
    L.nome("ENT_TIPO", f"AUX!$B${DAT}:$B${DAT + n_ent - 1}")
    L.nome("ENT_NOME", f"AUX!$C${DAT}:$C${DAT + n_ent - 1}")


# ---------------------------------------------------------------------------
# CADASTROS
# ---------------------------------------------------------------------------
INS_NOME = {i[0]: i[1] for i in D.INSTITUICOES}
CAT_NOME = {c[0]: c[1] for c in D.CATEGORIAS}
SUB_NOME = {s[0]: s[2] for s in D.SUBCATEGORIAS}


def _f_inst(r):
    return f'=IF($C{r}="","",IFERROR(INDEX(INS_ID,MATCH($C{r},INS_NOME,0)),"?"))'


def abas_cadastros(L):
    n = K.MAX_CAD

    # ---------------- INSTITUICOES ----------------
    cols = [
        Col("ID_Instituicao", 15, "auto"),
        Col("Nome", 32, "in", conv=lambda x: x[1]),
        Col("Tipo", 16, "in", conv=lambda x: x[2]),
        Col("Qtd_Contas", 12, "calc", formula='=IF($A{r}="","",COUNTIF(CAD_Contas!$D${d}:$D${f},$A{r}))'.replace("{d}", str(DAT)).replace("{f}", str(DAT + n - 1)), fmt="0"),
        Col("Qtd_Cartoes", 12, "calc", formula='=IF($A{r}="","",COUNTIF(CAD_Cartoes!$D${d}:$D${f},$A{r}))'.replace("{d}", str(DAT)).replace("{f}", str(DAT + n - 1)), fmt="0"),
        Col("Qtd_Investimentos", 16, "calc", formula='=IF($A{r}="","",COUNTIF(CAD_Investimentos!$D${d}:$D${f},$A{r}))'.replace("{d}", str(DAT)).replace("{f}", str(DAT + n - 1)), fmt="0"),
        Col("Observacao", 46, "in", conv=lambda x: x[3]),
    ]
    ws = esqueleto(L, "CAD_Instituicoes", "CAD_INSTITUICOES - BANCOS E EMISSORES", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.INSTITUICOES), n, "INS")

    # ---------------- CONTAS ----------------
    cols = [
        Col("ID_Conta", 13, "auto"),
        Col("Nome", 30, "in", conv=lambda x: x[1]),
        Col("Instituicao", 26, "in", conv=lambda x: INS_NOME[x[2]], dv="INS_NOME"),
        Col("ID_Instituicao", 14, "calc", formula=_f_inst("{r}")),
        Col("Tipo", 18, "in", conv=lambda x: x[3], dv="LST_TIPOCONTA"),
        Col("Saldo_Inicial", 17, "in", conv=lambda x: x[4], fmt=K.FMT_GS),
        Col("Data_Saldo_Inicial", 16, "in", conv=lambda x: x[5], fmt=K.FMT_DATA),
        Col("Entradas_Realizadas", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_DestTipo,"CONTA",L_DestID,$A{r},L_Status,"REALIZADO"))'),
        Col("Saidas_Realizadas", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_OrigTipo,"CONTA",L_OrigID,$A{r},L_Status,"REALIZADO"))'),
        Col("Saldo_Calculado", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$F{r}+$H{r}-$I{r})'),
        Col("Reservado_Metas", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(MM_VALSINAL,MM_CONTA,$A{r}))'),
        Col("Saldo_Disponivel", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$J{r}-$K{r})'),
        Col("Entradas_Projetadas", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_DestTipo,"CONTA",L_DestID,$A{r},L_Status,"PROJETADO"))'),
        Col("Saidas_Projetadas", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_OrigTipo,"CONTA",L_OrigID,$A{r},L_Status,"PROJETADO"))'),
        Col("Saldo_Projetado", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$J{r}+$M{r}-$N{r})'),
        Col("Status", 12, "in", conv=lambda x: x[6], dv="LST_STCAD"),
        Col("Observacao", 40, "in", conv=lambda x: x[7]),
    ]
    ws = esqueleto(L, "CAD_Contas", "CAD_CONTAS - CONTAS BANCARIAS E DINHEIRO FISICO", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.CONTAS), n, "CON")

    # ---------------- CARTOES ----------------
    base_ini = ('DATE(YEAR(CFG_HOJE),MONTH(CFG_HOJE)+IF(DAY(CFG_HOJE)>$H{r},1,0),1)')
    base_venc = ('DATE(YEAR($O{r}),MONTH($O{r})+IF($I{r}>$H{r},0,1),1)')
    cols = [
        Col("ID_Cartao", 13, "auto"),
        Col("Nome", 32, "in", conv=lambda x: x[1]),
        Col("Instituicao", 26, "in", conv=lambda x: INS_NOME[x[2]], dv="INS_NOME"),
        Col("ID_Instituicao", 14, "calc", formula=_f_inst("{r}")),
        Col("Limite", 16, "in", conv=lambda x: x[3], fmt=K.FMT_GS),
        Col("Divida_Inicial", 16, "in", conv=lambda x: x[4], fmt=K.FMT_GS),
        Col("Data_Divida_Inicial", 16, "in", conv=lambda x: x[5], fmt=K.FMT_DATA),
        Col("Dia_Fechamento", 14, "in", conv=lambda x: x[6], fmt="0"),
        Col("Dia_Vencimento", 14, "in", conv=lambda x: x[7], fmt="0"),
        Col("Compras_Realizadas", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_OrigTipo,"CARTAO",L_OrigID,$A{r},L_Status,"REALIZADO"))'),
        Col("Pagamentos_Realizados", 19, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_DestTipo,"CARTAO",L_DestID,$A{r},L_Status,"REALIZADO"))'),
        Col("Divida_Atual", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$F{r}+$J{r}-$K{r})'),
        Col("Pct_Limite_Utilizado", 16, "calc", fmt=K.FMT_PCT,
            formula='=IF(OR($A{r}="",$E{r}=0),0,$L{r}/$E{r})'),
        Col("Limite_Disponivel", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$E{r}-$L{r})'),
        Col("Fechamento_Ciclo_Atual", 17, "calc", fmt=K.FMT_DATA,
            formula='=IF($A{r}="","",MIN(EOMONTH(' + base_ini + ',0),' + base_ini + '+$H{r}-1))'),
        Col("Inicio_Ciclo_Atual", 16, "calc", fmt=K.FMT_DATA,
            formula='=IF($A{r}="","",EDATE($O{r},-1)+1)'),
        Col("Vencimento_Fatura_Atual", 18, "calc", fmt=K.FMT_DATA,
            formula='=IF($A{r}="","",MIN(EOMONTH(' + base_venc + ',0),' + base_venc + '+$I{r}-1))'),
        Col("Fatura_Atual", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_OrigTipo,"CARTAO",L_OrigID,$A{r},L_Data,">="&$P{r},L_Data,"<="&$O{r},L_Status,"REALIZADO")+SUMIFS(L_Valor,L_OrigTipo,"CARTAO",L_OrigID,$A{r},L_Data,">="&$P{r},L_Data,"<="&$O{r},L_Status,"PROJETADO"))'),
        Col("Faturas_Futuras", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_OrigTipo,"CARTAO",L_OrigID,$A{r},L_Data,">"&$O{r},L_Status,"PROJETADO"))'),
        Col("Parcelas_Futuras", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_OrigTipo,"CARTAO",L_OrigID,$A{r},L_Status,"PROJETADO",L_ParcID,"<>"))'),
        Col("Comprometimento_Futuro", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$L{r}+$S{r})'),
        Col("Status", 12, "in", conv=lambda x: x[8], dv="LST_STCAD"),
        Col("Observacao", 36, "in", conv=lambda x: x[9]),
    ]
    ws = esqueleto(L, "CAD_Cartoes", "CAD_CARTOES - CARTOES DE CREDITO", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.CARTOES), n, "CAR")

    # ---------------- INVESTIMENTOS ----------------
    cols = [
        Col("ID_Investimento", 15, "auto"),
        Col("Nome", 32, "in", conv=lambda x: x[1]),
        Col("Instituicao", 24, "in", conv=lambda x: INS_NOME[x[2]], dv="INS_NOME"),
        Col("ID_Instituicao", 14, "calc", formula=_f_inst("{r}")),
        Col("Tipo", 17, "in", conv=lambda x: x[3], dv="LST_TIPOINV"),
        Col("Valor_Inicial_Aplicado", 18, "in", conv=lambda x: x[4], fmt=K.FMT_GS),
        Col("Data_Inicial", 13, "in", conv=lambda x: x[5], fmt=K.FMT_DATA),
        Col("Base_Rendimento", 16, "in", conv=lambda x: x[6], dv="LST_BASEREND"),
        Col("Taxa", 11, "in", conv=lambda x: x[7], fmt=K.FMT_PCT2),
        Col("Prazo_Meses", 12, "in", conv=lambda x: x[8], fmt="0"),
        Col("Data_Vencimento", 15, "calc", fmt=K.FMT_DATA,
            formula='=IF($A{r}="","",IF($J{r}>0,EDATE($G{r},$J{r}),""))'),
        Col("Aportes_Realizados", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_DestTipo,"INVESTIMENTO",L_DestID,$A{r},L_Tipo,"APORTE",L_Status,"REALIZADO"))'),
        Col("Resgates_Realizados", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_OrigTipo,"INVESTIMENTO",L_OrigID,$A{r},L_Status,"REALIZADO"))'),
        Col("Rendimento_Realizado", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_DestTipo,"INVESTIMENTO",L_DestID,$A{r},L_Tipo,"RENDIMENTO",L_Status,"REALIZADO"))'),
        Col("Saldo_Atual", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$F{r}+$L{r}+$N{r}-$M{r})'),
        Col("Principal_Atual", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",MAX(0,$F{r}+$L{r}-$M{r}))'),
        Col("Rendimento_No_Saldo", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$O{r}-$P{r})'),
        Col("Meses_Ate_Horizonte", 16, "calc", fmt="0",
            formula='=IF($A{r}="","",IF($K{r}="",12,MAX(0,(YEAR($K{r})-YEAR(CFG_HOJE))*12+MONTH($K{r})-MONTH(CFG_HOJE))))'),
        Col("Rendimento_Projetado", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_DestTipo,"INVESTIMENTO",L_DestID,$A{r},L_Tipo,"RENDIMENTO",L_Status,"PROJETADO"))'),
        Col("Rendimento_Estimado_Taxa", 19, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,IF($H{r}="Taxa anual",ROUND($O{r}*((1+$I{r})^($R{r}/12)-1),0),IF($H{r}="Taxa mensal",ROUND($O{r}*((1+$I{r})^$R{r}-1),0),IF($H{r}="Valor fixo",$I{r},0))))'),
        Col("Valor_Estimado_Vencimento", 19, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$O{r}+$S{r}+$V{r}-$W{r})'),
        Col("Aportes_Projetados", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_DestTipo,"INVESTIMENTO",L_DestID,$A{r},L_Tipo,"APORTE",L_Status,"PROJETADO"))'),
        Col("Resgates_Projetados", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_OrigTipo,"INVESTIMENTO",L_OrigID,$A{r},L_Status,"PROJETADO"))'),
        Col("Status", 12, "in", conv=lambda x: x[9], dv="LST_STCAD"),
        Col("Observacao", 46, "in", conv=lambda x: x[10]),
    ]
    ws = esqueleto(L, "CAD_Investimentos", "CAD_INVESTIMENTOS - FONDOS MUTUOS, AHORRO A PLAZO E OUTROS", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.INVESTIMENTOS), n, "INV")

    # ---------------- CATEGORIAS ----------------
    def _tot(col_val, status):
        return ('SUMIFS(L_VDesp,L_CatID,$A{r},L_Ano,CFG_ANO,L_Status,"%s")'
                '+SUMIFS(L_VRec,L_CatID,$A{r},L_Ano,CFG_ANO,L_Status,"%s")'
                '+SUMIFS(L_VRend,L_CatID,$A{r},L_Ano,CFG_ANO,L_Status,"%s")' % (status, status, status))

    cols = [
        Col("ID_Categoria", 14, "auto"),
        Col("Nome", 30, "in", conv=lambda x: x[1]),
        Col("Tipo_Padrao", 14, "in", conv=lambda x: x[2]),
        Col("Qtd_Subcategorias", 16, "calc", fmt="0",
            formula='=IF($A{r}="","",COUNTIF(SUB_CATID,$A{r}))'),
        Col("Realizado_Ano", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,' + _tot("", "REALIZADO") + ")"),
        Col("Projetado_Ano", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,' + _tot("", "PROJETADO") + ")"),
        Col("Status", 12, "in", conv=lambda x: "Ativo", dv="LST_STCAD"),
        Col("Observacao", 34, "in", conv=lambda x: ""),
    ]
    ws = esqueleto(L, "CAD_Categorias", "CAD_CATEGORIAS - CLASSIFICACAO DE RECEITAS E DESPESAS", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.CATEGORIAS, True), n, "CAT")

    # ---------------- SUBCATEGORIAS ----------------
    def _tots(status):
        return ('SUMIFS(L_VDesp,L_SubID,$A{r},L_Ano,CFG_ANO,L_Status,"%s")'
                '+SUMIFS(L_VRec,L_SubID,$A{r},L_Ano,CFG_ANO,L_Status,"%s")'
                '+SUMIFS(L_VRend,L_SubID,$A{r},L_Ano,CFG_ANO,L_Status,"%s")' % (status, status, status))

    ns = K.MAX_CAD * 3
    cols = [
        Col("ID_Subcategoria", 15, "auto"),
        Col("Categoria", 26, "in", conv=lambda x: CAT_NOME[x[1]], dv="LST_CATEGORIAS"),
        Col("ID_Categoria", 13, "calc",
            formula='=IF($B{r}="","",IFERROR(INDEX(CAT_ID,MATCH($B{r},CAT_NOME,0)),"?"))'),
        Col("Nome", 30, "in", conv=lambda x: x[2]),
        Col("Realizado_Ano", 16, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,' + _tots("REALIZADO") + ")"),
        Col("Projetado_Ano", 16, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,' + _tots("PROJETADO") + ")"),
        Col("Ordem_Na_Categoria", 16, "calc", fmt="0",
            formula='=IF($A{r}="","",COUNTIFS($C$' + str(DAT) + ':$C{r},$C{r}))'),
        Col("Status", 12, "in", conv=lambda x: "Ativo", dv="LST_STCAD"),
        Col("Observacao", 30, "in", conv=lambda x: ""),
    ]
    ws = esqueleto(L, "CAD_Subcategorias", "CAD_SUBCATEGORIAS - DETALHAMENTO DAS CATEGORIAS", cols, "2E5C8A", ns)
    preencher(ws, cols, regs(D.SUBCATEGORIAS, True), ns, "SUB")

    # ---------------- COMPROMISSOS ----------------
    cols = [
        Col("ID_Compromisso", 15, "auto"),
        Col("Nome", 32, "in", conv=lambda x: x["nome"]),
        Col("Categoria", 24, "in", conv=lambda x: CAT_NOME[x["cat"]], dv="LST_CATEGORIAS"),
        Col("ID_Categoria", 13, "calc",
            formula='=IF($B{r}="","",IFERROR(INDEX(CAT_ID,MATCH($C{r},CAT_NOME,0)),"?"))'),
        Col("Subcategoria", 24, "in", conv=lambda x: SUB_NOME[x["sub"]]),
        Col("ID_Subcategoria", 14, "calc",
            formula='=IF($E{r}="","",IFERROR(INDEX(SUB_ID,MATCH(1,INDEX((SUB_NOME=$E{r})*(SUB_CATID=$D{r}),0),0)),"?"))'),
        Col("Qtd_Parcelas", 13, "in", conv=lambda x: x["qtd"], fmt="0"),
        Col("Valor_Parcela", 16, "in", conv=lambda x: x["valor"], fmt=K.FMT_GS),
        Col("Valor_Total", 17, "calc", fmt=K.FMT_GS, formula='=IF($A{r}="","",$G{r}*$H{r})'),
        Col("Data_Primeira_Parcela", 17, "in", conv=lambda x: x["primeira"], fmt=K.FMT_DATA),
        Col("Dia_Vencimento", 13, "in", conv=lambda x: x["dia"], fmt="0"),
        Col("Periodicidade", 14, "in", conv=lambda x: x["period"], dv="LST_PERIOD"),
        Col("Parcelas_Pagas_Antes", 17, "in", conv=lambda x: x["pagas_antes"], fmt="0"),
        Col("Entidade_Tipo", 14, "in", conv=lambda x: x["ent_tipo"], dv="LST_ENT_TIPO"),
        Col("Entidade_ID", 13, "in", conv=lambda x: x["ent_id"], dv="ENT_ID"),
        Col("Parcelas_Pagas", 14, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(P_CMP,$A{r},P_STATUS,"PAGA"))'),
        Col("Parcelas_Abertas", 15, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(P_CMP,$A{r},P_STATUS,"ABERTA"))'),
        Col("Parcelas_Atrasadas", 16, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(P_CMP,$A{r},P_STATUS,"ATRASADA"))'),
        Col("Parcelas_Projetadas", 16, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(P_CMP,$A{r},P_STATUS,"PROJETADA"))'),
        Col("Parcelas_Restantes", 16, "calc", fmt="0",
            formula='=IF($A{r}="","",$G{r}-$P{r})'),
        Col("Valor_Pago", 16, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(P_VALOR,P_CMP,$A{r},P_STATUS,"PAGA"))'),
        Col("Saldo_Devedor", 16, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$I{r}-$U{r})'),
        Col("Parcela_Atual", 13, "calc",
            formula='=IF($A{r}="","",MIN($G{r},$P{r}+1)&"/"&$G{r})'),
        Col("Proximo_Vencimento", 16, "calc", fmt=K.FMT_DATA,
            formula='=IF($A{r}="","",IFERROR(INDEX(P_VENC,MATCH(1,INDEX((P_CMP=$A{r})*(P_STATUS<>"PAGA")*(P_STATUS<>"CANCELADA"),0),0)),""))'),
        Col("Status", 12, "in", conv=lambda x: x["status"], dv="LST_STCAD"),
        Col("Observacao", 56, "in", conv=lambda x: x["obs"]),
    ]
    ws = esqueleto(L, "CAD_Compromissos", "CAD_COMPROMISSOS - PARCELAMENTOS E FINANCIAMENTOS", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.COMPROMISSOS), n, "CMP")

    # ---------------- METAS ----------------
    cols = [
        Col("ID_Meta", 12, "auto"),
        Col("Nome", 28, "in", conv=lambda x: x["nome"]),
        Col("Valor_Objetivo", 17, "in", conv=lambda x: x["objetivo"], fmt=K.FMT_GS),
        Col("Data_Prazo", 13, "in", conv=lambda x: x["prazo"], fmt=K.FMT_DATA),
        Col("Conta_Vinculada_ID", 16, "in", conv=lambda x: x["conta"], dv="C_ID"),
        Col("Investimento_Vinculado_ID", 20, "in", conv=lambda x: x["inv"], dv="I_ID"),
        Col("Considerar_No_Fluxo", 17, "in", conv=lambda x: x["fluxo"], dv="LST_SIMNAO"),
        Col("Valor_Reservado", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(MM_VALSINAL,MM_META,$A{r}))'),
        Col("Pct_Concluido", 13, "calc", fmt=K.FMT_PCT,
            formula='=IF(OR($A{r}="",$C{r}=0),0,$H{r}/$C{r})'),
        Col("Valor_Faltante", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",MAX(0,$C{r}-$H{r}))'),
        Col("Meses_Restantes", 15, "calc", fmt="0",
            formula='=IF(OR($A{r}="",$D{r}=""),0,MAX(0,(YEAR($D{r})-YEAR(CFG_HOJE))*12+MONTH($D{r})-MONTH(CFG_HOJE)))'),
        Col("Aporte_Mensal_Necessario", 19, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",IF($K{r}<=0,$J{r},ROUND($J{r}/$K{r},0)))'),
        Col("Status", 12, "in", conv=lambda x: x["status"], dv="LST_STMETA"),
        Col("Data_Conclusao", 14, "in", conv=lambda x: x["conclusao"], fmt=K.FMT_DATA),
        Col("Situacao", 18, "calc",
            formula='=IF($A{r}="","",IF($M{r}="Cancelada","CANCELADA",IF($M{r}="Concluida","CONCLUIDA",'
                    'IF($H{r}>=$C{r},"OBJETIVO ATINGIDO",IF($D{r}<CFG_HOJE,"PRAZO VENCIDO",'
                    'IF($K{r}<=3,"EM RISCO","EM DIA"))))))'),
        Col("Conta_Vinculada_Nome", 24, "calc",
            formula='=IF($E{r}="","",IFERROR(INDEX(C_NOME,MATCH($E{r},C_ID,0)),"?"))'),
        Col("Observacao", 56, "in", conv=lambda x: x["obs"]),
    ]
    ws = esqueleto(L, "CAD_Metas", "CAD_METAS - METAS FINANCEIRAS (RESERVA LOGICA)", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.METAS), n, "MET")

    # ---------------- RECORRENCIAS ----------------
    passo = ('IF($E{r}="Mensal",1,IF($E{r}="Bimestral",2,IF($E{r}="Trimestral",3,'
             'IF($E{r}="Semestral",6,IF($E{r}="Anual",12,0)))))')
    dias = 'IF($E{r}="Diario",1,IF($E{r}="Semanal",7,IF($E{r}="Quinzenal",15,0)))'
    cols = [
        Col("ID_Recorrencia", 15, "auto"),
        Col("Descricao", 34, "in", conv=lambda x: x["desc"]),
        Col("Tipo_Operacao", 18, "in", conv=lambda x: x["tipo"], dv="MAP_TIPO"),
        Col("Valor", 16, "in", conv=lambda x: x["valor"], fmt=K.FMT_GS),
        Col("Periodicidade", 14, "in", conv=lambda x: x["period"], dv="LST_PERIOD"),
        Col("Dia", 7, "in", conv=lambda x: x["dia"], fmt="0"),
        Col("Data_Inicio", 13, "in", conv=lambda x: x["inicio"], fmt=K.FMT_DATA),
        Col("Qtd_Ocorrencias", 15, "in", conv=lambda x: x["ocorr"], fmt="0"),
        Col("Data_Fim_Prevista", 16, "calc", fmt=K.FMT_DATA,
            formula='=IF(OR($A{r}="",$H{r}=0),"",IF(' + passo + '>0,EDATE($G{r},($H{r}-1)*' + passo
                    + '),IF(' + dias + '>0,$G{r}+($H{r}-1)*' + dias + ',"")))'),
        Col("Categoria_ID", 13, "in", conv=lambda x: x["cat"], dv="CAT_ID"),
        Col("Subcategoria_ID", 14, "in", conv=lambda x: x["sub"], dv="SUB_ID"),
        Col("Origem_Tipo", 13, "in", conv=lambda x: x["o_tipo"], dv="LST_ENT_TIPO"),
        Col("Origem_ID", 13, "in", conv=lambda x: x["o_id"], dv="ENT_ID"),
        Col("Destino_Tipo", 13, "in", conv=lambda x: x["d_tipo"], dv="LST_ENT_TIPO"),
        Col("Destino_ID", 13, "in", conv=lambda x: x["d_id"], dv="ENT_ID"),
        Col("Forma_Pagamento", 16, "in", conv=lambda x: x["forma"], dv="LST_FORMA"),
        Col("Status", 12, "in", conv=lambda x: x["status"]),
        Col("Lanc_Gerados", 13, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(L_RecID,$A{r}))'),
        Col("Lanc_Realizados", 15, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(L_RecID,$A{r},L_Status,"REALIZADO"))'),
        Col("Lanc_Projetados", 15, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(L_RecID,$A{r},L_Status,"PROJETADO"))'),
        Col("Proxima_Ocorrencia", 16, "calc", fmt=K.FMT_DATA,
            formula='=IF($A{r}="","",IFERROR(INDEX(L_Data,MATCH(1,INDEX((L_RecID=$A{r})*(L_Status="PROJETADO"),0),0)),""))'),
        Col("Observacao", 30, "in", conv=lambda x: ""),
    ]
    ws = esqueleto(L, "CAD_Recorrencias", "CAD_RECORRENCIAS - LANCAMENTOS REPETITIVOS", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.RECORRENCIAS), n, "REC")

    # ---------------- BENS ----------------
    cols = [
        Col("ID_Bem", 12, "auto"),
        Col("Nome", 34, "in", conv=lambda x: x[1]),
        Col("Tipo", 16, "in", conv=lambda x: x[2]),
        Col("Valor_Atual", 18, "in", conv=lambda x: x[3], fmt=K.FMT_GS),
        Col("Data_Aquisicao", 15, "in", conv=lambda x: x[4], fmt=K.FMT_DATA),
        Col("ID_Compromisso_Vinculado", 20, "in", conv=lambda x: x[5], dv="CM_ID"),
        Col("Saldo_Devedor_Vinculado", 19, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,IF($F{r}="",0,IFERROR(INDEX(CM_DEVEDOR,MATCH($F{r},CM_ID,0)),0)))'),
        Col("Valor_Liquido", 18, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="","",$D{r}-$G{r})'),
        Col("Status", 12, "in", conv=lambda x: "Ativo", dv="LST_STCAD"),
        Col("Observacao", 52, "in", conv=lambda x: x[6]),
    ]
    ws = esqueleto(L, "CAD_Bens", "CAD_BENS - ATIVOS NAO FINANCEIROS (COMPLEMENTO DO PATRIMONIO)", cols, "2E5C8A", n)
    preencher(ws, cols, regs(D.BENS), n, "BEM")
    L.nome("LST_CATEGORIAS", f"CAD_Categorias!$B${DAT}:$B${DAT + n - 1}")
    L.nome("LST_COMPROMISSOS", f"CAD_Compromissos!$B${DAT}:$B${DAT + n - 1}")
    L.nome("LST_METAS_NOME", f"CAD_Metas!$B${DAT}:$B${DAT + n - 1}")


# ---------------------------------------------------------------------------
# AUX - LISTAS DERIVADAS E LISTAS DEPENDENTES
# ---------------------------------------------------------------------------
AUX_CAT_COLS = 20          # quantas categorias recebem lista dependente de subcategorias
AUX_SUB_LINHAS = 25        # subcategorias por categoria na lista dependente


def aba_aux(L):
    n = K.MAX_CAD
    ws = L.aba("AUX", "808080")
    titulo(ws, 1, "AUX - TABELAS AUXILIARES (NAO EDITAR)", 30,
           "Gerada por formulas a partir dos cadastros. Alimenta as listas suspensas, "
           "as listas dependentes de subcategoria e a resolucao de nomes em LANCAMENTOS.")
    cabecalho(ws, HDR, [("ID_Entidade", 15), ("Tipo_Entidade", 15), ("Nome_Entidade", 34),
                        ("Indice", 8)])
    blocos = [("CAD_Contas", "CONTA"), ("CAD_Cartoes", "CARTAO"),
              ("CAD_Investimentos", "INVESTIMENTO")]
    for b, (aba, tipo) in enumerate(blocos):
        for i in range(n):
            r = DAT + b * n + i
            o = DAT + i
            ws.cell(r, 1, f'=IF({aba}!$A{o}="","",{aba}!$A{o})').font = K.F_NORMAL
            ws.cell(r, 2, f'=IF({aba}!$A{o}="","","{tipo}")').font = K.F_NORMAL
            ws.cell(r, 3, f'=IF({aba}!$A{o}="","",{aba}!$B{o})').font = K.F_NORMAL
    for i in range(AUX_SUB_LINHAS):
        c = ws.cell(DAT + i, 4, i + 1)
        c.font = K.F_NORMAL
        c.alignment = K.AL_C

    # nomes por tipo de entidade (usados pela tela LANCAR via INDIRECT)
    for j, (aba, nome_rng) in enumerate([("CAD_Contas", "LST_NOMES_CONTA"),
                                         ("CAD_Cartoes", "LST_NOMES_CARTAO"),
                                         ("CAD_Investimentos", "LST_NOMES_INVESTIMENTO")]):
        col = 6 + j
        ws.cell(HDR, col, nome_rng).font = K.F_HEADER
        ws.cell(HDR, col).fill = K.FILL_HEADER
        ws.column_dimensions[CL(col)].width = 30
        for i in range(n):
            r, o = DAT + i, DAT + i
            ws.cell(r, col, f'=IF({aba}!$A{o}="","",{aba}!$B{o})').font = K.F_NORMAL
        L.nome(nome_rng, f"AUX!${CL(col)}${DAT}:${CL(col)}${DAT + n - 1}")
    ws.cell(HDR, 9, "LST_NOMES_EXTERNO").font = K.F_HEADER
    ws.cell(HDR, 9).fill = K.FILL_HEADER
    ws.cell(DAT, 9, "(Externo)").font = K.F_NORMAL
    ws.column_dimensions["I"].width = 18
    L.nome("LST_NOMES_EXTERNO", f"AUX!$I${DAT}:$I${DAT}")

    # listas dependentes: uma coluna por categoria
    col0 = 11
    for j in range(AUX_CAT_COLS):
        col = col0 + j
        cat_row = DAT + j
        ws.column_dimensions[CL(col)].width = 24
        ws.cell(3, col, f'=IF(CAD_Categorias!$A{cat_row}="","",CAD_Categorias!$B{cat_row})').font = K.F_PEQ
        c = ws.cell(HDR, col, f'=IF(CAD_Categorias!$A{cat_row}="","",CAD_Categorias!$A{cat_row})')
        c.font = K.F_HEADER
        c.fill = K.FILL_HEADER
        c.alignment = K.AL_C
        for i in range(AUX_SUB_LINHAS):
            r = DAT + i
            ws.cell(r, col,
                    f'=IFERROR(INDEX(SUB_NOME,MATCH(1,INDEX((SUB_CATID=${CL(col)}${HDR})*'
                    f'(SUB_ORDEM=$D{r}),0),0)),"")').font = K.F_NORMAL
        L.nome(f"SUB_CAT_{j + 1:06d}", f"AUX!${CL(col)}${DAT}:${CL(col)}${DAT + AUX_SUB_LINHAS - 1}")
    ws.freeze_panes = f"E{DAT}"
    return ws


# ---------------------------------------------------------------------------
# LANCAMENTOS
# ---------------------------------------------------------------------------
def aba_lancamentos(L):
    n, nf = K.MAX_LANC, K.FORM_LANC
    valid = ('=IF($B{r}="","",'
             'IF(NOT(ISNUMBER($F{r})),"ERRO: valor nao numerico",'
             'IF($F{r}<=0,"ERRO: valor deve ser maior que zero",'
             'IF(COUNTIF(MAP_TIPO,$D{r})=0,"ERRO: tipo de operacao nao cadastrado",'
             'IF(COUNTIF(LST_STATUS,$E{r})=0,"ERRO: status invalido",'
             'IF(AND($G{r}="EXTERNO",$I{r}="EXTERNO"),"ERRO: origem e destino externos",'
             'IF(AND($G{r}<>"EXTERNO",COUNTIFS(ENT_ID,$H{r},ENT_TIPO,$G{r})=0),"ERRO: origem inexistente",'
             'IF(AND($I{r}<>"EXTERNO",COUNTIFS(ENT_ID,$J{r},ENT_TIPO,$I{r})=0),"ERRO: destino inexistente",'
             'IF(AND($K{r}<>"",COUNTIF(CAT_ID,$K{r})=0),"ERRO: categoria inexistente",'
             'IF(AND($L{r}<>"",COUNTIF(SUB_ID,$L{r})=0),"ERRO: subcategoria inexistente",'
             'IF(AND($L{r}<>"",$K{r}<>"",IFERROR(INDEX(SUB_CATID,MATCH($L{r},SUB_ID,0)),"")<>$K{r}),'
             '"ERRO: subcategoria nao pertence a categoria",'
             'IF(AND($O{r}<>"",COUNTIF(P_ID,$O{r})=0),"ERRO: parcela inexistente",'
             'IF(AND($N{r}<>"",COUNTIF(CM_ID,$N{r})=0),"ERRO: compromisso inexistente",'
             '"OK")))))))))))))')
    vinc = ('=IF($B{r}="","",IF(AND($N{r}="",$O{r}="",$P{r}="",$Q{r}="",$R{r}=""),"",'
            '"VINCULADO: "&TRIM(IF($N{r}<>"",$N{r}&" ","")&IF($O{r}<>"",$O{r}&" ","")'
            '&IF($P{r}<>"",$P{r}&" ","")&IF($Q{r}<>"",$Q{r}&" ","")&IF($R{r}<>"",$R{r},""))))')
    cols = [
        Col("ID_Lancamento", 13, "auto"),
        Col("Data", 11, "in", key="data", fmt=K.FMT_DATA),
        Col("Descricao", 44, "in", key="desc"),
        Col("Tipo_Operacao", 19, "in", key="tipo", dv="MAP_TIPO"),
        Col("Status", 12, "in", key="status", dv="LST_STATUS"),
        Col("Valor", 16, "in", key="valor", fmt=K.FMT_GS),
        Col("Origem_Tipo", 14, "in", key="o_tipo", dv="LST_ENT_TIPO"),
        Col("Origem_ID", 13, "in", key="o_id", dv="ENT_ID"),
        Col("Destino_Tipo", 14, "in", key="d_tipo", dv="LST_ENT_TIPO"),
        Col("Destino_ID", 13, "in", key="d_id", dv="ENT_ID"),
        Col("Categoria_ID", 13, "in", key="cat", dv="CAT_ID"),
        Col("Subcategoria_ID", 15, "in", key="sub", dv="SUB_ID"),
        Col("Forma_Pagamento", 17, "in", key="forma", dv="LST_FORMA"),
        Col("ID_Compromisso", 15, "in", key="cmp", dv="CM_ID"),
        Col("ID_Parcela", 12, "in", key="parcela"),
        Col("ID_Meta", 11, "in", key="meta", dv="M_ID"),
        Col("ID_Recorrencia", 15, "in", key="rec", dv="R_ID"),
        Col("ID_Relacionado", 15, "in", key="rel"),
        Col("Observacao", 52, "in", key="obs"),
        Col("Ano", 7, "calc", fmt="0", formula='=IF($B{r}="","",YEAR($B{r}))'),
        Col("Mes", 6, "calc", fmt="0", formula='=IF($B{r}="","",MONTH($B{r}))'),
        Col("AnoMes", 9, "calc", fmt="0",
            formula='=IF($B{r}="","",YEAR($B{r})*100+MONTH($B{r}))'),
        Col("Classificacao", 20, "calc",
            formula='=IF($D{r}="","",IFERROR(INDEX(MAP_CLASS,MATCH($D{r},MAP_TIPO,0)),"TIPO INVALIDO"))'),
        Col("Categoria", 24, "calc",
            formula='=IF($K{r}="","",IFERROR(INDEX(CAT_NOME,MATCH($K{r},CAT_ID,0)),"?"))'),
        Col("Subcategoria", 24, "calc",
            formula='=IF($L{r}="","",IFERROR(INDEX(SUB_NOME,MATCH($L{r},SUB_ID,0)),"?"))'),
        Col("Origem_Nome", 28, "calc",
            formula='=IF($G{r}="","",IF($G{r}="EXTERNO","(Externo)",IFERROR(INDEX(ENT_NOME,MATCH($H{r},ENT_ID,0)),"?")))'),
        Col("Destino_Nome", 28, "calc",
            formula='=IF($I{r}="","",IF($I{r}="EXTERNO","(Externo)",IFERROR(INDEX(ENT_NOME,MATCH($J{r},ENT_ID,0)),"?")))'),
        Col("Entrada_Caixa", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($B{r}="",0,IF(AND($I{r}="CONTA",$G{r}<>"CONTA"),$F{r},0))'),
        Col("Saida_Caixa", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($B{r}="",0,IF(AND($G{r}="CONTA",$I{r}<>"CONTA"),$F{r},0))'),
        Col("Vlr_Receita", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($W{r}="RECEITA",$F{r},0)'),
        Col("Vlr_Despesa", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($W{r}="DESPESA",$F{r},0)'),
        Col("Vlr_Rend_Fin", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($W{r}="RECEITA_FINANCEIRA",$F{r},0)'),
        Col("Efeito_Conta", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($B{r}="",0,IF($I{r}="CONTA",$F{r},0)-IF($G{r}="CONTA",$F{r},0))'),
        Col("Efeito_Cartao", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($B{r}="",0,IF($G{r}="CARTAO",$F{r},0)-IF($I{r}="CARTAO",$F{r},0))'),
        Col("Efeito_Investim", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($B{r}="",0,IF($I{r}="INVESTIMENTO",$F{r},0)-IF($G{r}="INVESTIMENTO",$F{r},0))'),
        Col("Validacao", 36, "calc", formula=valid),
        Col("Vinculos", 34, "calc", formula=vinc),
    ]
    aviso = (AVISO + "  As colunas cinzas sao calculadas - nao digite nelas. Para anular um "
             "lancamento troque o Status para CANCELADO (o historico permanece e o efeito "
             "financeiro desaparece).")
    ws = esqueleto(L, "LANCAMENTOS", "LANCAMENTOS - LIVRO UNICO DE OPERACOES", cols, "C00000", n, aviso)
    preencher(ws, cols, regs(D.LANCAMENTOS), nf, "LAN")
    # formatacao condicional
    fim = DAT + nf - 1
    ws.conditional_formatting.add(f"E{DAT}:E{fim}", CellIsRule(
        operator="equal", formula=['"CANCELADO"'], fill=K.FILL_ERRO))
    ws.conditional_formatting.add(f"E{DAT}:E{fim}", CellIsRule(
        operator="equal", formula=['"PROJETADO"'], fill=K.FILL_ALERTA))
    ws.conditional_formatting.add(f"AJ{DAT}:AJ{fim}", FormulaRule(
        formula=[f'AND($AJ{DAT}<>"",$AJ{DAT}<>"OK")'], fill=K.FILL_ERRO, font=Font(
            name=K.FONTE, size=10, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"AK{DAT}:AK{fim}", FormulaRule(
        formula=[f'$AK{DAT}<>""'], fill=K.FILL_ALERTA))
    return ws


# ---------------------------------------------------------------------------
# PARCELAS
# ---------------------------------------------------------------------------
def aba_parcelas(L):
    n, nf = K.MAX_PARC, K.FORM_PARC
    cols = [
        Col("ID_Parcela", 12, "auto"),
        Col("ID_Compromisso", 15, "in", key="cmp", dv="CM_ID"),
        Col("Num_Parcela", 12, "in", key="num", fmt="0"),
        Col("Data_Vencimento", 15, "in", key="venc", fmt=K.FMT_DATA),
        Col("Valor_Parcela", 16, "in", key="valor", fmt=K.FMT_GS),
        Col("Observacao", 26, "in", key="obs"),
        Col("Compromisso", 32, "calc",
            formula='=IF($B{r}="","",IFERROR(INDEX(CM_NOME,MATCH($B{r},CM_ID,0)),"?"))'),
        Col("Total_Parcelas", 14, "calc", fmt="0",
            formula='=IF($B{r}="","",IFERROR(INDEX(CM_QTD,MATCH($B{r},CM_ID,0)),""))'),
        Col("Rotulo", 10, "calc", formula='=IF($A{r}="","",$C{r}&"/"&$H{r})'),
        Col("Pago_Antes_Sistema", 17, "calc",
            formula='=IF($A{r}="","",IF($C{r}<=IFERROR(INDEX(CM_PAGASANTES,MATCH($B{r},CM_ID,0)),0),"SIM","NAO"))'),
        Col("Pagto_Realizado", 16, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,SUMIFS(L_Valor,L_ParcID,$A{r},L_Status,"REALIZADO"))'),
        Col("Data_Pagamento", 15, "calc", fmt=K.FMT_DATA,
            formula='=IF($A{r}="","",IFERROR(INDEX(L_Data,MATCH(1,INDEX((L_ParcID=$A{r})*(L_Status="REALIZADO"),0),0)),""))'),
        Col("Qtd_Lanc_Realizados", 17, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(L_ParcID,$A{r},L_Status,"REALIZADO"))'),
        Col("Qtd_Lanc_Projetados", 17, "calc", fmt="0",
            formula='=IF($A{r}="",0,COUNTIFS(L_ParcID,$A{r},L_Status,"PROJETADO"))'),
        Col("Status_Parcela", 15, "calc",
            formula='=IF($A{r}="","",IF($J{r}="SIM","PAGA",IF($M{r}>0,"PAGA",'
                    'IF($D{r}<CFG_HOJE,"ATRASADA",IF($D{r}<=EOMONTH(CFG_HOJE,0),"ABERTA","PROJETADA")))))'),
        Col("Dias_Atraso", 12, "calc", fmt="0",
            formula='=IF($O{r}="ATRASADA",CFG_HOJE-$D{r},0)'),
        Col("Diferenca_Valor", 15, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,IF(AND($O{r}="PAGA",$J{r}="NAO"),$K{r}-$E{r},0))'),
    ]
    aviso = (AVISO + "  O cronograma e gerado uma unica vez por compromisso. Use a aba "
             "GERADOR para produzir o cronograma de um novo compromisso e cole o resultado "
             "(colunas B ate E) no fim desta lista.")
    ws = esqueleto(L, "PARCELAS", "PARCELAS - CRONOGRAMA DE TODOS OS COMPROMISSOS", cols, "C00000", n, aviso)
    preencher(ws, cols, regs(D.PARCELAS), nf, "PAR")
    fim = DAT + nf - 1
    for st, fill in [("PAGA", K.FILL_OK), ("ATRASADA", K.FILL_ERRO), ("ABERTA", K.FILL_ALERTA)]:
        ws.conditional_formatting.add(f"O{DAT}:O{fim}", CellIsRule(
            operator="equal", formula=[f'"{st}"'], fill=fill))
    return ws


# ---------------------------------------------------------------------------
# MOV_METAS
# ---------------------------------------------------------------------------
def aba_mov_metas(L):
    n, nf = K.MAX_MOV, K.FORM_MOV
    cols = [
        Col("ID_Movimento", 13, "auto"),
        Col("Data", 11, "in", key="data", fmt=K.FMT_DATA),
        Col("ID_Meta", 11, "in", key="meta", dv="M_ID"),
        Col("Tipo_Movimento", 22, "in", key="tipo", dv="LST_TIPOMOV"),
        Col("Valor", 16, "in", key="valor", fmt=K.FMT_GS),
        Col("Destino_Retirada", 20, "in", key="destino", dv="LST_DESTRET"),
        Col("ID_Meta_Destino", 16, "in", key="meta_dest", dv="M_ID"),
        Col("Observacao", 58, "in", key="obs"),
        Col("Meta_Nome", 26, "calc",
            formula='=IF($C{r}="","",IFERROR(INDEX(M_NOME,MATCH($C{r},M_ID,0)),"?"))'),
        Col("Sinal", 8, "calc", fmt="0",
            formula='=IF($D{r}="",0,IFERROR(INDEX(MAP_SINALMOV,MATCH($D{r},LST_TIPOMOV,0)),0))'),
        Col("Valor_Com_Sinal", 17, "calc", fmt=K.FMT_GS,
            formula='=IF($A{r}="",0,$E{r}*$J{r})'),
        Col("Conta_Vinculada", 15, "calc",
            formula='=IF($C{r}="","",IFERROR(INDEX(M_CONTA,MATCH($C{r},M_ID,0)),""))'),
        Col("Validacao", 34, "calc",
            formula='=IF($B{r}="","",IF($E{r}<=0,"ERRO: valor deve ser maior que zero",'
                    'IF(COUNTIF(M_ID,$C{r})=0,"ERRO: meta inexistente",'
                    'IF(COUNTIF(LST_TIPOMOV,$D{r})=0,"ERRO: tipo de movimento invalido",'
                    'IF(AND($D{r}="TRANSF_SAIDA",COUNTIF(M_ID,$G{r})=0),"ERRO: informe a meta de destino",'
                    'IF(AND($J{r}=-1,$F{r}=""),"ATENCAO: informe o destino da retirada","OK"))))))'),
    ]
    aviso = (AVISO + "  Esta aba registra apenas RESERVAS LOGICAS. O dinheiro continua na "
             "conta; nenhuma linha daqui altera o saldo bancario.")
    ws = esqueleto(L, "MOV_METAS", "MOV_METAS - MOVIMENTACAO DAS RESERVAS DE METAS", cols, "C00000", n, aviso)
    preencher(ws, cols, regs(D.MOV_METAS), nf, "MOV")
    return ws


# ---------------------------------------------------------------------------
# FLUXO DE CAIXA (diario, semanal, mensal, anual)
# ---------------------------------------------------------------------------
SALDO_BASE = ("SUM(C_SALDOINI)"
              '+SUMIFS(L_Ent,L_Data,"<"&{d},L_Status,"REALIZADO")'
              '-SUMIFS(L_Sai,L_Data,"<"&{d},L_Status,"REALIZADO")'
              '+SUMIFS(L_Ent,L_Data,"<"&{d},L_Status,"PROJETADO")'
              '-SUMIFS(L_Sai,L_Data,"<"&{d},L_Status,"PROJETADO")')


def _sum_per(campo, ini, fim, status):
    return f'SUMIFS({campo},L_Data,">="&{ini},L_Data,"<="&{fim},L_Status,"{status}")'


def aba_fluxo(L):
    ws = L.aba("FLUXO_CAIXA", "375623")
    MESES, DIAS, SEMANAS, ANOS = 30, 60, 30, 6
    r = titulo(ws, 1, "FLUXO_CAIXA - ENTRADAS E SAIDAS DE CAIXA", 13,
               "Somente movimentos que atravessam uma CONTA. Transferencias entre contas "
               "proprias nao aparecem como entrada nem como saida (nao ha duplicidade). "
               "Compra no cartao so vira saida de caixa quando a fatura e paga.")
    ws.column_dimensions["A"].width = 15
    for c in "BCDEFGHIJKLM":
        ws.column_dimensions[c].width = 17

    # ---- parametros do bloco mensal
    r += 1
    rotulo(ws, r, 1, "Mes inicial da grade")
    ws.cell(r, 2, "=EOMONTH(CFG_INICIO,0)+1").number_format = K.FMT_DATA
    ws.cell(r, 2).fill = K.FILL_CALC
    base = f"$B${r}"
    rotulo(ws, r, 4, "Reservas de metas consideradas no fluxo")
    ws.cell(r, 6, '=IF(CFG_RESERVA_FLUXO="SIM",SUMIFS(M_RESERV,M_FLUXO,"SIM",M_STATUS,"Ativa"),0)'
            ).number_format = K.FMT_GS
    ws.cell(r, 6).fill = K.FILL_CALC
    reserva = f"$F${r}"
    r += 2

    # ---- 1) MENSAL
    r = secao(ws, r, "1. FLUXO DE CAIXA MENSAL (REALIZADO + PROJETADO)", 13)
    hdr = [("Competencia", 15), ("Data_Inicio", 13), ("Data_Fim", 13), ("Saldo_Inicial", 17),
           ("Entradas_Realizadas", 17), ("Saidas_Realizadas", 17), ("Liquido_Realizado", 17),
           ("Entradas_Projetadas", 17), ("Saidas_Projetadas", 17), ("Liquido_Projetado", 17),
           ("Saldo_Final", 17), ("Reservado_Metas", 15), ("Saldo_Livre", 17)]
    r = cabecalho(ws, r, hdr)
    r_mes_ini = r
    for i in range(MESES):
        rr = r + i
        ws.cell(rr, 1, f'=YEAR($B{rr})*100+MONTH($B{rr})').number_format = "0"
        ws.cell(rr, 2, f"=EDATE({base},{i})").number_format = K.FMT_DATA
        ws.cell(rr, 3, f"=EOMONTH($B{rr},0)").number_format = K.FMT_DATA
        if i == 0:
            ws.cell(rr, 4, "=" + SALDO_BASE.format(d=f"$B{rr}"))
        else:
            ws.cell(rr, 4, f"=$K{rr - 1}")
        ws.cell(rr, 5, "=" + _sum_per("L_Ent", f"$B{rr}", f"$C{rr}", "REALIZADO"))
        ws.cell(rr, 6, "=" + _sum_per("L_Sai", f"$B{rr}", f"$C{rr}", "REALIZADO"))
        ws.cell(rr, 7, f"=$E{rr}-$F{rr}")
        ws.cell(rr, 8, "=" + _sum_per("L_Ent", f"$B{rr}", f"$C{rr}", "PROJETADO"))
        ws.cell(rr, 9, "=" + _sum_per("L_Sai", f"$B{rr}", f"$C{rr}", "PROJETADO"))
        ws.cell(rr, 10, f"=$H{rr}-$I{rr}")
        ws.cell(rr, 11, f"=$D{rr}+$G{rr}+$J{rr}")
        ws.cell(rr, 12, f"={reserva}")
        ws.cell(rr, 13, f"=$K{rr}-$L{rr}")
        for c in range(4, 14):
            ws.cell(rr, c).number_format = K.FMT_GS
        for c in range(1, 14):
            ws.cell(rr, c).font = K.F_NORMAL
            ws.cell(rr, c).border = K.BORDA
            ws.cell(rr, c).fill = K.FILL_CALC
    ws.conditional_formatting.add(f"K{r}:K{r + MESES - 1}", CellIsRule(
        operator="lessThan", formula=["0"], fill=K.FILL_ERRO))
    L.nome("FX_COMP", f"FLUXO_CAIXA!$A${r}:$A${r + MESES - 1}")
    L.nome("FX_INI", f"FLUXO_CAIXA!$B${r}:$B${r + MESES - 1}")
    L.nome("FX_SALDOFIM", f"FLUXO_CAIXA!$K${r}:$K${r + MESES - 1}")
    L.nome("FX_ENTR", f"FLUXO_CAIXA!$E${r}:$E${r + MESES - 1}")
    L.nome("FX_SAID", f"FLUXO_CAIXA!$F${r}:$F${r + MESES - 1}")
    r += MESES + 1

    # ---- 2) DIARIO
    r = secao(ws, r, "2. FLUXO DE CAIXA DIARIO (15 DIAS ANTES ATE 44 DIAS DEPOIS DA DATA DE REFERENCIA)", 13)
    hdr = [("Data", 15), ("Entradas_Realizadas", 17), ("Saidas_Realizadas", 17),
           ("Entradas_Projetadas", 17), ("Saidas_Projetadas", 17), ("Liquido_Dia", 17),
           ("Saldo_Acumulado", 17)]
    r = cabecalho(ws, r, hdr)
    for i in range(DIAS):
        rr = r + i
        ws.cell(rr, 1, f"=CFG_HOJE-15+{i}").number_format = K.FMT_DATA
        ws.cell(rr, 2, '=SUMIFS(L_Ent,L_Data,$A%d,L_Status,"REALIZADO")' % rr)
        ws.cell(rr, 3, '=SUMIFS(L_Sai,L_Data,$A%d,L_Status,"REALIZADO")' % rr)
        ws.cell(rr, 4, '=SUMIFS(L_Ent,L_Data,$A%d,L_Status,"PROJETADO")' % rr)
        ws.cell(rr, 5, '=SUMIFS(L_Sai,L_Data,$A%d,L_Status,"PROJETADO")' % rr)
        ws.cell(rr, 6, f"=$B{rr}-$C{rr}+$D{rr}-$E{rr}")
        if i == 0:
            ws.cell(rr, 7, "=" + SALDO_BASE.format(d=f"$A{rr}") + f"+$F{rr}")
        else:
            ws.cell(rr, 7, f"=$G{rr - 1}+$F{rr}")
        for c in range(2, 8):
            ws.cell(rr, c).number_format = K.FMT_GS
        for c in range(1, 8):
            ws.cell(rr, c).font = K.F_NORMAL
            ws.cell(rr, c).border = K.BORDA
            ws.cell(rr, c).fill = K.FILL_CALC
    ws.conditional_formatting.add(f"G{r}:G{r + DIAS - 1}", CellIsRule(
        operator="lessThan", formula=["0"], fill=K.FILL_ERRO))
    r += DIAS + 1

    # ---- 3) SEMANAL
    r = secao(ws, r, "3. FLUXO DE CAIXA SEMANAL (30 SEMANAS A PARTIR DE 2 SEMANAS ANTES DA REFERENCIA)", 13)
    hdr = [("Semana_Inicio", 15), ("Semana_Fim", 15), ("Entradas", 17), ("Saidas", 17),
           ("Liquido", 17), ("Saldo_Acumulado", 17)]
    r = cabecalho(ws, r, hdr)
    for i in range(SEMANAS):
        rr = r + i
        ws.cell(rr, 1, f"=CFG_HOJE-WEEKDAY(CFG_HOJE,2)+1-14+{i * 7}").number_format = K.FMT_DATA
        ws.cell(rr, 2, f"=$A{rr}+6").number_format = K.FMT_DATA
        ws.cell(rr, 3, "=" + _sum_per("L_Ent", f"$A{rr}", f"$B{rr}", "REALIZADO")
                + "+" + _sum_per("L_Ent", f"$A{rr}", f"$B{rr}", "PROJETADO"))
        ws.cell(rr, 4, "=" + _sum_per("L_Sai", f"$A{rr}", f"$B{rr}", "REALIZADO")
                + "+" + _sum_per("L_Sai", f"$A{rr}", f"$B{rr}", "PROJETADO"))
        ws.cell(rr, 5, f"=$C{rr}-$D{rr}")
        if i == 0:
            ws.cell(rr, 6, "=" + SALDO_BASE.format(d=f"$A{rr}") + f"+$E{rr}")
        else:
            ws.cell(rr, 6, f"=$F{rr - 1}+$E{rr}")
        for c in range(3, 7):
            ws.cell(rr, c).number_format = K.FMT_GS
        for c in range(1, 7):
            ws.cell(rr, c).font = K.F_NORMAL
            ws.cell(rr, c).border = K.BORDA
            ws.cell(rr, c).fill = K.FILL_CALC
    r += SEMANAS + 1

    # ---- 4) ANUAL
    r = secao(ws, r, "4. FLUXO DE CAIXA ANUAL", 13)
    hdr = [("Ano", 15), ("Entradas_Realizadas", 17), ("Saidas_Realizadas", 17),
           ("Entradas_Projetadas", 17), ("Saidas_Projetadas", 17), ("Liquido_Ano", 17),
           ("Saldo_Final_Ano", 17)]
    r = cabecalho(ws, r, hdr)
    for i in range(ANOS):
        rr = r + i
        ws.cell(rr, 1, f"=YEAR(CFG_INICIO)+{i}").number_format = "0"
        ini, fim = f"DATE($A{rr},1,1)", f"DATE($A{rr},12,31)"
        ws.cell(rr, 2, "=" + _sum_per("L_Ent", ini, fim, "REALIZADO"))
        ws.cell(rr, 3, "=" + _sum_per("L_Sai", ini, fim, "REALIZADO"))
        ws.cell(rr, 4, "=" + _sum_per("L_Ent", ini, fim, "PROJETADO"))
        ws.cell(rr, 5, "=" + _sum_per("L_Sai", ini, fim, "PROJETADO"))
        ws.cell(rr, 6, f"=$B{rr}-$C{rr}+$D{rr}-$E{rr}")
        ws.cell(rr, 7, "=" + SALDO_BASE.format(d=f"{fim}+1"))
        for c in range(2, 8):
            ws.cell(rr, c).number_format = K.FMT_GS
        for c in range(1, 8):
            ws.cell(rr, c).font = K.F_NORMAL
            ws.cell(rr, c).border = K.BORDA
            ws.cell(rr, c).fill = K.FILL_CALC
    ws.freeze_panes = "A6"
    return ws


def _grade(ws, r, hdr, nlin, gen, largura_a=15):
    """Escreve um cabecalho e nlin linhas de formulas. gen(i, rr) -> lista de (col, formula, fmt)."""
    r = cabecalho(ws, r, hdr)
    for i in range(nlin):
        rr = r + i
        for col, f, fmt in gen(i, rr):
            c = ws.cell(rr, col, f)
            c.font = K.F_NORMAL
            c.border = K.BORDA
            c.fill = K.FILL_CALC
            if fmt:
                c.number_format = fmt
    return r + nlin


# ---------------------------------------------------------------------------
# PROJECAO
# ---------------------------------------------------------------------------
def aba_projecao(L):
    ws = L.aba("PROJECAO", "375623")
    N = 24
    r = titulo(ws, 1, "PROJECAO - FUTURO FINANCEIRO MES A MES", 14,
               "Considera receitas e despesas projetadas, parcelas, faturas de cartao, "
               "aportes, resgates, vencimentos de investimento e recorrencias. As colunas "
               "de projecao nunca se misturam com o realizado: o que ja aconteceu no mes "
               "corrente entra apenas na coluna Liquido_Realizado_no_Mes, para que o saldo "
               "acumulado parta da posicao verdadeira.")
    r += 1
    hdr = [("Competencia", 14), ("Data_Inicio", 13), ("Data_Fim", 13),
           ("Receitas_Proj", 16), ("Rendimentos_Proj", 16), ("Despesas_Proj", 16),
           ("Parcelas_Proj", 16), ("Faturas_Cartao_Proj", 17), ("Aportes_Proj", 15),
           ("Resgates_Proj", 15), ("Entradas_Caixa_Proj", 16), ("Saidas_Caixa_Proj", 16),
           ("Liquido_Projetado", 16), ("Liquido_Realizado_no_Mes", 18),
           ("Saldo_Projetado", 17)]

    def gen(i, rr):
        ini, fim = f"$B{rr}", f"$C{rr}"
        base_mes = "DATE(YEAR(CFG_HOJE),MONTH(CFG_HOJE),1)"
        out = [
            (1, f"=YEAR($B{rr})*100+MONTH($B{rr})", "0"),
            (2, f"=EDATE({base_mes},{i})", K.FMT_DATA),
            (3, f"=EOMONTH($B{rr},0)", K.FMT_DATA),
            (4, "=" + _sum_per("L_VRec", ini, fim, "PROJETADO"), K.FMT_GS),
            (5, "=" + _sum_per("L_VRend", ini, fim, "PROJETADO"), K.FMT_GS),
            (6, "=" + _sum_per("L_VDesp", ini, fim, "PROJETADO"), K.FMT_GS),
            (7, f'=SUMIFS(L_Valor,L_Data,">="&{ini},L_Data,"<="&{fim},L_Status,"PROJETADO",L_ParcID,"<>")', K.FMT_GS),
            (8, f'=SUMIFS(L_Valor,L_Data,">="&{ini},L_Data,"<="&{fim},L_Status,"PROJETADO",L_Tipo,"PAGAMENTO_CARTAO")', K.FMT_GS),
            (9, f'=SUMIFS(L_Valor,L_Data,">="&{ini},L_Data,"<="&{fim},L_Status,"PROJETADO",L_Tipo,"APORTE")', K.FMT_GS),
            (10, f'=SUMIFS(L_Valor,L_Data,">="&{ini},L_Data,"<="&{fim},L_Status,"PROJETADO",L_Tipo,"RESGATE")', K.FMT_GS),
            (11, "=" + _sum_per("L_Ent", ini, fim, "PROJETADO"), K.FMT_GS),
            (12, "=" + _sum_per("L_Sai", ini, fim, "PROJETADO"), K.FMT_GS),
            (13, f"=$K{rr}-$L{rr}", K.FMT_GS),
            (14, "=" + _sum_per("L_Ent", ini, fim, "REALIZADO") + "-"
             + _sum_per("L_Sai", ini, fim, "REALIZADO"), K.FMT_GS),
        ]
        if i == 0:
            out.append((15, "=" + SALDO_BASE.format(d=f"$B{rr}") + f"+$N{rr}+$M{rr}", K.FMT_GS))
        else:
            out.append((15, f"=$O{rr - 1}+$N{rr}+$M{rr}", K.FMT_GS))
        return out

    r = _grade(ws, r, hdr, N, gen)
    ws.conditional_formatting.add(f"O{r - N}:O{r - 1}", CellIsRule(
        operator="lessThan", formula=["0"], fill=K.FILL_ERRO))
    L.nome("PJ_COMP", f"PROJECAO!$A${r - N}:$A${r - 1}")
    L.nome("PJ_SALDO", f"PROJECAO!$O${r - N}:$O${r - 1}")
    L.nome("PJ_INI", f"PROJECAO!$B${r - N}:$B${r - 1}")
    L.nome("PJ_FIM", f"PROJECAO!$C${r - N}:$C${r - 1}")
    L.nome("PJ_SAIDAS", f"PROJECAO!$L${r - N}:$L${r - 1}")
    L.nome("PJ_PARCELAS", f"PROJECAO!$G${r - N}:$G${r - 1}")
    L.nome("PJ_FATURAS", f"PROJECAO!$H${r - N}:$H${r - 1}")
    r += 1

    r = secao(ws, r, "COMPROMISSOS EM ANDAMENTO", 14)
    hdr = [("Compromisso", 32), ("Parcela_Atual", 14), ("Parcelas_Restantes", 16),
           ("Valor_Parcela", 16), ("Saldo_Devedor", 17), ("Proximo_Vencimento", 17),
           ("Status", 12)]
    ini = DAT
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Compromissos!$A{ini + i}="","",CAD_Compromissos!$B{ini + i})', None),
        (2, f'=IF(CAD_Compromissos!$A{ini + i}="","",CAD_Compromissos!$W{ini + i})', None),
        (3, f'=IF(CAD_Compromissos!$A{ini + i}="","",CAD_Compromissos!$T{ini + i})', "0"),
        (4, f'=IF(CAD_Compromissos!$A{ini + i}="","",CAD_Compromissos!$H{ini + i})', K.FMT_GS),
        (5, f'=IF(CAD_Compromissos!$A{ini + i}="","",CAD_Compromissos!$V{ini + i})', K.FMT_GS),
        (6, f'=IF(CAD_Compromissos!$A{ini + i}="","",CAD_Compromissos!$X{ini + i})', K.FMT_DATA),
        (7, f'=IF(CAD_Compromissos!$A{ini + i}="","",CAD_Compromissos!$Y{ini + i})', None),
    ])
    r += 1
    r = secao(ws, r, "INVESTIMENTOS E VENCIMENTOS", 14)
    hdr = [("Investimento", 32), ("Tipo", 16), ("Saldo_Atual", 17), ("Data_Vencimento", 15),
           ("Rendimento_Projetado", 18), ("Rendimento_Estimado_Taxa", 20),
           ("Valor_Estimado_Vencimento", 20)]
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Investimentos!$A{ini + i}="","",CAD_Investimentos!$B{ini + i})', None),
        (2, f'=IF(CAD_Investimentos!$A{ini + i}="","",CAD_Investimentos!$E{ini + i})', None),
        (3, f'=IF(CAD_Investimentos!$A{ini + i}="","",CAD_Investimentos!$O{ini + i})', K.FMT_GS),
        (4, f'=IF(CAD_Investimentos!$A{ini + i}="","",CAD_Investimentos!$K{ini + i})', K.FMT_DATA),
        (5, f'=IF(CAD_Investimentos!$A{ini + i}="","",CAD_Investimentos!$S{ini + i})', K.FMT_GS),
        (6, f'=IF(CAD_Investimentos!$A{ini + i}="","",CAD_Investimentos!$T{ini + i})', K.FMT_GS),
        (7, f'=IF(CAD_Investimentos!$A{ini + i}="","",CAD_Investimentos!$U{ini + i})', K.FMT_GS),
    ])
    ws.freeze_panes = "A6"
    return ws


# ---------------------------------------------------------------------------
# PATRIMONIO
# ---------------------------------------------------------------------------
def aba_patrimonio(L):
    ws = L.aba("PATRIMONIO", "375623")
    N = 30
    for c, w in [("A", 40), ("B", 20), ("C", 20), ("D", 20), ("E", 20), ("F", 20),
                 ("G", 20), ("H", 20), ("I", 20), ("J", 20)]:
        ws.column_dimensions[c].width = w
    r = titulo(ws, 1, "PATRIMONIO - ATIVOS, PASSIVOS E PATRIMONIO LIQUIDO", 10,
               "Patrimonio Financeiro considera apenas contas, investimentos, cartoes e "
               "compromissos. Patrimonio Total soma tambem os bens cadastrados em CAD_Bens.")
    r += 1
    r = secao(ws, r, "1. POSICAO ATUAL (REALIZADO)", 10)
    linhas = [
        ("ATIVOS", None, None, True),
        ("Contas bancarias e dinheiro fisico", "=SUM(C_SALDO)", K.FMT_GS, False),
        ("Investimentos", "=SUM(I_SALDO)", K.FMT_GS, False),
        ("Subtotal - ativos financeiros", "=$B{a}+$B{b}", K.FMT_GS, True),
        ("Bens (ativos nao financeiros)", "=SUM(B_VALOR)", K.FMT_GS, False),
        ("TOTAL DE ATIVOS", "=$B{c}+$B{d}", K.FMT_GS, True),
        ("PASSIVOS", None, None, True),
        ("Dividas de cartao de credito", "=SUM(K_DIVIDA)", K.FMT_GS, False),
        ("Saldo devedor de compromissos", "=SUM(CM_DEVEDOR)", K.FMT_GS, False),
        ("TOTAL DE PASSIVOS", "=$B{e}+$B{f}", K.FMT_GS, True),
        ("RESULTADO", None, None, True),
        ("Patrimonio liquido FINANCEIRO", "=$B{c}-$B{g}", K.FMT_GS, True),
        ("Patrimonio liquido TOTAL", "=$B{h}-$B{g}", K.FMT_GS, True),
        ("Memoria: valor reservado em metas", "=SUM(M_RESERV)", K.FMT_GS, False),
        ("Memoria: saldo livre (nao reservado)", "=SUM(C_DISP)", K.FMT_GS, False),
    ]
    base = r + 1
    ref = dict(a=base + 1, b=base + 2, c=base + 3, d=base + 4, e=base + 7, f=base + 8,
               g=base + 9, h=base + 5)
    r = cabecalho(ws, r, [("Componente", 40), ("Valor", 20)])
    for i, (txt, f, fmt, bold) in enumerate(linhas):
        rr = r + i
        c1 = ws.cell(rr, 1, txt)
        c1.font = K.F_BOLD if bold else K.F_NORMAL
        c1.border = K.BORDA
        if f is None:
            c1.fill = K.FILL_SUB
            ws.cell(rr, 2).fill = K.FILL_SUB
            ws.cell(rr, 2).border = K.BORDA
            continue
        c2 = ws.cell(rr, 2, f.format(**ref))
        c2.number_format = fmt
        c2.font = K.F_BOLD if bold else K.F_NORMAL
        c2.border = K.BORDA
        c2.fill = K.FILL_CALC
    L.nome("PAT_ATIVOS", f"PATRIMONIO!$B${ref['h']}")
    L.nome("PAT_PASSIVOS", f"PATRIMONIO!$B${ref['g']}")
    L.nome("PAT_PL_FIN", f"PATRIMONIO!$B${base + 11}")
    L.nome("PAT_PL_TOTAL", f"PATRIMONIO!$B${base + 12}")
    r += len(linhas) + 1

    r = secao(ws, r, "2. EVOLUCAO PATRIMONIAL MES A MES (REALIZADO + PROJETADO)", 10)
    hdr = [("Competencia", 14), ("Fim_do_Mes", 13), ("Contas", 18), ("Investimentos", 18),
           ("Bens", 18), ("Cartoes", 18), ("Compromissos", 18), ("PL_Financeiro", 19),
           ("PL_Total", 19), ("Variacao_PL", 18)]
    b0 = "EOMONTH(CFG_INICIO,0)+1"

    def acum(campo, ini_soma, rr):
        return (f"{ini_soma}"
                f'+SUMIFS({campo},L_Data,"<="&$B{rr},L_Status,"REALIZADO")'
                f'+SUMIFS({campo},L_Data,"<="&$B{rr},L_Status,"PROJETADO")')

    def gen(i, rr):
        return [
            (1, f"=YEAR($B{rr})*100+MONTH($B{rr})", "0"),
            (2, f"=EOMONTH(EDATE({b0},{i}),0)", K.FMT_DATA),
            (3, "=" + acum("L_EfConta", "SUM(C_SALDOINI)", rr), K.FMT_GS),
            (4, "=" + acum("L_EfInv", "SUM(I_VALINI)", rr), K.FMT_GS),
            (5, "=SUM(B_VALOR)", K.FMT_GS),
            (6, "=" + acum("L_EfCartao", "SUM(K_DIVINI)", rr), K.FMT_GS),
            (7, f'=SUMIFS(P_VALOR,P_VENC,">"&$B{rr})+SUMIFS(P_VALOR,P_VENC,"<="&$B{rr},P_STATUS,"ATRASADA")', K.FMT_GS),
            (8, f"=$C{rr}+$D{rr}-$F{rr}-$G{rr}", K.FMT_GS),
            (9, f"=$H{rr}+$E{rr}", K.FMT_GS),
            (10, f"=$I{rr}-$I{rr - 1}" if i else "=0", K.FMT_GS),
        ]

    r = _grade(ws, r, hdr, N, gen)
    L.nome("PT_COMP", f"PATRIMONIO!$A${r - N}:$A${r - 1}")
    L.nome("PT_PLTOTAL", f"PATRIMONIO!$I${r - N}:$I${r - 1}")
    L.nome("PT_PLFIN", f"PATRIMONIO!$H${r - N}:$H${r - 1}")
    ws.freeze_panes = "A4"
    return ws


def kv(ws, r, label, formula, fmt=None, bold=False, col=1, largura=44, obs=None):
    c = ws.cell(r, col, label)
    c.font = K.F_BOLD if bold else K.F_NORMAL
    c.border = K.BORDA
    ws.column_dimensions[CL(col)].width = largura
    v = ws.cell(r, col + 1, formula)
    v.number_format = fmt or K.FMT_GS
    v.font = K.F_BOLD if bold else K.F_NORMAL
    v.border = K.BORDA
    v.fill = K.FILL_CALC
    if obs:
        o = ws.cell(r, col + 2, obs)
        o.font = K.F_PEQ
        o.alignment = K.AL_LW
    return r + 1


# ---------------------------------------------------------------------------
# RELATORIO MENSAL
# ---------------------------------------------------------------------------
def aba_rel_mensal(L):
    ws = L.aba("REL_Mensal", "7030A0")
    for c, w in [("A", 34), ("B", 22), ("C", 22), ("D", 22), ("E", 20), ("F", 22), ("G", 22)]:
        ws.column_dimensions[c].width = w
    r = titulo(ws, 1, "REL_MENSAL - RELATORIO DO MES", 7,
               "Escolha o ano e o mes nas celulas amarelas. Todo o relatorio se recalcula.")
    r += 1
    rotulo(ws, r, 1, "Ano de analise")
    valor(ws, r, 2, K.ANO_BASE, "0", entrada=True)
    rotulo(ws, r, 3, "Mes de analise (1 a 12)")
    valor(ws, r, 4, K.HOJE.month, "0", entrada=True)
    ws.cell(r, 5, "=INDEX(LST_MESES,$D%d)" % r).font = K.F_BOLD
    L.nome("REL_ANO", f"REL_Mensal!$B${r}")
    L.nome("REL_MES", f"REL_Mensal!$D${r}")
    r += 1
    rotulo(ws, r, 1, "Periodo")
    ws.cell(r, 2, "=DATE(REL_ANO,REL_MES,1)").number_format = K.FMT_DATA
    ws.cell(r, 4, "=EOMONTH($B%d,0)" % r).number_format = K.FMT_DATA
    for cc in (2, 4):
        ws.cell(r, cc).fill = K.FILL_CALC
        ws.cell(r, cc).border = K.BORDA
    L.nome("REL_INI", f"REL_Mensal!$B${r}")
    L.nome("REL_FIM", f"REL_Mensal!$D${r}")
    r += 2

    def S(campo, status, extra=""):
        return (f'SUMIFS({campo},L_Data,">="&REL_INI,L_Data,"<="&REL_FIM,'
                f'L_Status,"{status}"{extra})')

    def linha3(rr, label, campo, extra="", bold=False):
        c = ws.cell(rr, 1, label)
        c.font = K.F_BOLD if bold else K.F_NORMAL
        c.border = K.BORDA
        for j, st in enumerate(("REALIZADO", "PROJETADO")):
            v = ws.cell(rr, 2 + j, "=" + S(campo, st, extra))
            v.number_format = K.FMT_GS
            v.font = K.F_BOLD if bold else K.F_NORMAL
            v.border = K.BORDA
            v.fill = K.FILL_CALC
        v = ws.cell(rr, 4, f"=$B{rr}+$C{rr}")
        v.number_format = K.FMT_GS
        v.font = K.F_BOLD
        v.border = K.BORDA
        v.fill = K.FILL_SUB
        return rr + 1

    r = secao(ws, r, "1. RESUMO DO MES", 7)
    r = cabecalho(ws, r, [("Indicador", 34), ("Realizado", 22), ("Projetado", 22), ("Total", 22)])
    r0 = r
    r = linha3(r, "Receitas operacionais", "L_VRec")
    r = linha3(r, "Receitas financeiras (rendimentos)", "L_VRend")
    r = linha3(r, "Despesas", "L_VDesp")
    for j, cl in enumerate("BCD"):
        c = ws.cell(r, 2 + j, f"={cl}{r0}+{cl}{r0 + 1}-{cl}{r0 + 2}")
        c.number_format = K.FMT_GS
        c.font = K.F_BOLD
        c.border = K.BORDA
        c.fill = K.FILL_SUB
    ws.cell(r, 1, "(=) RESULTADO DO MES (competencia)").font = K.F_BOLD
    ws.cell(r, 1).border = K.BORDA
    L.nome("REL_RESULTADO", f"REL_Mensal!$D${r}")
    r += 1
    r_ent = r
    r = linha3(r, "Entradas de caixa", "L_Ent")
    r = linha3(r, "Saidas de caixa", "L_Sai")
    for j, cl in enumerate("BCD"):
        c = ws.cell(r, 2 + j, f"={cl}{r_ent}-{cl}{r_ent + 1}")
        c.number_format = K.FMT_GS
        c.font = K.F_BOLD
        c.border = K.BORDA
        c.fill = K.FILL_SUB
    ws.cell(r, 1, "(=) FLUXO LIQUIDO DE CAIXA").font = K.F_BOLD
    ws.cell(r, 1).border = K.BORDA
    r += 1
    for lab, tp in [("Transferencias entre contas (neutro)", "TRANSFERENCIA"),
                    ("Aportes em investimentos (neutro)", "APORTE"),
                    ("Resgates de investimentos (neutro)", "RESGATE"),
                    ("Compras no cartao (despesa)", "COMPRA_CARTAO"),
                    ("Pagamentos de fatura (neutro)", "PAGAMENTO_CARTAO"),
                    ("Parcelas de compromissos (despesa)", "PAGAMENTO_PARCELA")]:
        r = linha3(r, lab, "L_Valor", f',L_Tipo,"{tp}"')
    r = kv(ws, r, "Saldo total das contas no fim do mes",
           "=" + SALDO_BASE.format(d="REL_FIM+1"), K.FMT_GS, bold=True)
    r += 1

    r = secao(ws, r, "2. DESPESAS POR CATEGORIA", 7)
    hdr = [("Categoria", 34), ("Realizado", 22), ("Projetado", 22), ("Total", 22),
           ("% do realizado", 16)]
    tot_ref = "SUMIFS(L_VDesp,L_Data,\">=\"&REL_INI,L_Data,\"<=\"&REL_FIM,L_Status,\"REALIZADO\")"
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Categorias!$A{DAT + i}="","",CAD_Categorias!$B{DAT + i})', None),
        (2, f'=IF(CAD_Categorias!$A{DAT + i}="",0,' + S("L_VDesp", "REALIZADO", f",L_CatID,CAD_Categorias!$A{DAT + i}") + ")", K.FMT_GS),
        (3, f'=IF(CAD_Categorias!$A{DAT + i}="",0,' + S("L_VDesp", "PROJETADO", f",L_CatID,CAD_Categorias!$A{DAT + i}") + ")", K.FMT_GS),
        (4, f"=$B{rr}+$C{rr}", K.FMT_GS),
        (5, f"=IF({tot_ref}=0,0,$B{rr}/{tot_ref})", K.FMT_PCT),
    ])
    L.nome("RM_CAT_NOME", f"REL_Mensal!$A${r - K.MAX_CAD}:$A${r - 1}")
    L.nome("RM_CAT_REAL", f"REL_Mensal!$B${r - K.MAX_CAD}:$B${r - 1}")
    r += 1

    r = secao(ws, r, "3. DESPESAS POR SUBCATEGORIA", 7)
    ns = K.MAX_CAD * 3
    hdr = [("Categoria", 34), ("Subcategoria", 22), ("Realizado", 22), ("Projetado", 22),
           ("Total", 22)]
    r = _grade(ws, r, hdr, ns, lambda i, rr: [
        (1, f'=IF(CAD_Subcategorias!$A{DAT + i}="","",CAD_Subcategorias!$B{DAT + i})', None),
        (2, f'=IF(CAD_Subcategorias!$A{DAT + i}="","",CAD_Subcategorias!$D{DAT + i})', None),
        (3, f'=IF(CAD_Subcategorias!$A{DAT + i}="",0,' + S("L_VDesp", "REALIZADO", f",L_SubID,CAD_Subcategorias!$A{DAT + i}") + ")", K.FMT_GS),
        (4, f'=IF(CAD_Subcategorias!$A{DAT + i}="",0,' + S("L_VDesp", "PROJETADO", f",L_SubID,CAD_Subcategorias!$A{DAT + i}") + ")", K.FMT_GS),
        (5, f"=$C{rr}+$D{rr}", K.FMT_GS),
    ])
    r += 1

    r = secao(ws, r, "4. MOVIMENTO POR CONTA", 7)
    hdr = [("Conta", 34), ("Entradas_Realizadas", 22), ("Saidas_Realizadas", 22),
           ("Liquido", 22), ("Saldo_no_Fim_do_Mes", 22)]
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Contas!$A{DAT + i}="","",CAD_Contas!$B{DAT + i})', None),
        (2, f'=IF(CAD_Contas!$A{DAT + i}="",0,' + S("L_Valor", "REALIZADO", f',L_DestTipo,"CONTA",L_DestID,CAD_Contas!$A{DAT + i}') + ")", K.FMT_GS),
        (3, f'=IF(CAD_Contas!$A{DAT + i}="",0,' + S("L_Valor", "REALIZADO", f',L_OrigTipo,"CONTA",L_OrigID,CAD_Contas!$A{DAT + i}') + ")", K.FMT_GS),
        (4, f"=$B{rr}-$C{rr}", K.FMT_GS),
        (5, f'=IF(CAD_Contas!$A{DAT + i}="",0,CAD_Contas!$F{DAT + i}'
            f'+SUMIFS(L_Valor,L_DestTipo,"CONTA",L_DestID,CAD_Contas!$A{DAT + i},L_Data,"<="&REL_FIM,L_Status,"REALIZADO")'
            f'-SUMIFS(L_Valor,L_OrigTipo,"CONTA",L_OrigID,CAD_Contas!$A{DAT + i},L_Data,"<="&REL_FIM,L_Status,"REALIZADO")'
            f'+SUMIFS(L_Valor,L_DestTipo,"CONTA",L_DestID,CAD_Contas!$A{DAT + i},L_Data,"<="&REL_FIM,L_Status,"PROJETADO")'
            f'-SUMIFS(L_Valor,L_OrigTipo,"CONTA",L_OrigID,CAD_Contas!$A{DAT + i},L_Data,"<="&REL_FIM,L_Status,"PROJETADO"))', K.FMT_GS),
    ])
    r += 1

    r = secao(ws, r, "5. MOVIMENTO POR CARTAO", 7)
    hdr = [("Cartao", 34), ("Compras_no_Mes", 22), ("Pagamentos_no_Mes", 22),
           ("Divida_no_Fim_do_Mes", 22), ("Limite_Disponivel_Atual", 22)]
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Cartoes!$A{DAT + i}="","",CAD_Cartoes!$B{DAT + i})', None),
        (2, f'=IF(CAD_Cartoes!$A{DAT + i}="",0,' + S("L_Valor", "REALIZADO", f',L_OrigTipo,"CARTAO",L_OrigID,CAD_Cartoes!$A{DAT + i}') + ")", K.FMT_GS),
        (3, f'=IF(CAD_Cartoes!$A{DAT + i}="",0,' + S("L_Valor", "REALIZADO", f',L_DestTipo,"CARTAO",L_DestID,CAD_Cartoes!$A{DAT + i}') + ")", K.FMT_GS),
        (4, f'=IF(CAD_Cartoes!$A{DAT + i}="",0,CAD_Cartoes!$F{DAT + i}'
            f'+SUMIFS(L_Valor,L_OrigTipo,"CARTAO",L_OrigID,CAD_Cartoes!$A{DAT + i},L_Data,"<="&REL_FIM,L_Status,"REALIZADO")'
            f'-SUMIFS(L_Valor,L_DestTipo,"CARTAO",L_DestID,CAD_Cartoes!$A{DAT + i},L_Data,"<="&REL_FIM,L_Status,"REALIZADO"))', K.FMT_GS),
        (5, f'=IF(CAD_Cartoes!$A{DAT + i}="",0,CAD_Cartoes!$N{DAT + i})', K.FMT_GS),
    ])
    r += 1

    r = secao(ws, r, "6. PARCELAS COM VENCIMENTO NO MES", 7)
    hdr = [("Status_Parcela", 34), ("Quantidade", 22), ("Valor", 22)]
    sts = K.STATUS_PARCELA
    r = _grade(ws, r, hdr, len(sts) + 1, lambda i, rr: ([
        (1, sts[i], None),
        (2, f'=COUNTIFS(P_VENC,">="&REL_INI,P_VENC,"<="&REL_FIM,P_STATUS,"{sts[i]}")', "0"),
        (3, f'=SUMIFS(P_VALOR,P_VENC,">="&REL_INI,P_VENC,"<="&REL_FIM,P_STATUS,"{sts[i]}")', K.FMT_GS),
    ] if i < len(sts) else [
        (1, "TOTAL", None),
        (2, f"=SUM($B{rr - len(sts)}:$B{rr - 1})", "0"),
        (3, f"=SUM($C{rr - len(sts)}:$C{rr - 1})", K.FMT_GS),
    ]))
    r += 1

    r = secao(ws, r, "7. INVESTIMENTOS NO MES", 7)
    hdr = [("Investimento", 34), ("Aportes", 22), ("Resgates", 22), ("Rendimentos", 22),
           ("Saldo_Atual", 22)]
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Investimentos!$A{DAT + i}="","",CAD_Investimentos!$B{DAT + i})', None),
        (2, f'=IF(CAD_Investimentos!$A{DAT + i}="",0,' + S("L_Valor", "REALIZADO", f',L_Tipo,"APORTE",L_DestID,CAD_Investimentos!$A{DAT + i}') + ")", K.FMT_GS),
        (3, f'=IF(CAD_Investimentos!$A{DAT + i}="",0,' + S("L_Valor", "REALIZADO", f',L_Tipo,"RESGATE",L_OrigID,CAD_Investimentos!$A{DAT + i}') + ")", K.FMT_GS),
        (4, f'=IF(CAD_Investimentos!$A{DAT + i}="",0,' + S("L_Valor", "REALIZADO", f',L_Tipo,"RENDIMENTO",L_DestID,CAD_Investimentos!$A{DAT + i}') + ")", K.FMT_GS),
        (5, f'=IF(CAD_Investimentos!$A{DAT + i}="",0,CAD_Investimentos!$O{DAT + i})', K.FMT_GS),
    ])
    ws.freeze_panes = "A6"
    return ws


# ---------------------------------------------------------------------------
# RELATORIO ANUAL
# ---------------------------------------------------------------------------
def aba_rel_anual(L):
    ws = L.aba("REL_Anual", "7030A0")
    for c, w in [("A", 16), ("B", 13), ("C", 13)] + [(CL(i), 18) for i in range(4, 15)]:
        ws.column_dimensions[c].width = w
    r = titulo(ws, 1, "REL_ANUAL - RELATORIO DOS 12 MESES", 14,
               "Valores somam REALIZADO + PROJETADO. Escolha o ano na celula amarela.")
    r += 1
    rotulo(ws, r, 1, "Ano")
    valor(ws, r, 2, K.ANO_BASE, "0", entrada=True)
    L.nome("RELA_ANO", f"REL_Anual!$B${r}")
    r += 2
    r = secao(ws, r, "1. RESULTADO E FLUXO MES A MES", 14)
    hdr = [("Mes", 16), ("Data_Inicio", 13), ("Data_Fim", 13), ("Receitas", 18),
           ("Rendimentos", 18), ("Despesas", 18), ("Resultado", 18), ("Aportes", 18),
           ("Resgates", 18), ("Investimento_Liquido", 18), ("Entradas_Caixa", 18),
           ("Saidas_Caixa", 18), ("Fluxo_Liquido", 18), ("Saldo_Final_Contas", 18)]

    def amb(campo, ini, fim, extra=""):
        return (f'SUMIFS({campo},L_Data,">="&{ini},L_Data,"<="&{fim},L_Status,"REALIZADO"{extra})'
                f'+SUMIFS({campo},L_Data,">="&{ini},L_Data,"<="&{fim},L_Status,"PROJETADO"{extra})')

    def gen(i, rr):
        ini, fim = f"$B{rr}", f"$C{rr}"
        return [
            (1, f"=INDEX(LST_MESES,{i + 1})", None),
            (2, f"=DATE(RELA_ANO,{i + 1},1)", K.FMT_DATA),
            (3, f"=EOMONTH($B{rr},0)", K.FMT_DATA),
            (4, "=" + amb("L_VRec", ini, fim), K.FMT_GS),
            (5, "=" + amb("L_VRend", ini, fim), K.FMT_GS),
            (6, "=" + amb("L_VDesp", ini, fim), K.FMT_GS),
            (7, f"=$D{rr}+$E{rr}-$F{rr}", K.FMT_GS),
            (8, "=" + amb("L_Valor", ini, fim, ',L_Tipo,"APORTE"'), K.FMT_GS),
            (9, "=" + amb("L_Valor", ini, fim, ',L_Tipo,"RESGATE"'), K.FMT_GS),
            (10, f"=$H{rr}-$I{rr}", K.FMT_GS),
            (11, "=" + amb("L_Ent", ini, fim), K.FMT_GS),
            (12, "=" + amb("L_Sai", ini, fim), K.FMT_GS),
            (13, f"=$K{rr}-$L{rr}", K.FMT_GS),
            (14, "=" + SALDO_BASE.format(d=f"$C{rr}+1"), K.FMT_GS),
        ]

    r0 = r + 1
    r = _grade(ws, r, hdr, 12, gen)
    L.nome("RA_MES", f"REL_Anual!$A${r0}:$A${r0 + 11}")
    L.nome("RA_DESP", f"REL_Anual!$F${r0}:$F${r0 + 11}")
    L.nome("RA_REC", f"REL_Anual!$D${r0}:$D${r0 + 11}")
    L.nome("RA_RESULT", f"REL_Anual!$G${r0}:$G${r0 + 11}")
    L.nome("RA_SALDO", f"REL_Anual!$N${r0}:$N${r0 + 11}")
    for lab, fn in [("TOTAL DO ANO", "SUM"), ("MEDIA MENSAL", "AVERAGE")]:
        ws.cell(r, 1, lab).font = K.F_BOLD
        ws.cell(r, 1).fill = K.FILL_SUB
        ws.cell(r, 1).border = K.BORDA
        for col in range(4, 14):
            c = ws.cell(r, col, f"={fn}({CL(col)}{r0}:{CL(col)}{r0 + 11})")
            c.number_format = K.FMT_GS
            c.font = K.F_BOLD
            c.fill = K.FILL_SUB
            c.border = K.BORDA
        r += 1
    r += 1

    r = secao(ws, r, "2. DESPESAS DO ANO POR CATEGORIA", 14)
    hdr = [("Categoria", 30), ("Despesa_no_Ano", 20), ("% do total", 14)]
    tot = 'SUM(RA_DESP)'
    rcat = r + 1
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Categorias!$A{DAT + i}="","",CAD_Categorias!$B{DAT + i})', None),
        (2, f'=IF(CAD_Categorias!$A{DAT + i}="",0,'
            f'SUMIFS(L_VDesp,L_Ano,RELA_ANO,L_CatID,CAD_Categorias!$A{DAT + i},L_Status,"REALIZADO")'
            f'+SUMIFS(L_VDesp,L_Ano,RELA_ANO,L_CatID,CAD_Categorias!$A{DAT + i},L_Status,"PROJETADO"))', K.FMT_GS),
        (3, f"=IF({tot}=0,0,$B{rr}/{tot})", K.FMT_PCT),
    ])
    L.nome("RA_CAT_NOME", f"REL_Anual!$A${rcat}:$A${rcat + K.MAX_CAD - 1}")
    L.nome("RA_CAT_TOT", f"REL_Anual!$B${rcat}:$B${rcat + K.MAX_CAD - 1}")
    r += 1

    r = secao(ws, r, "3. INDICADORES DO ANO", 14)
    r = kv(ws, r, "Total de receitas", "=SUM(RA_REC)", K.FMT_GS, bold=True, largura=46)
    r = kv(ws, r, "Total de despesas", "=SUM(RA_DESP)", K.FMT_GS, bold=True, largura=46)
    r = kv(ws, r, "Resultado do ano", "=SUM(RA_RESULT)", K.FMT_GS, bold=True, largura=46)
    r = kv(ws, r, "Despesa media mensal", "=AVERAGE(RA_DESP)", K.FMT_GS, largura=46)
    r = kv(ws, r, "Maior despesa mensal", "=MAX(RA_DESP)", K.FMT_GS, largura=46)
    r = kv(ws, r, "Mes da maior despesa", "=INDEX(RA_MES,MATCH(MAX(RA_DESP),RA_DESP,0))",
           "General", largura=46)
    r = kv(ws, r, "Menor despesa mensal", "=MIN(RA_DESP)", K.FMT_GS, largura=46)
    r = kv(ws, r, "Mes da menor despesa", "=INDEX(RA_MES,MATCH(MIN(RA_DESP),RA_DESP,0))",
           "General", largura=46)
    r = kv(ws, r, "Categoria de maior gasto no ano",
           "=IF(MAX(RA_CAT_TOT)=0,\"-\",INDEX(RA_CAT_NOME,MATCH(MAX(RA_CAT_TOT),RA_CAT_TOT,0)))",
           "General", largura=46)
    r = kv(ws, r, "Valor da categoria de maior gasto", "=MAX(RA_CAT_TOT)", K.FMT_GS, largura=46)
    r = kv(ws, r, "Taxa de poupanca do ano (resultado / receitas)",
           "=IF(SUM(RA_REC)=0,0,SUM(RA_RESULT)/SUM(RA_REC))", K.FMT_PCT, largura=46)
    r = kv(ws, r, "Saldo das contas em 31/12",
           "=" + SALDO_BASE.format(d="DATE(RELA_ANO,12,31)+1"), K.FMT_GS, largura=46)
    ev = ('SUM(I_VALINI)+SUMIFS(L_EfInv,L_Data,"<="&{d},L_Status,"REALIZADO")'
          '+SUMIFS(L_EfInv,L_Data,"<="&{d},L_Status,"PROJETADO")')
    dv_ = ('SUM(K_DIVINI)+SUMIFS(L_EfCartao,L_Data,"<="&{d},L_Status,"REALIZADO")'
           '+SUMIFS(L_EfCartao,L_Data,"<="&{d},L_Status,"PROJETADO")')
    r = kv(ws, r, "Investimentos em 31/12", "=" + ev.format(d="DATE(RELA_ANO,12,31)"),
           K.FMT_GS, largura=46)
    r = kv(ws, r, "Evolucao dos investimentos no ano",
           "=" + ev.format(d="DATE(RELA_ANO,12,31)") + "-(" + ev.format(d="DATE(RELA_ANO,1,1)-1") + ")",
           K.FMT_GS, largura=46)
    r = kv(ws, r, "Dividas de cartao em 31/12", "=" + dv_.format(d="DATE(RELA_ANO,12,31)"),
           K.FMT_GS, largura=46)
    r = kv(ws, r, "Evolucao das dividas de cartao no ano",
           "=" + dv_.format(d="DATE(RELA_ANO,12,31)") + "-(" + dv_.format(d="DATE(RELA_ANO,1,1)-1") + ")",
           K.FMT_GS, largura=46)
    r = kv(ws, r, "Saldo devedor de compromissos em 31/12",
           '=SUMIFS(P_VALOR,P_VENC,">"&DATE(RELA_ANO,12,31))', K.FMT_GS, largura=46)
    r = kv(ws, r, "Patrimonio liquido total em 31/12 (evolucao)",
           '=IFERROR(INDEX(PT_PLTOTAL,MATCH(RELA_ANO*100+12,PT_COMP,0)),"fora da grade")',
           K.FMT_GS, bold=True, largura=46)
    ws.freeze_panes = "A6"
    return ws


# ---------------------------------------------------------------------------
# INDICADORES
# ---------------------------------------------------------------------------
REC_MED = ('SUMIFS(L_VRec,L_Data,">="&EDATE(CFG_HOJE,-6),L_Data,"<="&CFG_HOJE,L_Status,"REALIZADO")/6')
DESP_MED = ('SUMIFS(L_VDesp,L_Data,">="&EDATE(CFG_HOJE,-6),L_Data,"<"&DATE(YEAR(CFG_HOJE),MONTH(CFG_HOJE),1),L_Status,"REALIZADO")/6')
DESP_MES = ('SUMIFS(L_VDesp,L_Data,">="&DATE(YEAR(CFG_HOJE),MONTH(CFG_HOJE),1),L_Data,"<="&EOMONTH(CFG_HOJE,0),L_Status,"REALIZADO")')
COMPROM_12 = ('SUMIFS(L_Valor,L_Data,">"&CFG_HOJE,L_Data,"<="&EDATE(CFG_HOJE,12),L_Status,"PROJETADO",L_ParcID,"<>")'
              '+SUMIFS(L_Valor,L_Data,">"&CFG_HOJE,L_Data,"<="&EDATE(CFG_HOJE,12),L_Status,"PROJETADO",L_Tipo,"PAGAMENTO_CARTAO")')


def aba_indicadores(L):
    ws = L.aba("INDICADORES", "BF8F00")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 78
    r = titulo(ws, 1, "INDICADORES - INTELIGENCIA FINANCEIRA POR REGRAS", 3,
               "Calculado apenas com formulas, sem inteligencia artificial. "
               "Base: data de referencia e ano/mes definidos em CFG_Sistema.")
    r += 1
    r = secao(ws, r, "1. ONDE O DINHEIRO ESTA INDO", 3)
    r = kv(ws, r, "Categoria de maior gasto no mes selecionado (REL_Mensal)",
           '=IF(MAX(RM_CAT_REAL)=0,"-",INDEX(RM_CAT_NOME,MATCH(MAX(RM_CAT_REAL),RM_CAT_REAL,0)))',
           "General", bold=True, largura=52,
           obs="Depende do ano/mes escolhidos na aba REL_Mensal.")
    r = kv(ws, r, "Valor gasto nessa categoria", "=MAX(RM_CAT_REAL)", K.FMT_GS, largura=52)
    r = kv(ws, r, "Categoria de maior gasto no ano (REL_Anual)",
           '=IF(MAX(RA_CAT_TOT)=0,"-",INDEX(RA_CAT_NOME,MATCH(MAX(RA_CAT_TOT),RA_CAT_TOT,0)))',
           "General", bold=True, largura=52)
    r = kv(ws, r, "Valor gasto nessa categoria no ano", "=MAX(RA_CAT_TOT)", K.FMT_GS, largura=52)
    r += 1
    r = secao(ws, r, "2. RITMO DE RECEITAS E DESPESAS", 3)
    r = kv(ws, r, "Receita media mensal realizada (ultimos 6 meses)", "=" + REC_MED,
           K.FMT_GS, largura=52)
    r = kv(ws, r, "Despesa media mensal realizada (6 meses anteriores)", "=" + DESP_MED,
           K.FMT_GS, largura=52)
    r = kv(ws, r, "Despesa realizada no mes corrente", "=" + DESP_MES, K.FMT_GS, largura=52)
    r = kv(ws, r, "Variacao da despesa do mes sobre a media",
           f"=IF({DESP_MED}=0,0,{DESP_MES}/({DESP_MED})-1)", K.FMT_PCT, largura=52,
           obs="Compara apenas dias ja decorridos do mes corrente com a media dos 6 meses anteriores.")
    r = kv(ws, r, "Taxa de poupanca dos ultimos 6 meses",
           f'=IF({REC_MED}=0,0,1-(SUMIFS(L_VDesp,L_Data,">="&EDATE(CFG_HOJE,-6),L_Data,"<="&CFG_HOJE,L_Status,"REALIZADO")/6)/({REC_MED}))',
           K.FMT_PCT, bold=True, largura=52)
    r += 1
    r = secao(ws, r, "3. COMPROMETIMENTO FUTURO", 3)
    r = kv(ws, r, "Parcelas + faturas dos proximos 12 meses", "=" + COMPROM_12,
           K.FMT_GS, bold=True, largura=52)
    r = kv(ws, r, "Percentual da renda comprometida (12 meses)",
           f"=IF({REC_MED}=0,0,({COMPROM_12})/(({REC_MED})*12))", K.FMT_PCT, bold=True,
           largura=52, obs="Comparado com o limite definido em CFG_Sistema.")
    r = kv(ws, r, "Total de parcelas ainda nao pagas (todos os compromissos)",
           '=SUMIFS(P_VALOR,P_STATUS,"PROJETADA")+SUMIFS(P_VALOR,P_STATUS,"ABERTA")+SUMIFS(P_VALOR,P_STATUS,"ATRASADA")',
           K.FMT_GS, largura=52)
    r = kv(ws, r, "Maior saida de caixa projetada em um mes", "=MAX(PJ_SAIDAS)", K.FMT_GS, largura=52)
    r = kv(ws, r, "Competencia do mes de maior saida projetada",
           '=IF(MAX(PJ_SAIDAS)=0,"-",INDEX(PJ_COMP,MATCH(MAX(PJ_SAIDAS),PJ_SAIDAS,0)))',
           "0", largura=52)
    r = kv(ws, r, "Menor saldo projetado no horizonte", "=MIN(PJ_SALDO)", K.FMT_GS, bold=True, largura=52)
    r = kv(ws, r, "Competencia do menor saldo projetado",
           "=INDEX(PJ_COMP,MATCH(MIN(PJ_SALDO),PJ_SALDO,0))", "0", largura=52)
    r += 1
    r = secao(ws, r, "4. PATRIMONIO", 3)
    r = kv(ws, r, "Patrimonio liquido financeiro atual", "=PAT_PL_FIN", K.FMT_GS, bold=True, largura=52)
    r = kv(ws, r, "Patrimonio liquido total atual", "=PAT_PL_TOTAL", K.FMT_GS, bold=True, largura=52)
    r = kv(ws, r, "Patrimonio total 6 meses atras",
           '=IFERROR(INDEX(PT_PLTOTAL,MATCH(YEAR(EDATE(CFG_HOJE,-6))*100+MONTH(EDATE(CFG_HOJE,-6)),PT_COMP,0)),0)',
           K.FMT_GS, largura=52)
    r = kv(ws, r, "Variacao do patrimonio em 6 meses",
           '=IFERROR(INDEX(PT_PLTOTAL,MATCH(YEAR(CFG_HOJE)*100+MONTH(CFG_HOJE),PT_COMP,0)),0)'
           '-IFERROR(INDEX(PT_PLTOTAL,MATCH(YEAR(EDATE(CFG_HOJE,-6))*100+MONTH(EDATE(CFG_HOJE,-6)),PT_COMP,0)),0)',
           K.FMT_GS, largura=52)
    r = kv(ws, r, "Tendencia do patrimonio",
           f'=IF($B{r - 1}>0,"CRESCENTE",IF($B{r - 1}<0,"DECRESCENTE","ESTAVEL"))',
           "General", bold=True, largura=52)
    r += 1
    r = secao(ws, r, "5. INVESTIMENTOS E METAS", 3)
    r = kv(ws, r, "Total investido (saldo atual)", "=SUM(I_SALDO)", K.FMT_GS, bold=True, largura=52)
    r = kv(ws, r, "Rendimento ja realizado", "=SUM(I_REND)", K.FMT_GS, largura=52)
    r = kv(ws, r, "Rendimento ainda projetado", "=SUM(I_RENDPROJ)", K.FMT_GS, largura=52,
           obs="Projecao. Nunca somada ao rendimento realizado nos relatorios de resultado.")
    r = kv(ws, r, "Investimentos vencendo dentro do prazo de alerta",
           '=COUNTIFS(I_VENC,">="&CFG_HOJE,I_VENC,"<="&CFG_HOJE+CFG_DIAS_INV)', "0", largura=52)
    r = kv(ws, r, "Metas ativas", '=COUNTIF(M_STATUS,"Ativa")', "0", largura=52)
    r = kv(ws, r, "Metas com 80% ou mais concluido",
           '=SUMPRODUCT((M_STATUS="Ativa")*(M_PCT>=0.8))', "0", largura=52)
    r = kv(ws, r, "Total reservado em metas", "=SUM(M_RESERV)", K.FMT_GS, largura=52)
    r = kv(ws, r, "Total ainda faltante nas metas ativas",
           '=SUMIFS(M_FALTA,M_STATUS,"Ativa")', K.FMT_GS, largura=52)
    return ws


# ---------------------------------------------------------------------------
# ALERTAS
# ---------------------------------------------------------------------------
def aba_alertas(L):
    ws = L.aba("ALERTAS", "C00000")
    for c, w in [("A", 38), ("B", 20), ("C", 20), ("D", 16), ("E", 18), ("F", 14), ("G", 16)]:
        ws.column_dimensions[c].width = w
    r = titulo(ws, 1, "ALERTAS - SITUACOES QUE EXIGEM ATENCAO", 7,
               "Os limites de cada alerta ficam na aba CFG_Sistema. ALERTA = fora do limite; "
               "OK = dentro do limite.")
    r += 1
    r_total = r
    rotulo(ws, r, 1, "TOTAL DE ALERTAS ATIVOS")
    r += 2
    faixas = []

    r = secao(ws, r, "1. SALDO DAS CONTAS", 7)
    hdr = [("Conta", 38), ("Saldo_Disponivel", 20), ("Saldo_Minimo", 20), ("Situacao", 16)]
    ini = r + 1
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Contas!$A{DAT + i}="","",CAD_Contas!$B{DAT + i})', None),
        (2, f'=IF(CAD_Contas!$A{DAT + i}="","",CAD_Contas!$L{DAT + i})', K.FMT_GS),
        (3, f'=IF(CAD_Contas!$A{DAT + i}="","",CFG_SALDO_MIN)', K.FMT_GS),
        (4, f'=IF($A{rr}="","",IF($B{rr}<$C{rr},"ALERTA","OK"))', None),
    ])
    faixas.append(f"$D${ini}:$D${r - 1}")
    ws.conditional_formatting.add(f"D{ini}:D{r - 1}", CellIsRule(
        operator="equal", formula=['"ALERTA"'], fill=K.FILL_ERRO))
    r += 1

    r = secao(ws, r, "2. CARTOES DE CREDITO", 7)
    hdr = [("Cartao", 38), ("Divida_Atual", 20), ("Limite", 20), ("Uso_do_Limite", 16),
           ("Vencimento_Fatura", 18), ("Dias", 14), ("Situacao", 16)]
    ini = r + 1
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Cartoes!$A{DAT + i}="","",CAD_Cartoes!$B{DAT + i})', None),
        (2, f'=IF(CAD_Cartoes!$A{DAT + i}="","",CAD_Cartoes!$L{DAT + i})', K.FMT_GS),
        (3, f'=IF(CAD_Cartoes!$A{DAT + i}="","",CAD_Cartoes!$E{DAT + i})', K.FMT_GS),
        (4, f'=IF(CAD_Cartoes!$A{DAT + i}="","",CAD_Cartoes!$M{DAT + i})', K.FMT_PCT),
        (5, f'=IF(CAD_Cartoes!$A{DAT + i}="","",CAD_Cartoes!$Q{DAT + i})', K.FMT_DATA),
        (6, f'=IF($A{rr}="","",$E{rr}-CFG_HOJE)', "0"),
        (7, f'=IF($A{rr}="","",IF($D{rr}>=CFG_PCT_CARTAO,"ALERTA",'
            f'IF(AND($F{rr}>=0,$F{rr}<=CFG_DIAS_ALERTA),"ALERTA","OK")))', None),
    ])
    faixas.append(f"$G${ini}:$G${r - 1}")
    ws.conditional_formatting.add(f"G{ini}:G{r - 1}", CellIsRule(
        operator="equal", formula=['"ALERTA"'], fill=K.FILL_ERRO))
    r += 1

    r = secao(ws, r, "3. PARCELAS E COMPROMISSOS", 7)
    hdr = [("Situacao verificada", 38), ("Quantidade", 20), ("Valor", 20), ("Situacao", 16)]
    regras = [
        ("Parcelas ATRASADAS", '=COUNTIF(P_STATUS,"ATRASADA")', '=SUMIFS(P_VALOR,P_STATUS,"ATRASADA")'),
        ("Parcelas vencendo dentro do prazo de alerta",
         '=COUNTIFS(P_VENC,">="&CFG_HOJE,P_VENC,"<="&CFG_HOJE+CFG_DIAS_ALERTA,P_STATUS,"ABERTA")',
         '=SUMIFS(P_VALOR,P_VENC,">="&CFG_HOJE,P_VENC,"<="&CFG_HOJE+CFG_DIAS_ALERTA,P_STATUS,"ABERTA")'),
        ("Parcelas ABERTAS no mes corrente", '=COUNTIF(P_STATUS,"ABERTA")',
         '=SUMIFS(P_VALOR,P_STATUS,"ABERTA")'),
    ]
    ini = r + 1
    r = _grade(ws, r, hdr, len(regras), lambda i, rr: [
        (1, regras[i][0], None), (2, regras[i][1], "0"), (3, regras[i][2], K.FMT_GS),
        (4, f'=IF($B{rr}>0,"ALERTA","OK")', None),
    ])
    faixas.append(f"$D${ini}:$D${r - 1}")
    ws.conditional_formatting.add(f"D{ini}:D{r - 1}", CellIsRule(
        operator="equal", formula=['"ALERTA"'], fill=K.FILL_ERRO))
    r += 1

    r = secao(ws, r, "4. INVESTIMENTOS", 7)
    hdr = [("Investimento", 38), ("Saldo_Atual", 20), ("Data_Vencimento", 20),
           ("Dias_para_Vencer", 16), ("Situacao", 16)]
    ini = r + 1
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Investimentos!$A{DAT + i}="","",CAD_Investimentos!$B{DAT + i})', None),
        (2, f'=IF(CAD_Investimentos!$A{DAT + i}="","",CAD_Investimentos!$O{DAT + i})', K.FMT_GS),
        (3, f'=IF(CAD_Investimentos!$A{DAT + i}="","",CAD_Investimentos!$K{DAT + i})', K.FMT_DATA),
        (4, f'=IF(OR($A{rr}="",$C{rr}=""),"",$C{rr}-CFG_HOJE)', "0"),
        (5, f'=IF(OR($A{rr}="",$C{rr}=""),"",IF(AND($D{rr}>=0,$D{rr}<=CFG_DIAS_INV),"ALERTA","OK"))', None),
    ])
    faixas.append(f"$E${ini}:$E${r - 1}")
    ws.conditional_formatting.add(f"E{ini}:E{r - 1}", CellIsRule(
        operator="equal", formula=['"ALERTA"'], fill=K.FILL_ERRO))
    r += 1

    r = secao(ws, r, "5. METAS", 7)
    hdr = [("Meta", 38), ("Pct_Concluido", 20), ("Prazo", 20), ("Situacao_da_Meta", 20),
           ("Situacao", 16)]
    ini = r + 1
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$B{DAT + i})', None),
        (2, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$I{DAT + i})', K.FMT_PCT),
        (3, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$D{DAT + i})', K.FMT_DATA),
        (4, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$O{DAT + i})', None),
        (5, f'=IF($A{rr}="","",IF(OR($D{rr}="PRAZO VENCIDO",$D{rr}="EM RISCO"),"ALERTA","OK"))', None),
    ])
    faixas.append(f"$E${ini}:$E${r - 1}")
    ws.conditional_formatting.add(f"E{ini}:E{r - 1}", CellIsRule(
        operator="equal", formula=['"ALERTA"'], fill=K.FILL_ERRO))
    r += 1

    r = secao(ws, r, "6. FLUXO, COMPROMETIMENTO E INTEGRIDADE", 7)
    hdr = [("Verificacao", 38), ("Valor_Apurado", 20), ("Limite", 20), ("Situacao", 16)]
    regras = [
        ("Menor saldo projetado no horizonte", "=MIN(PJ_SALDO)", "=0",
         f'=IF($B{{rr}}<$C{{rr}},"ALERTA","OK")'),
        ("Menor saldo mensal do fluxo de caixa", "=MIN(FX_SALDOFIM)", "=0",
         f'=IF($B{{rr}}<$C{{rr}},"ALERTA","OK")'),
        ("Percentual da renda comprometida (12 meses)",
         f"=IF({REC_MED}=0,0,({COMPROM_12})/(({REC_MED})*12))", "=CFG_PCT_RENDA",
         f'=IF($B{{rr}}>$C{{rr}},"ALERTA","OK")'),
        ("Aumento da despesa do mes sobre a media",
         f"=IF({DESP_MED}=0,0,{DESP_MES}/({DESP_MED})-1)", "=CFG_PCT_DESP",
         f'=IF($B{{rr}}>$C{{rr}},"ALERTA","OK")'),
        ("Lancamentos com erro de validacao",
         '=SUMPRODUCT((LEFT(L_Valid,4)="ERRO")*1)', "=0", f'=IF($B{{rr}}>$C{{rr}},"ALERTA","OK")'),
        ("Movimentos de meta com erro de validacao",
         '=SUMPRODUCT((LEFT(MM_VALID,4)="ERRO")*1)', "=0", f'=IF($B{{rr}}>$C{{rr}},"ALERTA","OK")'),
        ("Parcelas com mais de um pagamento realizado",
         '=SUMPRODUCT((P_QTDREAL>1)*1)', "=0", f'=IF($B{{rr}}>$C{{rr}},"ALERTA","OK")'),
        ("Parcelas ja pagas que ainda tem lancamento projetado",
         '=SUMPRODUCT((P_QTDREAL>0)*(P_QTDPROJ>0)*1)', "=0",
         f'=IF($B{{rr}}>$C{{rr}},"ALERTA","OK")'),
    ]
    fmts = [K.FMT_GS, K.FMT_GS, K.FMT_PCT, K.FMT_PCT, "0", "0", "0", "0"]
    ini = r + 1
    r = _grade(ws, r, hdr, len(regras), lambda i, rr: [
        (1, regras[i][0], None), (2, regras[i][1], fmts[i]), (3, regras[i][2], fmts[i]),
        (4, regras[i][3].replace("{rr}", str(rr)), None),
    ])
    faixas.append(f"$D${ini}:$D${r - 1}")
    ws.conditional_formatting.add(f"D{ini}:D{r - 1}", CellIsRule(
        operator="equal", formula=['"ALERTA"'], fill=K.FILL_ERRO))

    total = "=" + "+".join(f'COUNTIF({f},"ALERTA")' for f in faixas)
    c = ws.cell(r_total, 2, total)
    c.number_format = "0"
    c.font = K.F_KPI
    c.fill = K.FILL_ALERTA
    c.border = K.BORDA
    L.nome("AL_TOTAL", f"ALERTAS!$B${r_total}")
    ws.freeze_panes = "A6"
    return ws


# ---------------------------------------------------------------------------
# DASHBOARD
# ---------------------------------------------------------------------------
def _kpi(ws, r, col, label, formula, fmt=None, cor=None):
    ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
    c = ws.cell(r, col, label)
    c.font = Font(name=K.FONTE, size=8, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=cor or K.C_TITULO)
    c.alignment = K.AL_C
    ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 1)
    v = ws.cell(r + 1, col, formula)
    v.font = K.F_KPI
    v.number_format = fmt or K.FMT_GS
    v.alignment = K.AL_C
    v.fill = PatternFill("solid", fgColor="EDF2F9")
    v.border = K.BORDA
    ws.row_dimensions[r + 1].height = 24


def aba_dashboard(L):
    ws = L.aba("DASHBOARD", "1F3864", zoom=85)
    for i in range(1, 9):
        ws.column_dimensions[CL(i)].width = 21
    r = titulo(ws, 1, "PAINEL FINANCEIRO PESSOAL - GUARANI (Gs.)", 8,
               "Data de referencia e parametros na aba CFG_Sistema. Valores REALIZADOS salvo "
               "quando o titulo do indicador disser projetado.")
    r += 1
    blocos = [
        ("DINHEIRO", [
            ("SALDO TOTAL DAS CONTAS", "=SUM(C_SALDO)", None),
            ("RESERVADO PARA METAS", "=SUM(C_RESERV)", None),
            ("SALDO DISPONIVEL (LIVRE)", "=SUM(C_DISP)", None),
            ("SALDO PROJETADO (FIM DO HORIZONTE)", "=INDEX(PJ_SALDO,COUNT(PJ_SALDO))", None),
        ]),
        ("INVESTIMENTOS E PATRIMONIO", [
            ("TOTAL INVESTIDO", "=SUM(I_SALDO)", None),
            ("RENDIMENTO REALIZADO", "=SUM(I_REND)", None),
            ("RENDIMENTO PROJETADO", "=SUM(I_RENDPROJ)", None),
            ("PATRIMONIO LIQUIDO TOTAL", "=PAT_PL_TOTAL", None),
        ]),
        ("DIVIDAS", [
            ("DIVIDA DE CARTOES", "=SUM(K_DIVIDA)", None),
            ("LIMITE DISPONIVEL EM CARTOES", "=SUM(K_DISP)", None),
            ("SALDO DEVEDOR DE COMPROMISSOS", "=SUM(CM_DEVEDOR)", None),
            ("PARCELAS + FATURAS PROXIMOS 12 MESES", "=" + COMPROM_12, None),
        ]),
        ("FLUXO DO MES CORRENTE", [
            ("RECEITAS REALIZADAS NO MES",
             '=SUMIFS(L_VRec,L_Data,">="&DATE(YEAR(CFG_HOJE),MONTH(CFG_HOJE),1),L_Data,"<="&EOMONTH(CFG_HOJE,0),L_Status,"REALIZADO")', None),
            ("DESPESAS REALIZADAS NO MES", "=" + DESP_MES, None),
            ("RESULTADO DO MES",
             '=SUMIFS(L_VRec,L_Data,">="&DATE(YEAR(CFG_HOJE),MONTH(CFG_HOJE),1),L_Data,"<="&EOMONTH(CFG_HOJE,0),L_Status,"REALIZADO")'
             '+SUMIFS(L_VRend,L_Data,">="&DATE(YEAR(CFG_HOJE),MONTH(CFG_HOJE),1),L_Data,"<="&EOMONTH(CFG_HOJE,0),L_Status,"REALIZADO")'
             "-(" + DESP_MES + ")", None),
            ("MENOR SALDO PROJETADO NO HORIZONTE", "=MIN(PJ_SALDO)", None),
        ]),
        ("METAS", [
            ("METAS ATIVAS", '=COUNTIF(M_STATUS,"Ativa")', "0"),
            ("VALOR ACUMULADO NAS METAS", '=SUMIFS(M_RESERV,M_STATUS,"Ativa")', None),
            ("VALOR RESTANTE NAS METAS", '=SUMIFS(M_FALTA,M_STATUS,"Ativa")', None),
            ("CONCLUSAO MEDIA DAS METAS ATIVAS",
             '=IF(COUNTIF(M_STATUS,"Ativa")=0,0,SUMIFS(M_RESERV,M_STATUS,"Ativa")/SUMIFS(M_OBJ,M_STATUS,"Ativa"))', K.FMT_PCT),
        ]),
        ("CONTROLE", [
            ("ALERTAS ATIVOS", "=AL_TOTAL", "0"),
            ("LANCAMENTOS REGISTRADOS", '=COUNTIF(L_Status,"REALIZADO")+COUNTIF(L_Status,"PROJETADO")', "0"),
            ("LANCAMENTOS COM ERRO", '=SUMPRODUCT((LEFT(L_Valid,4)="ERRO")*1)', "0"),
            ("TESTES COM FALHA", "=TS_FALHAS", "0"),
        ]),
    ]
    for nome, itens in blocos:
        r = secao(ws, r, nome, 8)
        for j, (lab, f, fmt) in enumerate(itens):
            _kpi(ws, r, 1 + j * 2, lab, f, fmt)
        r += 3
    r = secao(ws, r, "SALDO DAS CONTAS", 8)
    hdr = [("Conta", 26), ("Saldo_Calculado", 20), ("Reservado", 18), ("Disponivel", 20),
           ("Saldo_Projetado", 20)]
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Contas!$A{DAT + i}="","",CAD_Contas!$B{DAT + i})', None),
        (2, f'=IF(CAD_Contas!$A{DAT + i}="","",CAD_Contas!$J{DAT + i})', K.FMT_GS),
        (3, f'=IF(CAD_Contas!$A{DAT + i}="","",CAD_Contas!$K{DAT + i})', K.FMT_GS),
        (4, f'=IF(CAD_Contas!$A{DAT + i}="","",CAD_Contas!$L{DAT + i})', K.FMT_GS),
        (5, f'=IF(CAD_Contas!$A{DAT + i}="","",CAD_Contas!$O{DAT + i})', K.FMT_GS),
    ])
    r += 1
    r = secao(ws, r, "METAS FINANCEIRAS", 8)
    hdr = [("Meta", 26), ("Objetivo", 20), ("Acumulado", 18), ("Faltante", 20),
           ("Concluido", 14), ("Situacao", 18)]
    r = _grade(ws, r, hdr, K.MAX_CAD, lambda i, rr: [
        (1, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$B{DAT + i})', None),
        (2, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$C{DAT + i})', K.FMT_GS),
        (3, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$H{DAT + i})', K.FMT_GS),
        (4, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$J{DAT + i})', K.FMT_GS),
        (5, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$I{DAT + i})', K.FMT_PCT),
        (6, f'=IF(CAD_Metas!$A{DAT + i}="","",CAD_Metas!$O{DAT + i})', None),
    ])
    return ws


# ---------------------------------------------------------------------------
# LANCAR - TELA DE LANCAMENTO RAPIDO
# ---------------------------------------------------------------------------
def aba_lancar(L):
    ws = L.aba("LANCAR", "ED7D31")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20
    for i in range(5, 26):
        ws.column_dimensions[CL(i)].width = 17
    r = titulo(ws, 1, "LANCAR - TELA DE LANCAMENTO RAPIDO", 8,
               "Preencha as celulas amarelas. O sistema resolve os IDs e monta a linha pronta. "
               "Copie a LINHA PRONTA e cole na primeira linha vazia de LANCAMENTOS (colunas B ate S). "
               "O ID do lancamento e gerado automaticamente pela planilha.")
    r += 1
    campos = [
        ("Data", K.HOJE, K.FMT_DATA, None),
        ("Descricao", "Exemplo: Supermercado da semana", None, None),
        ("Tipo de operacao", "DESPESA", None, "MAP_TIPO"),
        ("Status", "REALIZADO", None, "LST_STATUS"),
        ("Valor (Gs.)", 250000, K.FMT_GS, None),
        ("Origem - tipo", "CONTA", None, "LST_ENT_TIPO"),
        ("Origem - nome", ("" if MODO_LIMPO else "Cuenta Corriente Ueno"), None, "IND_ORIGEM"),
        ("Destino - tipo", "EXTERNO", None, "LST_ENT_TIPO"),
        ("Destino - nome", "(Externo)", None, "IND_DESTINO"),
        ("Categoria", "Alimentacao", None, "LST_CATEGORIAS"),
        ("Subcategoria", "Supermercado", None, "IND_SUB"),
        ("Forma de pagamento", "Debito", None, "LST_FORMA"),
        ("Compromisso (opcional)", "", None, "LST_COMPROMISSOS"),
        ("Numero da parcela (vazio = proxima em aberto)", "", "0", None),
        ("Meta (opcional)", "", None, "LST_METAS_NOME"),
        ("Recorrencia (opcional)", "", None, "R_ID"),
        ("Lancamento relacionado (estorno, opcional)", "", None, None),
        ("Observacao", "", None, None),
    ]
    r_ini = r
    for i, (lab, val, fmt, dv) in enumerate(campos):
        rr = r + i
        rotulo(ws, rr, 1, lab)
        c = valor(ws, rr, 2, val, fmt, entrada=True)
        c.alignment = K.AL_L
        if dv and not dv.startswith("IND_"):
            dv_lista(ws, dv, f"B{rr}")
    R = {lab: r_ini + i for i, (lab, *_x) in enumerate(campos)}
    rD, rDesc, rTipo, rSt, rVal = R["Data"], R["Descricao"], R["Tipo de operacao"], R["Status"], R["Valor (Gs.)"]
    rOT, rON = R["Origem - tipo"], R["Origem - nome"]
    rDT, rDN = R["Destino - tipo"], R["Destino - nome"]
    rCat, rSub, rForma = R["Categoria"], R["Subcategoria"], R["Forma de pagamento"]
    rCmp, rNum = R["Compromisso (opcional)"], R["Numero da parcela (vazio = proxima em aberto)"]
    rMeta, rRec = R["Meta (opcional)"], R["Recorrencia (opcional)"]
    rRel, rObs = R["Lancamento relacionado (estorno, opcional)"], R["Observacao"]

    # listas dependentes
    for celula, formula in [(f"B{rON}", f'=INDIRECT("LST_NOMES_"&$B${rOT})'),
                            (f"B{rDN}", f'=INDIRECT("LST_NOMES_"&$B${rDT})'),
                            (f"B{rSub}", f'=INDIRECT("SUB_"&SUBSTITUTE($D${rCat},"-","_"))')]:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            showDropDown=False, errorStyle="warning")
        dv.error = "Valor fora da lista. Cadastre-o antes ou escolha um da lista."
        ws.add_data_validation(dv)
        dv.add(celula)

    # resolucao de IDs (coluna D)
    res = [
        (rOT, "ID resolvido", f'=IF($B${rOT}="EXTERNO","",IFERROR(INDEX(ENT_ID,MATCH(1,INDEX((ENT_NOME=$B${rON})*(ENT_TIPO=$B${rOT}),0),0)),"NAO ENCONTRADO"))'),
        (rDT, "ID resolvido", f'=IF($B${rDT}="EXTERNO","",IFERROR(INDEX(ENT_ID,MATCH(1,INDEX((ENT_NOME=$B${rDN})*(ENT_TIPO=$B${rDT}),0),0)),"NAO ENCONTRADO"))'),
        (rCat, "ID resolvido", f'=IF($B${rCat}="","",IFERROR(INDEX(CAT_ID,MATCH($B${rCat},CAT_NOME,0)),"NAO ENCONTRADO"))'),
        (rSub, "ID resolvido", f'=IF($B${rSub}="","",IFERROR(INDEX(SUB_ID,MATCH(1,INDEX((SUB_NOME=$B${rSub})*(SUB_CATID=$D${rCat}),0),0)),"NAO ENCONTRADO"))'),
        (rCmp, "ID resolvido", f'=IF($B${rCmp}="","",IFERROR(INDEX(CM_ID,MATCH($B${rCmp},CM_NOME,0)),"NAO ENCONTRADO"))'),
        (rNum, "Parcela resolvida", f'=IF($D${rCmp}="","",IF($B${rNum}="",'
                                   f'IFERROR(INDEX(P_ID,MATCH(1,INDEX((P_CMP=$D${rCmp})*(P_STATUS<>"PAGA")*(P_STATUS<>"CANCELADA"),0),0)),""),'
                                   f'IFERROR(INDEX(P_ID,MATCH(1,INDEX((P_CMP=$D${rCmp})*(P_NUM=$B${rNum}),0),0)),"NAO ENCONTRADA")))'),
        (rMeta, "ID resolvido", f'=IF($B${rMeta}="","",IFERROR(INDEX(M_ID,MATCH($B${rMeta},M_NOME,0)),"NAO ENCONTRADO"))'),
    ]
    for rr, lab, f in res:
        ws.cell(rr, 3, lab).font = K.F_PEQ
        c = ws.cell(rr, 4, f)
        c.font = K.F_LINK
        c.fill = K.FILL_CALC
        c.border = K.BORDA
    # informacoes da parcela escolhida
    ws.cell(rNum + 1, 3, "Rotulo / vencimento / valor da parcela").font = K.F_PEQ
    ws.cell(rNum + 1, 4, f'=IF($D${rNum}="","",IFERROR(INDEX(P_ROT,MATCH($D${rNum},P_ID,0))&"  venc "&DAY(INDEX(P_VENC,MATCH($D${rNum},P_ID,0)))&"/"&MONTH(INDEX(P_VENC,MATCH($D${rNum},P_ID,0)))&"/"&YEAR(INDEX(P_VENC,MATCH($D${rNum},P_ID,0)))&"  Gs. "&TEXT(INDEX(P_VALOR,MATCH($D${rNum},P_ID,0)),"#,##0"),"?"))').font = K.F_LINK
    r = r_ini + len(campos) + 1

    r = secao(ws, r, "VERIFICACAO ANTES DE COLAR", 8)
    ver = [
        ("Classificacao economica do lancamento",
         f'=IFERROR(INDEX(MAP_CLASS,MATCH($B${rTipo},MAP_TIPO,0)),"TIPO INVALIDO")'),
        ("Efeito no caixa",
         f'=IF(AND($B${rDT}="CONTA",$B${rOT}<>"CONTA"),"ENTRADA de caixa",'
         f'IF(AND($B${rOT}="CONTA",$B${rDT}<>"CONTA"),"SAIDA de caixa","NAO afeta o caixa"))'),
        ("Situacao do preenchimento",
         f'=IF($B${rVal}<=0,"ERRO: valor deve ser maior que zero",'
         f'IF(AND($B${rOT}="EXTERNO",$B${rDT}="EXTERNO"),"ERRO: origem e destino externos",'
         f'IF(AND($B${rOT}<>"EXTERNO",$D${rOT}=""),"ERRO: origem nao resolvida",'
         f'IF(AND($B${rDT}<>"EXTERNO",$D${rDT}=""),"ERRO: destino nao resolvido",'
         f'IF($D${rCat}="","ERRO: categoria nao resolvida",'
         f'IF(AND($B${rSub}<>"",$D${rSub}="NAO ENCONTRADO"),"ERRO: subcategoria nao pertence a categoria",'
         f'IF(AND($D${rCmp}<>"",$D${rNum}=""),"ATENCAO: compromisso sem parcela em aberto","PRONTO PARA COLAR")))))))'),
        ("Aviso de vinculo",
         f'=IF(AND($D${rCmp}="",$D${rMeta}="",$B${rRec}="",$B${rRel}=""),"Sem vinculos",'
         f'"ATENCAO: este lancamento fica vinculado a "&TRIM(IF($D${rCmp}<>"","compromisso/parcela ","")'
         f'&IF($D${rMeta}<>"","meta ","")&IF($B${rRec}<>"","recorrencia ","")&IF($B${rRel}<>"","estorno","")))'),
    ]
    for i, (lab, f) in enumerate(ver):
        rotulo(ws, r + i, 1, lab)
        c = ws.cell(r + i, 2, f)
        c.font = K.F_BOLD
        c.fill = K.FILL_CALC
        c.border = K.BORDA
        ws.merge_cells(start_row=r + i, start_column=2, end_row=r + i, end_column=5)
    r += len(ver) + 1

    r = secao(ws, r, "LINHA PRONTA PARA COPIAR (COLE EM LANCAMENTOS, COLUNAS B ATE S)", 20)
    hdr = [h for h, _w, _k in COLS_LANC[1:19]]
    r = cabecalho(ws, r, [(h, 17) for h in hdr])
    vals = [f"=$B${rD}", f"=$B${rDesc}", f"=$B${rTipo}", f"=$B${rSt}", f"=$B${rVal}",
            f"=$B${rOT}", f"=$D${rOT}", f"=$B${rDT}", f"=$D${rDT}", f"=$D${rCat}",
            f"=$D${rSub}", f"=$B${rForma}", f"=$D${rCmp}", f"=$D${rNum}", f"=$D${rMeta}",
            f"=$B${rRec}", f"=$B${rRel}", f"=$B${rObs}"]
    for j, f in enumerate(vals, start=1):
        c = ws.cell(r, j, f)
        c.font = K.F_BOLD
        c.border = K.BORDA
        c.fill = K.FILL_OK
        if j == 1:
            c.number_format = K.FMT_DATA
        if j == 5:
            c.number_format = K.FMT_GS
    r += 2

    r = secao(ws, r, "ASSISTENTE DE RESGATE DE INVESTIMENTO (GERA DOIS LANCAMENTOS)", 20)
    r = nota(ws, r, "Regra: o rendimento e reconhecido como RECEITA FINANCEIRA (lancamento 1) e "
                    "o resgate transfere o dinheiro do investimento para a conta sem virar receita "
                    "(lancamento 2). Se o rendimento ja foi lancado mes a mes, informe 0 no campo "
                    "Rendimento apurado e use apenas o lancamento 2.", 20)
    rg = r
    for lab, val, fmt, dv in [("Investimento (nome)", ("" if MODO_LIMPO else "Fondo Mutuo Atlas Renta"), None, "LST_NOMES_INVESTIMENTO"),
                              ("Conta de destino (nome)", ("" if MODO_LIMPO else "Cuenta Corriente Ueno"), None, "LST_NOMES_CONTA"),
                              ("Data do resgate", K.HOJE, K.FMT_DATA, None),
                              ("Valor total resgatado", 3000000, K.FMT_GS, None),
                              ("Rendimento apurado ainda nao lancado", 0, K.FMT_GS, None)]:
        rotulo(ws, r, 1, lab)
        c = valor(ws, r, 2, val, fmt, entrada=True)
        c.alignment = K.AL_L
        if dv:
            dv_lista(ws, dv, f"B{r}")
        r += 1
    ws.cell(rg, 3, "ID").font = K.F_PEQ
    ws.cell(rg, 4, f'=IFERROR(INDEX(I_ID,MATCH($B${rg},I_NOME,0)),"NAO ENCONTRADO")').font = K.F_LINK
    ws.cell(rg + 1, 3, "ID").font = K.F_PEQ
    ws.cell(rg + 1, 4, f'=IFERROR(INDEX(C_ID,MATCH($B${rg + 1},C_NOME,0)),"NAO ENCONTRADO")').font = K.F_LINK
    ws.cell(rg + 4, 3, "Principal do resgate").font = K.F_PEQ
    ws.cell(rg + 4, 4, f"=MAX(0,$B${rg + 3}-$B${rg + 4})").number_format = K.FMT_GS
    r += 1
    r = cabecalho(ws, r, [(h, 17) for h in hdr])
    linhas = [
        [f"=$B${rg + 2}", '="Rendimento de "&$B$' + str(rg), '="RENDIMENTO"', '="REALIZADO"',
         f"=$B${rg + 4}", '="EXTERNO"', '=""', '="INVESTIMENTO"', f"=$D${rg}",
         '=IFERROR(INDEX(CAT_ID,MATCH("Receitas",CAT_NOME,0)),"")',
         '=IFERROR(INDEX(SUB_ID,MATCH(1,INDEX((SUB_NOME="Rendimento financeiro"),0),0)),"")',
         '="Transferencia"', '=""', '=""', '=""', '=""', '=""',
         '="Rendimento reconhecido no momento do resgate"'],
        [f"=$B${rg + 2}", '="Resgate de "&$B$' + str(rg), '="RESGATE"', '="REALIZADO"',
         f"=$B${rg + 3}", '="INVESTIMENTO"', f"=$D${rg}", '="CONTA"', f"=$D${rg + 1}",
         '=IFERROR(INDEX(CAT_ID,MATCH("Investimentos",CAT_NOME,0)),"")',
         '=IFERROR(INDEX(SUB_ID,MATCH(1,INDEX((SUB_NOME="Resgate"),0),0)),"")',
         '="Transferencia"', '=""', '=""', '=""', '=""', '=""',
         '="Operacao neutra: principal nao e receita"'],
    ]
    for i, linha in enumerate(linhas):
        for j, f in enumerate(linha, start=1):
            c = ws.cell(r + i, j, f)
            c.font = K.F_NORMAL
            c.border = K.BORDA
            c.fill = K.FILL_OK
            if j == 1:
                c.number_format = K.FMT_DATA
            if j == 5:
                c.number_format = K.FMT_GS
    r += len(linhas) + 2

    r = secao(ws, r, "ASSISTENTE DE MOVIMENTO DE META (COLE EM MOV_METAS, COLUNAS B ATE H)", 20)
    rm = r
    for lab, val, fmt, dv in [("Meta (nome)", ("" if MODO_LIMPO else "Comprar auto"), None, "LST_METAS_NOME"),
                              ("Tipo de movimento", "RESERVA", None, "LST_TIPOMOV"),
                              ("Data", K.HOJE, K.FMT_DATA, None),
                              ("Valor", 1500000, K.FMT_GS, None),
                              ("Destino da retirada (se for saida)", "", None, "LST_DESTRET"),
                              ("Meta de destino (se for transferencia)", "", None, "LST_METAS_NOME"),
                              ("Observacao", "Reserva mensal", None, None)]:
        rotulo(ws, r, 1, lab)
        c = valor(ws, r, 2, val, fmt, entrada=True)
        c.alignment = K.AL_L
        if dv:
            dv_lista(ws, dv, f"B{r}")
        r += 1
    r += 1
    r = cabecalho(ws, r, [(h, 18) for h, _w, _k in COLS_MOV[1:8]])
    mov = [f"=$B${rm + 2}", f'=IFERROR(INDEX(M_ID,MATCH($B${rm},M_NOME,0)),"NAO ENCONTRADO")',
           f"=$B${rm + 1}", f"=$B${rm + 3}", f"=$B${rm + 4}",
           f'=IF($B${rm + 5}="","",IFERROR(INDEX(M_ID,MATCH($B${rm + 5},M_NOME,0)),"NAO ENCONTRADO"))',
           f"=$B${rm + 6}"]
    for j, f in enumerate(mov, start=1):
        c = ws.cell(r, j, f)
        c.font = K.F_BOLD
        c.border = K.BORDA
        c.fill = K.FILL_OK
        if j == 1:
            c.number_format = K.FMT_DATA
        if j == 4:
            c.number_format = K.FMT_GS
    return ws


# ---------------------------------------------------------------------------
# GERADOR - CRONOGRAMAS E RECORRENCIAS
# ---------------------------------------------------------------------------
def aba_gerador(L):
    ws = L.aba("GERADOR", "ED7D31")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20
    for i in range(5, 26):
        ws.column_dimensions[CL(i)].width = 17
    r = titulo(ws, 1, "GERADOR - CRONOGRAMAS DE PARCELAS E OCORRENCIAS DE RECORRENCIAS", 20,
               "Escolha o registro nas celulas amarelas, confira as linhas geradas e copie o "
               "bloco verde para o fim da aba de destino. Nada e gravado automaticamente.")
    r += 1

    # ---- 1) cronograma de parcelas
    r = secao(ws, r, "1. CRONOGRAMA DE PARCELAS (COLE EM PARCELAS, COLUNAS B ATE E)", 20)
    rc = r
    rotulo(ws, r, 1, "Compromisso")
    c = valor(ws, r, 2, ("" if MODO_LIMPO else "Terreno Santa Rita"), None, entrada=True)
    c.alignment = K.AL_L
    dv_lista(ws, "LST_COMPROMISSOS", f"B{r}")
    ws.cell(r, 3, "ID").font = K.F_PEQ
    ws.cell(r, 4, f'=IFERROR(INDEX(CM_ID,MATCH($B${r},CM_NOME,0)),"NAO ENCONTRADO")').font = K.F_LINK
    r += 1
    for lab, f, fmt in [("Quantidade de parcelas", f'=IFERROR(INDEX(CM_QTD,MATCH($D${rc},CM_ID,0)),0)', "0"),
                        ("Valor da parcela", f'=IFERROR(INDEX(CM_VLRPARC,MATCH($D${rc},CM_ID,0)),0)', K.FMT_GS),
                        ("Data da primeira parcela", f'=IFERROR(INDEX(CAD_Compromissos!$J${DAT}:$J${DAT + K.MAX_CAD - 1},MATCH($D${rc},CM_ID,0)),"")', K.FMT_DATA),
                        ("Periodicidade", f'=IFERROR(INDEX(CAD_Compromissos!$L${DAT}:$L${DAT + K.MAX_CAD - 1},MATCH($D${rc},CM_ID,0)),"")', None)]:
        rotulo(ws, r, 1, lab)
        cc = ws.cell(r, 2, f)
        cc.number_format = fmt or "General"
        cc.fill = K.FILL_CALC
        cc.border = K.BORDA
        r += 1
    rq, rv, rp, rper = rc + 1, rc + 2, rc + 3, rc + 4
    passo = (f'IF($B${rper}="Mensal",1,IF($B${rper}="Bimestral",2,IF($B${rper}="Trimestral",3,'
             f'IF($B${rper}="Semestral",6,IF($B${rper}="Anual",12,0)))))')
    dias = f'IF($B${rper}="Diario",1,IF($B${rper}="Semanal",7,IF($B${rper}="Quinzenal",15,0)))'
    r += 1
    NP = 120
    r = cabecalho(ws, r, [("ID_Compromisso", 16), ("Num_Parcela", 13), ("Data_Vencimento", 16),
                          ("Valor_Parcela", 16)])
    for i in range(NP):
        rr = r + i
        cond = f'IF(OR({i + 1}>N($B${rq}),NOT(ISNUMBER($B${rp}))),"",'
        vals = [
            (1, f'={cond}$D${rc})', None),
            (2, f'={cond}{i + 1})', "0"),
            (3, f'={cond}IF({passo}>0,EDATE($B${rp},{i}*{passo}),$B${rp}+{i}*{dias}))', K.FMT_DATA),
            (4, f'={cond}$B${rv})', K.FMT_GS),
        ]
        for col, f, fmt in vals:
            cc = ws.cell(rr, col, f)
            cc.font = K.F_NORMAL
            cc.border = K.BORDA
            cc.fill = K.FILL_OK
            if fmt:
                cc.number_format = fmt
    r += NP + 1

    # ---- 2) ocorrencias de recorrencia
    r = secao(ws, r, "2. OCORRENCIAS DE UMA RECORRENCIA (COLE EM LANCAMENTOS, COLUNAS B ATE S)", 20)
    rr0 = r
    rotulo(ws, r, 1, "Recorrencia (ID)")
    c = valor(ws, r, 2, ("" if MODO_LIMPO else "REC-000001"), None, entrada=True)
    c.alignment = K.AL_L
    dv_lista(ws, "R_ID", f"B{r}")
    ws.cell(r, 3, "Descricao").font = K.F_PEQ
    ws.cell(r, 4, f'=IFERROR(INDEX(R_DESC,MATCH($B${r},R_ID,0)),"NAO ENCONTRADA")').font = K.F_LINK
    r += 1
    campos_rec = [("Tipo de operacao", "C"), ("Valor", "D"), ("Periodicidade", "E"),
                  ("Dia", "F"), ("Data de inicio", "G"), ("Qtd de ocorrencias", "H"),
                  ("Categoria_ID", "J"), ("Subcategoria_ID", "K"), ("Origem_Tipo", "L"),
                  ("Origem_ID", "M"), ("Destino_Tipo", "N"), ("Destino_ID", "O"),
                  ("Forma de pagamento", "P")]
    linhas_rec = {}
    for lab, col in campos_rec:
        rotulo(ws, r, 1, lab)
        cc = ws.cell(r, 2, f'=IFERROR(INDEX(CAD_Recorrencias!${col}${DAT}:${col}${DAT + K.MAX_CAD - 1},'
                           f'MATCH($B${rr0},R_ID,0)),"")')
        cc.fill = K.FILL_CALC
        cc.border = K.BORDA
        if lab == "Valor":
            cc.number_format = K.FMT_GS
        if lab == "Data de inicio":
            cc.number_format = K.FMT_DATA
        linhas_rec[lab] = r
        r += 1
    LR = linhas_rec
    passo2 = (f'IF($B${LR["Periodicidade"]}="Mensal",1,IF($B${LR["Periodicidade"]}="Bimestral",2,'
              f'IF($B${LR["Periodicidade"]}="Trimestral",3,IF($B${LR["Periodicidade"]}="Semestral",6,'
              f'IF($B${LR["Periodicidade"]}="Anual",12,0)))))')
    dias2 = (f'IF($B${LR["Periodicidade"]}="Diario",1,IF($B${LR["Periodicidade"]}="Semanal",7,'
             f'IF($B${LR["Periodicidade"]}="Quinzenal",15,0)))')
    r += 1
    NR = 36
    hdr = [h for h, _w, _k in COLS_LANC[1:19]]
    r = cabecalho(ws, r, [(h, 17) for h in hdr])
    for i in range(NR):
        rr = r + i
        cond = (f'IF(OR({i + 1}>N($B${LR["Qtd de ocorrencias"]}),NOT(ISNUMBER($B${LR["Data de inicio"]}))),"",')
        data = (f'IF({passo2}>0,EDATE($B${LR["Data de inicio"]},{i}*{passo2}),'
                f'$B${LR["Data de inicio"]}+{i}*{dias2})')
        vals = [
            f'={cond}{data})',
            f'={cond}$D${rr0})',
            f'={cond}$B${LR["Tipo de operacao"]})',
            f'={cond}IF({data}<=CFG_HOJE,"REALIZADO","PROJETADO"))',
            f'={cond}$B${LR["Valor"]})',
            f'={cond}$B${LR["Origem_Tipo"]})',
            f'={cond}$B${LR["Origem_ID"]})',
            f'={cond}$B${LR["Destino_Tipo"]})',
            f'={cond}$B${LR["Destino_ID"]})',
            f'={cond}$B${LR["Categoria_ID"]})',
            f'={cond}$B${LR["Subcategoria_ID"]})',
            f'={cond}$B${LR["Forma de pagamento"]})',
            '=""', '=""', '=""',
            f'={cond}$B${rr0})',
            '=""',
            f'={cond}"Gerado pela recorrencia "&$B${rr0})',
        ]
        for j, f in enumerate(vals, start=1):
            cc = ws.cell(rr, j, f)
            cc.font = K.F_NORMAL
            cc.border = K.BORDA
            cc.fill = K.FILL_OK
            if j == 1:
                cc.number_format = K.FMT_DATA
            if j == 5:
                cc.number_format = K.FMT_GS
    return ws


# ---------------------------------------------------------------------------
# TESTES
# ---------------------------------------------------------------------------
def aba_testes(L):
    ws = L.aba("TESTES", "00B050")
    for c, w in [("A", 10), ("B", 20), ("C", 56), ("D", 22), ("E", 22), ("F", 18), ("G", 14)]:
        ws.column_dimensions[c].width = w
    r = titulo(ws, 1, "TESTES - VALIDACAO MATEMATICA E DE INTEGRIDADE", 7,
               "Cada teste recalcula o valor esperado por um caminho diferente (SUMPRODUCT) e "
               "compara com o valor que o sistema apresenta (SUMIFS). Tolerancia: Gs. 0,50 - como o Guarani nao tem centavos, qualquer diferenca real reprova o teste.")
    r += 1
    r_tot = r
    rotulo(ws, r, 1, "TESTES COM FALHA")
    rotulo(ws, r, 4, "TESTES EXECUTADOS")
    r += 2
    faixas_res = []
    hdr = [("ID", 10), ("Grupo", 20), ("Teste", 56), ("Valor_Esperado", 22),
           ("Valor_Obtido", 22), ("Diferenca", 18), ("Resultado", 14)]

    def bloco(r, titulo_txt, nlin, gen, prefixo):
        nonlocal faixas_res
        r = secao(ws, r, titulo_txt, 7)
        ini = r + 1
        r = _grade(ws, r, hdr, nlin, gen)
        for i in range(nlin):
            rr = ini + i
            ws.cell(rr, 1, f"{prefixo}{i + 1:02d}").alignment = K.AL_C
            ws.cell(rr, 6, f'=IF($C{rr}="","",$E{rr}-$D{rr})').number_format = K.FMT_GS
            ws.cell(rr, 7, f'=IF($C{rr}="","",IF(ABS($F{rr})<=0.5,"OK","FALHA"))')
            for col in (6, 7):
                ws.cell(rr, col).border = K.BORDA
                ws.cell(rr, col).fill = K.FILL_CALC
                ws.cell(rr, col).font = K.F_BOLD
        ws.conditional_formatting.add(f"G{ini}:G{r - 1}", CellIsRule(
            operator="equal", formula=['"FALHA"'], fill=K.FILL_ERRO))
        ws.conditional_formatting.add(f"G{ini}:G{r - 1}", CellIsRule(
            operator="equal", formula=['"OK"'], fill=K.FILL_OK))
        faixas_res.append(f"$G${ini}:$G${r - 1}")
        return r + 1

    n = K.MAX_CAD
    SPD = '(L_DestTipo="{t}")*(L_DestID={id})*(L_Status="REALIZADO")*L_Valor'
    SPO = '(L_OrigTipo="{t}")*(L_OrigID={id})*(L_Status="REALIZADO")*L_Valor'

    # A - contas
    def genA(i, rr):
        a = f"CAD_Contas!$A{DAT + i}"
        return [
            (2, f'=IF({a}="","","Saldo de conta")', None),
            (3, f'=IF({a}="","","Saldo inicial + entradas - saidas = saldo calculado: "&CAD_Contas!$B{DAT + i})', None),
            (4, f'=IF({a}="","",CAD_Contas!$F{DAT + i}+SUMPRODUCT({SPD.format(t="CONTA", id=a)})-SUMPRODUCT({SPO.format(t="CONTA", id=a)}))', K.FMT_GS),
            (5, f'=IF({a}="","",CAD_Contas!$J{DAT + i})', K.FMT_GS),
        ]
    r = bloco(r, "A. SALDO DE CADA CONTA (saldo inicial + entradas - saidas)", n, genA, "A")

    # B - cartoes
    def genB(i, rr):
        a = f"CAD_Cartoes!$A{DAT + i}"
        return [
            (2, f'=IF({a}="","","Divida de cartao")', None),
            (3, f'=IF({a}="","","Divida inicial + compras - pagamentos = divida atual: "&CAD_Cartoes!$B{DAT + i})', None),
            (4, f'=IF({a}="","",CAD_Cartoes!$F{DAT + i}+SUMPRODUCT({SPO.format(t="CARTAO", id=a)})-SUMPRODUCT({SPD.format(t="CARTAO", id=a)}))', K.FMT_GS),
            (5, f'=IF({a}="","",CAD_Cartoes!$L{DAT + i})', K.FMT_GS),
        ]
    r = bloco(r, "B. DIVIDA DE CADA CARTAO (divida inicial + compras - pagamentos)", n, genB, "B")

    # C - investimentos (saldo)
    def genC(i, rr):
        a = f"CAD_Investimentos!$A{DAT + i}"
        return [
            (2, f'=IF({a}="","","Saldo de investimento")', None),
            (3, f'=IF({a}="","","Aplicado + aportes + rendimentos - resgates = saldo: "&CAD_Investimentos!$B{DAT + i})', None),
            (4, f'=IF({a}="","",CAD_Investimentos!$F{DAT + i}+SUMPRODUCT({SPD.format(t="INVESTIMENTO", id=a)})-SUMPRODUCT({SPO.format(t="INVESTIMENTO", id=a)}))', K.FMT_GS),
            (5, f'=IF({a}="","",CAD_Investimentos!$O{DAT + i})', K.FMT_GS),
        ]
    r = bloco(r, "C. SALDO DE CADA INVESTIMENTO", n, genC, "C")

    # D - investimentos (principal + rendimento)
    def genD(i, rr):
        a = f"CAD_Investimentos!$A{DAT + i}"
        return [
            (2, f'=IF({a}="","","Principal + rendimento")', None),
            (3, f'=IF({a}="","","Principal + rendimento no saldo = saldo atual: "&CAD_Investimentos!$B{DAT + i})', None),
            (4, f'=IF({a}="","",CAD_Investimentos!$P{DAT + i}+CAD_Investimentos!$Q{DAT + i})', K.FMT_GS),
            (5, f'=IF({a}="","",CAD_Investimentos!$O{DAT + i})', K.FMT_GS),
        ]
    r = bloco(r, "D. DECOMPOSICAO PRINCIPAL x RENDIMENTO", n, genD, "D")

    # E - metas
    def genE(i, rr):
        a = f"CAD_Metas!$A{DAT + i}"
        return [
            (2, f'=IF({a}="","","Saldo de meta")', None),
            (3, f'=IF({a}="","","Reservas - retiradas = valor reservado: "&CAD_Metas!$B{DAT + i})', None),
            (4, f'=IF({a}="","",SUMPRODUCT((MM_META={a})*MM_VALSINAL))', K.FMT_GS),
            (5, f'=IF({a}="","",CAD_Metas!$H{DAT + i})', K.FMT_GS),
        ]
    r = bloco(r, "E. SALDO DE CADA META (reserva logica)", n, genE, "E")

    # F - compromissos
    def genF(i, rr):
        a = f"CAD_Compromissos!$A{DAT + i}"
        return [
            (2, f'=IF({a}="","","Compromisso")', None),
            (3, f'=IF({a}="","","Valor pago + saldo devedor = valor total: "&CAD_Compromissos!$B{DAT + i})', None),
            (4, f'=IF({a}="","",CAD_Compromissos!$I{DAT + i})', K.FMT_GS),
            (5, f'=IF({a}="","",CAD_Compromissos!$U{DAT + i}+CAD_Compromissos!$V{DAT + i})', K.FMT_GS),
        ]
    r = bloco(r, "F. COMPROMISSOS (valor pago + saldo devedor = valor total)", n, genF, "F")

    # G - integridade financeira e duplicidade
    testes = [
        ("Integridade", "Transferencia entre contas nao gera receita",
         "=0", '=SUMIFS(L_VRec,L_Tipo,"TRANSFERENCIA")', "0"),
        ("Integridade", "Transferencia entre contas nao gera despesa",
         "=0", '=SUMIFS(L_VDesp,L_Tipo,"TRANSFERENCIA")', "0"),
        ("Integridade", "Aporte em investimento nao e despesa",
         "=0", '=SUMIFS(L_VDesp,L_Tipo,"APORTE")', "0"),
        ("Integridade", "Resgate de investimento nao e receita (principal)",
         "=0", '=SUMIFS(L_VRec,L_Tipo,"RESGATE")', "0"),
        ("Integridade", "Pagamento de fatura de cartao nao e despesa",
         "=0", '=SUMIFS(L_VDesp,L_Tipo,"PAGAMENTO_CARTAO")', "0"),
        ("Integridade", "Compra no cartao e integralmente classificada como despesa",
         '=SUMIFS(L_Valor,L_Tipo,"COMPRA_CARTAO")', '=SUMIFS(L_VDesp,L_Tipo,"COMPRA_CARTAO")', K.FMT_GS),
        ("Integridade", "Rendimento e integralmente receita financeira",
         '=SUMIFS(L_Valor,L_Tipo,"RENDIMENTO")', '=SUMIFS(L_VRend,L_Tipo,"RENDIMENTO")', K.FMT_GS),
        ("Integridade", "Parcela de compromisso e integralmente despesa",
         '=SUMIFS(L_Valor,L_Tipo,"PAGAMENTO_PARCELA")', '=SUMIFS(L_VDesp,L_Tipo,"PAGAMENTO_PARCELA")', K.FMT_GS),
        ("Realizado x projetado", "Saldo das contas usa apenas lancamentos REALIZADOS",
         '=SUM(C_SALDOINI)+SUMPRODUCT((L_DestTipo="CONTA")*(L_Status="REALIZADO")*L_Valor)'
         '-SUMPRODUCT((L_OrigTipo="CONTA")*(L_Status="REALIZADO")*L_Valor)',
         "=SUM(C_SALDO)", K.FMT_GS),
        ("Realizado x projetado", "Saldo projetado = saldo realizado + movimentos projetados",
         '=SUM(C_SALDO)+SUMPRODUCT((L_DestTipo="CONTA")*(L_Status="PROJETADO")*L_Valor)'
         '-SUMPRODUCT((L_OrigTipo="CONTA")*(L_Status="PROJETADO")*L_Valor)',
         "=SUM(C_PROJ)", K.FMT_GS),
        ("Realizado x projetado", "Total de todos os status menos projetados e cancelados = saldo realizado",
         "=SUM(C_SALDO)",
         '=SUM(C_SALDOINI)+SUMPRODUCT((L_DestTipo="CONTA")*L_Valor)-SUMPRODUCT((L_OrigTipo="CONTA")*L_Valor)'
         '-SUMPRODUCT((L_Status<>"REALIZADO")*((L_DestTipo="CONTA")-(L_OrigTipo="CONTA"))*L_Valor)', K.FMT_GS),
        ("Realizado x projetado", "Todo lancamento tem um status valido (realizado+projetado+cancelado)",
         "=SUM(L_Valor)",
         '=SUMIFS(L_Valor,L_Status,"REALIZADO")+SUMIFS(L_Valor,L_Status,"PROJETADO")'
         '+SUMIFS(L_Valor,L_Status,"CANCELADO")', K.FMT_GS),
        ("Patrimonio", "Ativos - passivos = patrimonio liquido total",
         "=PAT_ATIVOS-PAT_PASSIVOS", "=PAT_PL_TOTAL", K.FMT_GS),
        ("Patrimonio", "Patrimonio financeiro = patrimonio total - bens",
         "=PAT_PL_TOTAL-SUM(B_VALOR)", "=PAT_PL_FIN", K.FMT_GS),
        ("Metas", "Reserva de meta nao altera o saldo fisico da conta",
         "=SUM(C_SALDO)",
         '=SUM(C_SALDOINI)+SUMPRODUCT((L_DestTipo="CONTA")*(L_Status="REALIZADO")*L_Valor)'
         '-SUMPRODUCT((L_OrigTipo="CONTA")*(L_Status="REALIZADO")*L_Valor)', K.FMT_GS),
        ("Metas", "Saldo disponivel = saldo calculado - reservado",
         "=SUM(C_SALDO)-SUM(C_RESERV)", "=SUM(C_DISP)", K.FMT_GS),
        ("Metas", "Reservado nas contas = soma das metas vinculadas a contas",
         '=SUMPRODUCT((MM_CONTA<>"")*MM_VALSINAL)', "=SUM(C_RESERV)", K.FMT_GS),
        ("Metas", "Meta cancelada ou concluida fica com saldo reservado zerado",
         "=0",
         '=SUMIFS(M_RESERV,M_STATUS,"Cancelada")+SUMIFS(M_RESERV,M_STATUS,"Concluida")', K.FMT_GS),
        ("Parcelas", "Soma dos status de parcela = total de parcelas cadastradas",
         '=SUMPRODUCT((P_ID<>"")*1)',
         '=COUNTIF(P_STATUS,"PAGA")+COUNTIF(P_STATUS,"ABERTA")+COUNTIF(P_STATUS,"ATRASADA")'
         '+COUNTIF(P_STATUS,"PROJETADA")+COUNTIF(P_STATUS,"CANCELADA")', "0"),
        ("Parcelas", "Nenhuma parcela com mais de um pagamento realizado",
         "=0", "=SUMPRODUCT((P_QTDREAL>1)*1)", "0"),
        ("Parcelas", "Nenhuma parcela paga mantem lancamento projetado",
         "=0", "=SUMPRODUCT((P_QTDREAL>0)*(P_QTDPROJ>0)*1)", "0"),
        ("Parcelas", "Valor pago das parcelas confere com os lancamentos vinculados",
         '=SUMPRODUCT((P_ANTES="NAO")*(P_STATUS="PAGA")*P_VALOR)',
         '=SUMIFS(L_Valor,L_ParcID,"<>",L_Status,"REALIZADO")', K.FMT_GS),
        ("Duplicidade", "Todo lancamento preenchido possui ID gerado",
         '=SUMPRODUCT((L_Data<>"")*1)', '=SUMPRODUCT((L_ID<>"")*1)', "0"),
        ("Duplicidade", "Toda parcela preenchida possui ID gerado",
         '=SUMPRODUCT((P_VENC<>"")*1)', '=SUMPRODUCT((P_ID<>"")*1)', "0"),
        ("Duplicidade", "Nenhum lancamento com erro de validacao",
         "=0", '=SUMPRODUCT((LEFT(L_Valid,4)="ERRO")*1)', "0"),
        ("Duplicidade", "Nenhum movimento de meta com erro de validacao",
         "=0", '=SUMPRODUCT((LEFT(MM_VALID,4)="ERRO")*1)', "0"),
        ("Fluxo de caixa", "Entradas do fluxo mensal = entradas dos lancamentos na grade",
         '=SUMIFS(L_Ent,L_Data,">="&INDEX(FX_INI,1),L_Status,"REALIZADO")',
         "=SUM(FX_ENTR)", K.FMT_GS),
        ("Fluxo de caixa", "Saldo final da grade mensal = saldo projetado total das contas",
         "=SUM(C_PROJ)", "=INDEX(FX_SALDOFIM,COUNT(FX_SALDOFIM))", K.FMT_GS),
        ("Fluxo de caixa", "Transferencia entre contas proprias nao entra no fluxo de caixa",
         "=0",
         '=SUMPRODUCT((L_Tipo="TRANSFERENCIA")*(L_OrigTipo="CONTA")*(L_DestTipo="CONTA")*(L_Ent+L_Sai))', "0"),
        ("Investimentos", "Rendimento realizado e rendimento projetado ficam separados",
         '=SUMIFS(L_Valor,L_Tipo,"RENDIMENTO",L_Status,"REALIZADO")', "=SUM(I_REND)", K.FMT_GS),
        ("Investimentos", "Rendimento projetado nao entra no rendimento realizado",
         '=SUMIFS(L_Valor,L_Tipo,"RENDIMENTO",L_Status,"PROJETADO")', "=SUM(I_RENDPROJ)", K.FMT_GS),
        ("Projecao", "Saldo do 1o mes projetado = saldo real + projetado ate o fim do mes",
         "=" + SALDO_BASE.format(d="INDEX(PJ_FIM,1)+1"), "=INDEX(PJ_SALDO,1)", K.FMT_GS),
        ("Projecao", "Saldo do ultimo mes projetado = saldo projetado total das contas",
         "=SUM(C_PROJ)", "=INDEX(PJ_SALDO,COUNT(PJ_SALDO))", K.FMT_GS),
        ("Cartoes", "Fatura atual esta dentro do ciclo aberto de cada cartao",
         "=0", '=SUMPRODUCT((K_FATURA<0)*1)', "0"),
        ("Relatorios", "Resultado do mes = receitas + rendimentos - despesas",
         '=SUMIFS(L_VRec,L_Data,">="&REL_INI,L_Data,"<="&REL_FIM,L_Status,"REALIZADO")'
         '+SUMIFS(L_VRend,L_Data,">="&REL_INI,L_Data,"<="&REL_FIM,L_Status,"REALIZADO")'
         '-SUMIFS(L_VDesp,L_Data,">="&REL_INI,L_Data,"<="&REL_FIM,L_Status,"REALIZADO")',
         "=REL_RESULTADO-" +
         '(SUMIFS(L_VRec,L_Data,">="&REL_INI,L_Data,"<="&REL_FIM,L_Status,"PROJETADO")'
         '+SUMIFS(L_VRend,L_Data,">="&REL_INI,L_Data,"<="&REL_FIM,L_Status,"PROJETADO")'
         '-SUMIFS(L_VDesp,L_Data,">="&REL_INI,L_Data,"<="&REL_FIM,L_Status,"PROJETADO"))', K.FMT_GS),
        ("Relatorios", "Despesa do ano por categoria = despesa total do ano",
         '=SUMIFS(L_VDesp,L_Ano,RELA_ANO,L_Status,"REALIZADO")'
         '+SUMIFS(L_VDesp,L_Ano,RELA_ANO,L_Status,"PROJETADO")',
         "=SUM(RA_CAT_TOT)", K.FMT_GS),
        ("Relatorios", "Soma dos 12 meses do relatorio anual = despesa do ano",
         '=SUMIFS(L_VDesp,L_Ano,RELA_ANO,L_Status,"REALIZADO")'
         '+SUMIFS(L_VDesp,L_Ano,RELA_ANO,L_Status,"PROJETADO")',
         "=SUM(RA_DESP)", K.FMT_GS),
    ]

    def genG(i, rr):
        t = testes[i]
        return [(2, t[0], None), (3, t[1], None), (4, t[2], t[4]), (5, t[3], t[4])]

    r = bloco(r, "G. INTEGRIDADE FINANCEIRA, DUPLICIDADE E RELATORIOS", len(testes), genG, "G")

    total_falhas = "=" + "+".join(f'COUNTIF({f},"FALHA")' for f in faixas_res)
    total_exec = "=" + "+".join(f'COUNTIF({f},"OK")+COUNTIF({f},"FALHA")' for f in faixas_res)
    c = ws.cell(r_tot, 2, total_falhas)
    c.number_format = "0"
    c.font = K.F_KPI
    c.border = K.BORDA
    ws.conditional_formatting.add(f"B{r_tot}", CellIsRule(
        operator="greaterThan", formula=["0"], fill=K.FILL_ERRO))
    ws.conditional_formatting.add(f"B{r_tot}", CellIsRule(
        operator="equal", formula=["0"], fill=K.FILL_OK))
    c2 = ws.cell(r_tot, 5, total_exec)
    c2.number_format = "0"
    c2.font = K.F_KPI
    c2.border = K.BORDA
    L.nome("TS_FALHAS", f"TESTES!$B${r_tot}")
    L.nome("TS_EXEC", f"TESTES!$E${r_tot}")
    ws.freeze_panes = "A6"
    return ws


# ---------------------------------------------------------------------------
# MANUAL E ARQUITETURA
# ---------------------------------------------------------------------------
def _texto(L, nome, tit, subtitulo, secoes, cor, larg=(4, 46, 110)):
    ws = L.aba(nome, cor)
    ws.column_dimensions["A"].width = larg[0]
    ws.column_dimensions["B"].width = larg[1]
    ws.column_dimensions["C"].width = larg[2]
    r = titulo(ws, 1, tit, 3, subtitulo)
    r += 1
    for titulo_sec, itens in secoes:
        r = secao(ws, r, titulo_sec, 3)
        for a, b in itens:
            ca = ws.cell(r, 2, a)
            ca.font = K.F_BOLD
            ca.alignment = K.AL_LW
            ca.border = K.BORDA
            cb = ws.cell(r, 3, b)
            cb.font = K.F_NORMAL
            cb.alignment = K.AL_LW
            cb.border = K.BORDA
            ws.row_dimensions[r].height = max(15, 13 * (1 + len(b) // 95))
            r += 1
        r += 1
    return ws


def aba_manual(L):
    secoes = [
        ("0. PRIMEIROS PASSOS (LEIA ANTES DE MEXER)", [
            ("Se voce recebeu o arquivo LIMPO", "Ele ja vem sem nenhum registro: so as categorias e subcategorias padrao. Comece cadastrando suas contas em CAD_Contas, depois cartoes, investimentos, compromissos e metas. So entao comece a lancar."),
            ("Se voce recebeu o arquivo EXEMPLO", "Ele vem com dados ficticios para voce ver o sistema funcionando. Para comeca-lo do zero, siga a instrucao abaixo - NAO apague linha por linha."),
            ("COMO APAGAR OS DADOS DE EXEMPLO", "Em cada aba de dados, selecione o intervalo das celulas AMARELAS (a primeira linha de dados ate a ultima preenchida) e pressione DELETE. Isso limpa o conteudo e mantem as formulas e os IDs no lugar."),
            ("O QUE NUNCA FAZER", "Nunca use 'Excluir linha' (botao direito > Excluir) nem classifique (sort) as abas de dados. Os IDs sao gerados pela posicao da linha: excluir uma linha faz todos os IDs abaixo dela mudarem, e os lancamentos passam a apontar para o registro errado sem nenhum aviso."),
            ("Ordem para limpar", "Limpe primeiro LANCAMENTOS, depois PARCELAS e MOV_METAS, e so por ultimo os cadastros. Assim nenhum lancamento fica apontando para um cadastro que ja sumiu."),
            ("Como saber se algo quebrou", "Abra a aba TESTES: se TESTES COM FALHA estiver diferente de zero, algum calculo esta inconsistente. A aba ALERTAS tambem acusa lancamentos com erro de validacao."),
            ("Se ja quebrou", "Nao tente consertar celula por celula. Pegue uma copia nova do arquivo e recomece: leva menos tempo e nao deixa erro escondido."),
        ]),
        ("1. COMO O SISTEMA FUNCIONA", [
            ("Ideia central", "Existe UM unico livro de operacoes (aba LANCAMENTOS). Todo o resto - saldos, dividas, parcelas, metas, fluxo, relatorios e patrimonio - e calculado a partir dele. Nenhum numero e digitado duas vezes."),
            ("Como um lancamento move dinheiro", "Cada linha tem uma ORIGEM e um DESTINO. Se o destino e uma CONTA, entra dinheiro nela; se a origem e uma CONTA, sai. EXTERNO representa o mundo fora do seu controle (empregador, supermercado)."),
            ("Por que nao ha duplicidade", "Uma transferencia tem conta na origem E no destino: soma numa e subtrai da outra, e nao aparece como receita nem como despesa. Uma compra no cartao tem origem CARTAO, entao vira despesa mas nao mexe no caixa; o caixa so e afetado quando a fatura e paga."),
            ("Cores das celulas", "AMARELO = voce digita. CINZA = calculado, nao digite. VERDE = bloco pronto para copiar. VERMELHO = erro ou alerta."),
            ("Regra de ouro", "Nunca exclua nem classifique (sort) linhas das abas de dados. Os IDs sao gerados pela posicao da linha. Para anular um lancamento, mude o Status para CANCELADO."),
        ]),
        ("2. CADASTROS", [
            ("Cadastrar uma conta", "Aba CAD_Contas: escreva o Nome na primeira linha vazia. O ID (CON-000004, por exemplo) aparece sozinho. Escolha Instituicao e Tipo nas listas."),
            ("Configurar o saldo inicial", "Na mesma linha, preencha Saldo_Inicial e Data_Saldo_Inicial. Use o extrato da data em que voce comecou a usar o sistema. O Saldo_Calculado passa a ser saldo inicial + entradas - saidas."),
            ("Cadastrar um cartao", "Aba CAD_Cartoes: Nome, Instituicao, Limite, Dia_Fechamento e Dia_Vencimento. O sistema calcula sozinho o ciclo aberto, a fatura atual e o vencimento."),
            ("Cadastrar a divida inicial do cartao", "Preencha Divida_Inicial com o saldo em aberto na data de inicio. Depois registre o pagamento dessa fatura como um lancamento PAGAMENTO_CARTAO normal."),
            ("Cadastrar uma instituicao", "Aba CAD_Instituicoes. Assim que o nome existir ali, ele aparece nas listas de contas, cartoes e investimentos."),
            ("Criar categoria", "Aba CAD_Categorias: escreva o Nome. O ID sai sozinho e a categoria ja aparece nas listas de LANCAR."),
            ("Criar subcategoria", "Aba CAD_Subcategorias: escolha a Categoria na lista e escreva o Nome. A subcategoria passa a aparecer automaticamente na lista dependente da tela LANCAR."),
        ]),
        ("3. LANCAR NO DIA A DIA", [
            ("Registrar qualquer operacao", "Va para a aba LANCAR, preencha as celulas amarelas, confira a linha Situacao do preenchimento e copie a LINHA PRONTA para a primeira linha vazia de LANCAMENTOS (colunas B ate S)."),
            ("Registrar uma compra no cartao", "Tipo_Operacao = COMPRA_CARTAO, Origem = CARTAO + o cartao, Destino = EXTERNO. Isso ja e a despesa."),
            ("Pagar a fatura do cartao", "Tipo_Operacao = PAGAMENTO_CARTAO, Origem = CONTA, Destino = CARTAO. Sai dinheiro da conta e a divida do cartao cai. NAO gera uma segunda despesa."),
            ("Registrar uma receita", "Tipo_Operacao = RECEITA, Origem = EXTERNO, Destino = CONTA."),
            ("Fazer uma transferencia", "Tipo_Operacao = TRANSFERENCIA, Origem = CONTA A, Destino = CONTA B. Nao e receita nem despesa."),
            ("Usar PROJETADO e REALIZADO", "PROJETADO e o que voce espera que aconteca; nao altera saldo nenhum. Quando o dinheiro sair de fato, apenas troque o Status da MESMA linha para REALIZADO. Nunca crie uma linha nova - seria duplicidade."),
            ("Corrigir um lancamento", "Edite a propria linha. Se ela tiver algo na coluna Vinculos, leia o aviso antes: alterar pode mudar o status de uma parcela."),
            ("Cancelar / estornar", "Troque o Status para CANCELADO. A linha continua no historico e o efeito financeiro desaparece de todos os calculos. Para estornar com contrapartida, crie um lancamento inverso e aponte o ID original na coluna ID_Relacionado."),
            ("Excluir", "So exclua se for a ultima linha preenchida. Excluir uma linha do meio desloca os IDs de todas as linhas abaixo. Prefira sempre CANCELADO."),
        ]),
        ("4. COMPROMISSOS E PARCELAS", [
            ("Cadastrar um parcelamento", "Aba CAD_Compromissos: Nome, Categoria, Subcategoria, Qtd_Parcelas, Valor_Parcela, Data_Primeira_Parcela, Dia_Vencimento, Periodicidade e a conta ou cartao de pagamento."),
            ("Informar as parcelas ja pagas antes", "No campo Parcelas_Pagas_Antes_Sistema escreva quantas parcelas voce ja quitou antes de comecar a usar a planilha. Elas ficam com status PAGA sem gerar lancamento (o dinheiro ja saiu antes do saldo inicial)."),
            ("Gerar o cronograma", "Aba GERADOR, bloco 1: escolha o compromisso e copie as linhas verdes para o fim da aba PARCELAS (colunas B ate E). O cronograma inteiro fica pronto: 1/60, 2/60, ... 60/60."),
            ("Pagar uma parcela", "Na aba LANCAR escolha o Compromisso e deixe o numero da parcela em branco: o sistema aponta sozinho a proxima em aberto (por exemplo 44/60). Copie a linha pronta para LANCAMENTOS. A parcela muda para PAGA e a Parcela_Atual do compromisso avanca."),
            ("Parcela atrasada", "Uma parcela vencida e sem pagamento realizado aparece como ATRASADA em PARCELAS e na aba ALERTAS, com os dias de atraso."),
            ("Compra parcelada no cartao", "Cadastre como compromisso com Entidade_Tipo = CARTAO. Cada parcela vira um lancamento COMPRA_CARTAO na data do vencimento. Nao lance o valor total de uma vez - isso duplicaria a despesa."),
        ]),
        ("5. INVESTIMENTOS", [
            ("Criar um investimento", "Aba CAD_Investimentos: Nome, Instituicao, Tipo (Fondo Mutuo, Ahorro a Plazo, ...), Base_Rendimento, Taxa e Prazo_Meses. Instituicoes e tipos sao livres - nada esta fixo."),
            ("Fazer um aporte", "Tipo_Operacao = APORTE, Origem = CONTA, Destino = INVESTIMENTO. Sai da conta e entra no investimento. NAO e despesa."),
            ("Fazer um resgate", "Use o Assistente de resgate na aba LANCAR. Ele gera dois lancamentos: um RENDIMENTO (receita financeira, se ainda nao foi lancado) e um RESGATE (movimento neutro). O principal nunca vira receita."),
            ("Resgate parcial ou total", "E o mesmo procedimento; o que muda e o valor. O saldo do investimento cai e o da conta sobe."),
            ("Rendimento projetado x realizado", "Rendimento REALIZADO e o que ja foi creditado. Rendimento PROJETADO vem de lancamentos com status PROJETADO e aparece sempre em coluna separada. As duas colunas nunca se somam nos relatorios de resultado."),
            ("Ahorro a Plazo", "Informe Data_Inicial, Prazo_Meses e Taxa. O sistema calcula a Data_Vencimento e a estimativa pela taxa. Lance o rendimento e o resgate como PROJETADOS na data do vencimento para que entrem no fluxo futuro."),
        ]),
        ("6. METAS", [
            ("Criar uma meta", "Aba CAD_Metas: Nome, Valor_Objetivo, Data_Prazo e a conta onde o dinheiro esta guardado."),
            ("Vincular a meta a uma conta", "O campo Conta_Vinculada_ID faz a reserva aparecer na coluna Reservado_Metas daquela conta. O dinheiro continua fisicamente na conta; o que muda e o Saldo_Disponivel."),
            ("Reservar dinheiro", "Assistente de meta na aba LANCAR: Tipo = RESERVA. Copie a linha para MOV_METAS (colunas B ate H)."),
            ("Retirar dinheiro da meta", "Tipo = RETIRADA e preencha Destino_Retirada (saldo disponivel, outra meta, investimento, despesa, outra conta). O valor nunca e simplesmente apagado - fica registrado."),
            ("Transferir entre metas", "Duas linhas: TRANSF_SAIDA na meta de origem (informando a meta de destino) e TRANSF_ENTRADA na meta de destino, com o mesmo valor."),
            ("Concluir uma meta", "Mude o Status para Concluida e informe a Data_Conclusao. Depois decida o que fazer com o valor reservado: manter (nao faca nada), liberar (LIBERACAO_CONCLUSAO), transferir (TRANSF_SAIDA) ou investir (retirada com destino Investimento + um APORTE em LANCAMENTOS)."),
            ("Cancelar uma meta", "Status = Cancelada e registre um movimento de saida com o destino escolhido. O historico da meta permanece inteiro em MOV_METAS."),
            ("Reserva no fluxo projetado", "Em CFG_Sistema, o parametro Considerar reservas de metas no fluxo projetado decide se o valor reservado das metas marcadas com SIM e descontado da coluna Saldo_Livre do fluxo de caixa."),
        ]),
        ("7. RECORRENCIAS", [
            ("Criar uma recorrencia", "Aba CAD_Recorrencias: descricao, tipo de operacao, valor, periodicidade, dia, data de inicio e quantas ocorrencias."),
            ("Gerar os lancamentos", "Aba GERADOR, bloco 2: escolha a recorrencia e copie as linhas verdes para LANCAMENTOS. As datas futuras ja saem como PROJETADO."),
            ("Quando o pagamento acontecer", "Troque o Status daquela linha de PROJETADO para REALIZADO. Nunca crie uma nova linha."),
            ("Recorrencia sem data final", "Coloque uma quantidade grande de ocorrencias (por exemplo 60) e gere quando precisar."),
        ]),
        ("8. CONSULTAR", [
            ("Fluxo de caixa", "Aba FLUXO_CAIXA: quatro visoes na mesma aba - mensal, diario, semanal e anual. Realizado e projetado ficam em colunas separadas."),
            ("Projecao", "Aba PROJECAO: mes a mes para frente, detalhando o que compoe cada saida (parcelas, faturas, aportes)."),
            ("Relatorio mensal", "Aba REL_Mensal: escolha o ano e o mes nas celulas amarelas."),
            ("Relatorio anual", "Aba REL_Anual: escolha o ano. Traz os 12 meses, totais, medias, maior e menor gasto e a evolucao de investimentos e dividas."),
            ("Patrimonio", "Aba PATRIMONIO: posicao atual e evolucao mes a mes."),
            ("Alertas e indicadores", "Abas ALERTAS e INDICADORES. Os limites de cada alerta ficam em CFG_Sistema."),
            ("Conferencia", "Aba TESTES: se o numero de TESTES COM FALHA for diferente de zero, algum calculo esta inconsistente - procure a linha vermelha."),
        ]),
        ("9. LIMITES DESTA VERSAO", [
            ("Sem macros", "O arquivo nao usa VBA, para abrir em qualquer Excel sem aviso de seguranca. Por isso copiar e colar blocos prontos substitui os botoes."),
            ("Quantidade de linhas", "As formulas cobrem 3.000 lancamentos, 1.200 parcelas, 600 movimentos de meta e 40 registros por cadastro. Para ampliar, use Formulas > Gerenciador de Nomes."),
            ("Ordem das linhas", "Nao classifique (sort) as abas de dados. Use os filtros do cabecalho, que nao alteram a ordem fisica."),
            ("Data de referencia", "Fica fixa em CFG_Sistema para que os testes sejam reproduziveis. Troque por =HOJE() se quiser que o sistema acompanhe o dia corrente."),
        ]),
    ]
    return _texto(L, "MANUAL", "MANUAL DE USO", "Leia a secao 1 uma vez; depois use as demais como consulta.", secoes, "0070C0")


def aba_arquitetura(L):
    secoes = [
        ("1. CAMADAS", [
            ("Configuracao", "CFG_Sistema (parametros), CFG_Listas (dominios), AUX (tabelas derivadas e listas dependentes)."),
            ("Cadastros (dimensoes)", "CAD_Instituicoes, CAD_Contas, CAD_Cartoes, CAD_Investimentos, CAD_Categorias, CAD_Subcategorias, CAD_Compromissos, CAD_Metas, CAD_Recorrencias, CAD_Bens."),
            ("Operacoes (fatos)", "LANCAMENTOS (livro unico), PARCELAS (cronograma), MOV_METAS (reservas logicas)."),
            ("Calculos", "Ficam nas proprias tabelas, em colunas cinzas: saldos em CAD_Contas, dividas em CAD_Cartoes, saldos e rendimentos em CAD_Investimentos, status em PARCELAS, reservas em CAD_Metas."),
            ("Relatorios", "DASHBOARD, FLUXO_CAIXA, PROJECAO, REL_Mensal, REL_Anual, PATRIMONIO, INDICADORES, ALERTAS."),
            ("Apoio", "LANCAR (entrada de dados), GERADOR (cronogramas e recorrencias), TESTES, MANUAL, ARQUITETURA."),
        ]),
        ("2. CHAVES E RELACIONAMENTOS", [
            ("Prefixos de ID", "INS (instituicao), CON (conta), CAR (cartao), INV (investimento), CAT (categoria), SUB (subcategoria), CMP (compromisso), PAR (parcela), MET (meta), REC (recorrencia), BEM (bem), LAN (lancamento), MOV (movimento de meta)."),
            ("Geracao do ID", 'Formula na coluna A de cada tabela: ="PREFIXO-"&TEXT(ROW()-4,"000000"). O usuario nunca digita um ID.'),
            ("LANCAMENTOS -> cadastros", "Origem_ID e Destino_ID apontam para CON/CAR/INV (a coluna Origem_Tipo diz qual). Categoria_ID -> CAT, Subcategoria_ID -> SUB, ID_Compromisso -> CMP, ID_Parcela -> PAR, ID_Meta -> MET, ID_Recorrencia -> REC, ID_Relacionado -> LAN."),
            ("PARCELAS -> CMP", "Cada parcela pertence a um compromisso; o pagamento e o lancamento que traz o ID_Parcela dela."),
            ("MOV_METAS -> MET", "Cada movimento pertence a uma meta; a conta afetada vem da meta (Conta_Vinculada_ID)."),
            ("Integridade referencial", "A coluna Validacao de LANCAMENTOS e de MOV_METAS verifica, linha a linha, se cada ID referenciado existe e se a subcategoria pertence a categoria informada."),
        ]),
        ("3. REGRAS FINANCEIRAS", [
            ("Classificacao economica", "Cada Tipo_Operacao tem uma classificacao fixa em CFG_Listas: RECEITA, RECEITA_FINANCEIRA, DESPESA ou NEUTRO. E dela que saem as colunas Vlr_Receita, Vlr_Despesa e Vlr_Rend_Fin."),
            ("Efeito no caixa", "Entrada_Caixa = destino e CONTA e origem nao e CONTA. Saida_Caixa = origem e CONTA e destino nao e CONTA. Transferencia entre contas proprias fica fora dos dois - e a garantia de que nao ha duplicidade."),
            ("Saldo de conta", "Saldo_Inicial + entradas realizadas - saidas realizadas."),
            ("Divida de cartao", "Divida_Inicial + compras (origem CARTAO) - pagamentos (destino CARTAO)."),
            ("Saldo de investimento", "Valor inicial + aportes + rendimentos - resgates. Principal = valor inicial + aportes - resgates (minimo zero); o restante do saldo e rendimento acumulado."),
            ("Saldo de meta", "Soma dos movimentos com sinal: RESERVA e TRANSF_ENTRADA somam; RETIRADA, TRANSF_SAIDA e liberacoes subtraem. Nao toca em saldo bancario."),
            ("Status de parcela", "PAGA se foi quitada antes do sistema ou tem lancamento realizado; senao ATRASADA se venceu, ABERTA se vence no mes corrente, PROJETADA se vence depois."),
        ]),
        ("4. PREPARACAO PARA A VERSAO 2", [
            ("Formato de tabela", "Cada aba de dados e uma tabela retangular com cabecalho unico na linha 4 e uma linha por registro - le direto em pandas, SQL ou qualquer ORM."),
            ("Separacao dado x calculo", "As colunas amarelas sao o dado bruto (o que iria para o banco). As cinzas sao derivadas e seriam substituidas por views ou por codigo de aplicacao."),
            ("Modelo relacional pronto", "As tabelas ja tem chave primaria (ID) e chaves estrangeiras explicitas. A migracao para PostgreSQL, SQLite ou Firestore e um mapeamento direto, tabela por tabela."),
            ("Dominios controlados", "CFG_Listas equivale a tabelas de dominio ou enums; a classificacao economica por tipo de operacao ja e uma tabela de mapeamento."),
            ("Regras isoladas", "Cada regra financeira esta em uma unica coluna calculada, o que facilita reescreve-la como funcao pura na V2."),
            ("O que a V2 poderia trazer", "Aplicativo ou site com telas de cadastro, importacao automatica de extratos, leitura de comprovantes, conciliacao bancaria, alertas por notificacao, orcamento por envelope, multimoeda (Gs./US$/R$) e analise inteligente por IA."),
        ]),
    ]
    return _texto(L, "ARQUITETURA", "ARQUITETURA DO SISTEMA",
                  "Como os dados estao organizados e por que. Base para a evolucao futura.",
                  secoes, "808080")


# ---------------------------------------------------------------------------
# MONTAGEM
# ---------------------------------------------------------------------------
def construir(caminho, limpo=False):
    global MODO_LIMPO
    MODO_LIMPO = limpo
    L = Livro()
    aba_dashboard(L)
    aba_lancar(L)
    aba_lancamentos(L)
    aba_parcelas(L)
    aba_mov_metas(L)
    abas_cadastros(L)
    aba_fluxo(L)
    aba_projecao(L)
    aba_rel_mensal(L)
    aba_rel_anual(L)
    aba_patrimonio(L)
    aba_indicadores(L)
    aba_alertas(L)
    aba_gerador(L)
    aba_testes(L)
    aba_manual(L)
    aba_arquitetura(L)
    aba_cfg_sistema(L)
    aba_cfg_listas(L)
    aba_aux(L)
    definir_nomes(L)
    ordem = ["DASHBOARD", "LANCAR", "LANCAMENTOS", "PARCELAS", "MOV_METAS",
             "CAD_Instituicoes", "CAD_Contas", "CAD_Cartoes", "CAD_Investimentos",
             "CAD_Categorias", "CAD_Subcategorias", "CAD_Compromissos", "CAD_Metas",
             "CAD_Recorrencias", "CAD_Bens", "FLUXO_CAIXA", "PROJECAO", "REL_Mensal",
             "REL_Anual", "PATRIMONIO", "INDICADORES", "ALERTAS", "GERADOR", "TESTES",
             "MANUAL", "ARQUITETURA", "CFG_Sistema", "CFG_Listas", "AUX"]
    L.wb._sheets.sort(key=lambda s: ordem.index(s.title))
    L.wb.active = 0
    L.salvar(caminho)
    return L
