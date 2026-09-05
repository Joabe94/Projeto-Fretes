#!/usr/bin/env python3
"""Recalcula um .xlsx com LibreOffice (perfil persistente) e reporta celulas de erro."""
import json, os, subprocess, sys, tempfile, time
from pathlib import Path

SKILL = next(Path("/root/.claude/skills/synced").glob("*/xlsx/scripts"), None)
if SKILL: sys.path.insert(0, str(SKILL))
try:
    from office.soffice import get_soffice_env
except Exception:
    def get_soffice_env():
        e = os.environ.copy(); e["SAL_USE_VCLPLUGIN"] = "svp"; return e

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

PROFILE = Path(tempfile.gettempdir()) / "lo_profile_persist"
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
ERRS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "Err:")


def ensure_profile():
    std = PROFILE / "user" / "basic" / "Standard"
    if not std.exists():
        subprocess.run(["soffice", "--headless", "--terminate_after_init",
                        f"-env:UserInstallation={PROFILE.as_uri()}"],
                       env=get_soffice_env(), capture_output=True, timeout=180)
    std.mkdir(parents=True, exist_ok=True)
    (std / "Module1.xba").write_text(MACRO)


def recalc(path, timeout=900):
    p = Path(path).absolute()
    ensure_profile()
    t0 = time.time()
    cmd = ["soffice", "--headless", "--norestore", "--invisible",
           f"-env:UserInstallation={PROFILE.as_uri()}",
           f'macro:///Standard.Module1.RecalculateAndSave("{p}")', str(p)]
    try:
        subprocess.run(cmd, env=get_soffice_env(), capture_output=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        return {"error": f"timeout apos {timeout}s"}
    dur = round(time.time() - t0, 1)
    wb = openpyxl.load_workbook(p, data_only=True)
    wbf = openpyxl.load_workbook(p, data_only=False)
    nform, errors = 0, {}
    for sh in wbf.sheetnames:
        wsf, wsv = wbf[sh], wb[sh]
        for row in wsf.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula): v = v.text
                if isinstance(v, str) and v.startswith("="):
                    nform += 1
                    got = wsv[c.coordinate].value
                    if isinstance(got, str) and any(got.startswith(e) for e in ERRS):
                        errors.setdefault(got, []).append(f"{sh}!{c.coordinate}")
    total = sum(len(v) for v in errors.values())
    return {"status": "errors_found" if total else "success", "seconds": dur,
            "total_formulas": nform, "total_errors": total,
            "error_summary": {k: v[:40] for k, v in errors.items()}}


if __name__ == "__main__":
    f = sys.argv[1]
    t = int(sys.argv[2]) if len(sys.argv) > 2 else 900
    print(json.dumps(recalc(f, t), indent=2, ensure_ascii=False))
