#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Auditoria independente do arquivo Excel gerado.

Le os valores JA RECALCULADOS pelo LibreOffice e compara com um calculo feito
em Python a partir dos dados brutos (gerador/motor.py). Qualquer divergencia e
uma falha real do sistema, nao um erro de formula.
"""
from __future__ import annotations

import datetime as dt
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "gerador"))

import openpyxl

from comum import (ANO_BASE, DATA_INICIO, HOJE, LINHA_DADOS as DAT,
                   LINHA_HDR as HDR, gs)
import dados as D
import motor as M

TOL = 0.5
FALHAS: list[str] = []
OKS: list[str] = []


def cmp(nome, esperado, obtido, tol=TOL, fmt=gs):
    try:
        dif = abs(float(esperado or 0) - float(obtido or 0))
    except (TypeError, ValueError):
        dif = None
    if dif is None or dif > tol:
        FALHAS.append(f"{nome}: esperado {fmt(esperado)} | obtido {fmt(obtido)}")
        return False
    OKS.append(nome)
    return True


def cmp_txt(nome, esperado, obtido):
    if str(esperado) != str(obtido):
        FALHAS.append(f"{nome}: esperado {esperado!r} | obtido {obtido!r}")
        return False
    OKS.append(nome)
    return True


def linhas(ws, col, n, ini=DAT):
    return [ws.cell(ini + i, col).value for i in range(n)]


def auditar(caminho):
    wbv = openpyxl.load_workbook(caminho, data_only=True)
    wbf = openpyxl.load_workbook(caminho, data_only=False)
    L = D.LANCAMENTOS

    print("=" * 78)
    print("AUDITORIA INDEPENDENTE - SISTEMA FINANCEIRO PESSOAL V1")
    print("=" * 78)

    # ---------------------------------------------------------------- 1. TESTES
    ts = wbv["TESTES"]
    falhas_excel, exec_excel = None, None
    for r in range(1, 12):
        if ts.cell(r, 1).value == "TESTES COM FALHA":
            falhas_excel, exec_excel = ts.cell(r, 2).value, ts.cell(r, 5).value
    print(f"\n[1] Bateria interna da aba TESTES: {exec_excel} testes executados, "
          f"{falhas_excel} com falha")
    if falhas_excel:
        for r in range(1, ts.max_row + 1):
            if ts.cell(r, 7).value == "FALHA":
                print(f"    FALHA -> {ts.cell(r, 1).value} | {ts.cell(r, 3).value} | "
                      f"esperado {gs(ts.cell(r, 4).value)} obtido {gs(ts.cell(r, 5).value)}")
    cmp("Bateria interna sem falhas", 0, falhas_excel or 0, fmt=str)
    if not exec_excel:
        FALHAS.append("Bateria interna nao executou nenhum teste")

    # ---------------------------------------------------------------- 2. CONTAS
    print("\n[2] Saldos de contas (Excel x Python)")
    ws = wbv["CAD_Contas"]
    esp = M.saldos_contas(L, ("REALIZADO",))
    esp_proj = M.saldos_contas(L, ("REALIZADO", "PROJETADO"))
    reserv = M.reservado_por_conta(D.MOV_METAS, D.METAS)
    for i, c in enumerate(D.CONTAS):
        rr = DAT + i
        cid, nome = c[0], c[1]
        cmp_txt(f"ID da conta {nome}", cid, ws.cell(rr, 1).value)
        ok = cmp(f"Saldo calculado {nome}", esp[cid], ws.cell(rr, 10).value)
        cmp(f"Reservado em metas {nome}", reserv.get(cid, 0), ws.cell(rr, 11).value)
        cmp(f"Saldo disponivel {nome}", esp[cid] - reserv.get(cid, 0), ws.cell(rr, 12).value)
        cmp(f"Saldo projetado {nome}", esp_proj[cid], ws.cell(rr, 15).value)
        print(f"    {nome:32s} {gs(ws.cell(rr, 10).value):>20} "
              f"reservado {gs(ws.cell(rr, 11).value):>18} {'OK' if ok else 'FALHA'}")

    # ---------------------------------------------------------------- 3. CARTOES
    print("\n[3] Dividas de cartao")
    ws = wbv["CAD_Cartoes"]
    espk = M.dividas_cartoes(L, ("REALIZADO",))
    for i, c in enumerate(D.CARTOES):
        rr = DAT + i
        kid, nome, limite = c[0], c[1], c[3]
        ok = cmp(f"Divida atual {nome}", espk[kid], ws.cell(rr, 12).value)
        cmp(f"Limite disponivel {nome}", limite - espk[kid], ws.cell(rr, 14).value)
        print(f"    {nome:32s} divida {gs(ws.cell(rr, 12).value):>16} "
              f"uso {(ws.cell(rr, 13).value or 0) * 100:5.1f}%  fatura atual "
              f"{gs(ws.cell(rr, 18).value):>16} venc {ws.cell(rr, 17).value} "
              f"{'OK' if ok else 'FALHA'}")

    # ------------------------------------------------------------ 4. INVESTIMENTOS
    print("\n[4] Investimentos")
    ws = wbv["CAD_Investimentos"]
    espi = M.saldos_investimentos(L, ("REALIZADO",))
    proj = M.saldos_investimentos(L, ("PROJETADO",))
    for i, inv in enumerate(D.INVESTIMENTOS):
        rr = DAT + i
        iid, nome = inv[0], inv[1]
        e = espi[iid]
        ok = cmp(f"Saldo do investimento {nome}", e["saldo"], ws.cell(rr, 15).value)
        cmp(f"Principal {nome}", e["principal"], ws.cell(rr, 16).value)
        cmp(f"Rendimento no saldo {nome}", e["rend_no_saldo"], ws.cell(rr, 17).value)
        cmp(f"Rendimento realizado {nome}", e["rendimento"], ws.cell(rr, 14).value)
        cmp(f"Rendimento projetado {nome}", proj[iid]["rendimento"], ws.cell(rr, 19).value)
        cmp(f"Principal + rendimento = saldo {nome}",
            (ws.cell(rr, 16).value or 0) + (ws.cell(rr, 17).value or 0),
            ws.cell(rr, 15).value)
        print(f"    {nome:32s} saldo {gs(ws.cell(rr, 15).value):>18} "
              f"principal {gs(ws.cell(rr, 16).value):>18} rend.real {gs(ws.cell(rr, 14).value):>14} "
              f"rend.proj {gs(ws.cell(rr, 19).value):>14} venc {ws.cell(rr, 11).value} "
              f"{'OK' if ok else 'FALHA'}")

    # ---------------------------------------------------------------- 5. METAS
    print("\n[5] Metas")
    ws = wbv["CAD_Metas"]
    espm = M.saldos_metas(D.MOV_METAS)
    for i, m in enumerate(D.METAS):
        rr = DAT + i
        ok = cmp(f"Valor reservado na meta {m['nome']}", espm[m["id"]], ws.cell(rr, 8).value)
        cmp(f"Faltante na meta {m['nome']}", max(0, m["objetivo"] - espm[m["id"]]),
            ws.cell(rr, 10).value)
        print(f"    {m['nome']:26s} objetivo {gs(m['objetivo']):>18} acumulado "
              f"{gs(ws.cell(rr, 8).value):>18} {(ws.cell(rr, 9).value or 0) * 100:5.1f}% "
              f"{str(ws.cell(rr, 15).value):18s} {'OK' if ok else 'FALHA'}")

    # ---------------------------------------------------------------- 6. PARCELAS
    print("\n[6] Parcelas")
    ws = wbv["PARCELAS"]
    espp = M.status_parcelas(D.PARCELAS, L)
    cont_x, cont_p = {}, {}
    div = 0
    for i, p in enumerate(D.PARCELAS):
        rr = DAT + i
        st_x = ws.cell(rr, 15).value
        cont_x[st_x] = cont_x.get(st_x, 0) + 1
        cont_p[espp[p["id"]]] = cont_p.get(espp[p["id"]], 0) + 1
        if st_x != espp[p["id"]]:
            div += 1
            if div <= 5:
                FALHAS.append(f"Status da parcela {p['id']} ({p['rotulo']}): "
                              f"esperado {espp[p['id']]} obtido {st_x}")
    if div == 0:
        OKS.append("Status de todas as parcelas")
    print(f"    Excel : {cont_x}")
    print(f"    Python: {cont_p}")
    # compromisso do terreno
    wsc = wbv["CAD_Compromissos"]
    for i, c in enumerate(D.COMPROMISSOS):
        rr = DAT + i
        pagas = sum(1 for p in D.PARCELAS if p["cmp"] == c["id"] and espp[p["id"]] == "PAGA")
        cmp(f"Parcelas pagas de {c['nome']}", pagas, wsc.cell(rr, 16).value, fmt=str)
        cmp(f"Valor total de {c['nome']}", c["qtd"] * c["valor"], wsc.cell(rr, 9).value)
        cmp(f"Pago + devedor = total {c['nome']}",
            wsc.cell(rr, 9).value, (wsc.cell(rr, 21).value or 0) + (wsc.cell(rr, 22).value or 0))
        print(f"    {c['nome']:34s} parcela atual {str(wsc.cell(rr, 23).value):8s} "
              f"restantes {str(wsc.cell(rr, 20).value):>3} devedor {gs(wsc.cell(rr, 22).value):>18} "
              f"prox.venc {wsc.cell(rr, 24).value}")

    # -------------------------------------------------------------- 7. PATRIMONIO
    print("\n[7] Patrimonio")
    ws = wbv["PATRIMONIO"]
    pat = {}
    for r in range(1, 40):
        k = ws.cell(r, 1).value
        if isinstance(k, str):
            pat[k] = ws.cell(r, 2).value
    espp2 = M.patrimonio(L, D.MOV_METAS, D.PARCELAS, HOJE)
    cmp("Ativo: contas", espp2["contas"], pat.get("Contas bancarias e dinheiro fisico"))
    cmp("Ativo: investimentos", espp2["investimentos"], pat.get("Investimentos"))
    cmp("Passivo: cartoes", espp2["cartoes"], pat.get("Dividas de cartao de credito"))
    cmp("Passivo: compromissos", espp2["compromissos"], pat.get("Saldo devedor de compromissos"))
    cmp("Patrimonio liquido financeiro", espp2["pl_financeiro"],
        pat.get("Patrimonio liquido FINANCEIRO"))
    cmp("Patrimonio liquido total", espp2["pl_total"], pat.get("Patrimonio liquido TOTAL"))
    cmp("Ativos - passivos = PL total",
        (pat.get("TOTAL DE ATIVOS") or 0) - (pat.get("TOTAL DE PASSIVOS") or 0),
        pat.get("Patrimonio liquido TOTAL"))
    for k in ["TOTAL DE ATIVOS", "TOTAL DE PASSIVOS", "Patrimonio liquido FINANCEIRO",
              "Patrimonio liquido TOTAL"]:
        print(f"    {k:42s} {gs(pat.get(k)):>20}")

    # ------------------------------------------------------------- 8. FLUXO MENSAL
    print("\n[8] Fluxo de caixa mensal (primeiros 12 meses)")
    ws = wbv["FLUXO_CAIXA"]
    r_ini = None
    for r in range(1, 40):
        if ws.cell(r, 1).value == "Competencia":
            r_ini = r + 1
            break
    espf = M.fluxo_mensal(L, dt.date(2026, 1, 1), 30)
    for i in range(30):
        rr = r_ini + i
        e = espf[i]
        cmp(f"Fluxo {e['ano']}-{e['mes']:02d} entradas realizadas", e["ent_real"], ws.cell(rr, 5).value)
        cmp(f"Fluxo {e['ano']}-{e['mes']:02d} saidas realizadas", e["sai_real"], ws.cell(rr, 6).value)
        cmp(f"Fluxo {e['ano']}-{e['mes']:02d} saldo final", e["saldo_fim"], ws.cell(rr, 11).value)
        if i < 12:
            print(f"    {e['ano']}-{e['mes']:02d}  inicial {gs(ws.cell(rr, 4).value):>18}"
                  f"  ent {gs(ws.cell(rr, 5).value):>16}  sai {gs(ws.cell(rr, 6).value):>16}"
                  f"  final {gs(ws.cell(rr, 11).value):>18}")

    # ---------------------------------------------------------- 9. RELATORIO MENSAL
    print("\n[9] Relatorio mensal (ano/mes selecionados na aba)")
    ws = wbv["REL_Mensal"]
    ano = ws.cell(6, 2).value or ANO_BASE
    mes = ws.cell(6, 4).value or HOJE.month
    ini = dt.date(int(ano), int(mes), 1)
    fim = (dt.date(int(ano) + (mes == 12), int(mes) % 12 + 1, 1) - dt.timedelta(days=1))
    er = M.resultado_periodo(L, ini, fim, "REALIZADO")
    ep = M.resultado_periodo(L, ini, fim, "PROJETADO")
    mapa = {}
    for r in range(1, 60):
        k = ws.cell(r, 1).value
        if isinstance(k, str) and k:
            mapa[k] = (ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value)
    print(f"    Periodo: {ini} a {fim}")
    for nome, chave, campo in [("Receitas", "Receitas operacionais", "receita"),
                               ("Rendimentos", "Receitas financeiras (rendimentos)", "rec_fin"),
                               ("Despesas", "Despesas", "despesa"),
                               ("Entradas de caixa", "Entradas de caixa", "entrada"),
                               ("Saidas de caixa", "Saidas de caixa", "saida")]:
        v = mapa.get(chave)
        if v is None:
            FALHAS.append(f"Linha '{chave}' nao encontrada em REL_Mensal")
            continue
        cmp(f"REL_Mensal {nome} realizado", er[campo], v[0])
        cmp(f"REL_Mensal {nome} projetado", ep[campo], v[1])
        print(f"    {nome:22s} realizado {gs(v[0]):>18}  projetado {gs(v[1]):>18}")

    # ------------------------------------------------------------ 10. FORMATO/MOEDA
    print("\n[10] Moeda, validacoes e estrutura")
    wsf = wbf["LANCAMENTOS"]
    fmt_valor = wsf.cell(DAT, 6).number_format
    cmp("Coluna Valor formatada em Guarani", True, "Gs." in fmt_valor, fmt=str)
    print(f"    Formato da coluna Valor: {fmt_valor}")
    total_gs = total_cel = 0
    for aba in wbf.sheetnames:
        w = wbf[aba]
        for row in w.iter_rows():
            for c in row:
                if c.number_format and "Gs." in str(c.number_format):
                    total_gs += 1
                if isinstance(c.number_format, str) and ("R$" in c.number_format
                                                         or "$#" in c.number_format
                                                         or "USD" in c.number_format):
                    total_cel += 1
    print(f"    Celulas formatadas em Gs.: {total_gs}")
    cmp("Nenhuma celula em outra moeda", 0, total_cel, fmt=str)
    nomes = list(wbf.defined_names.keys())
    print(f"    Intervalos nomeados: {len(nomes)}")
    faltando = [n for n in ["L_Valor", "L_Status", "P_STATUS", "MM_VALSINAL", "C_SALDO",
                            "K_DIVIDA", "I_SALDO", "M_RESERV", "CFG_HOJE", "TS_FALHAS"]
                if n not in nomes]
    cmp("Intervalos nomeados essenciais presentes", 0, len(faltando), fmt=str)
    if faltando:
        FALHAS.append(f"Intervalos nomeados ausentes: {faltando}")
    dvs = sum(len(wbf[a].data_validations.dataValidation) for a in wbf.sheetnames)
    print(f"    Regras de validacao de dados: {dvs}")
    cmp("Existem validacoes de dados", True, dvs > 20, fmt=str)
    print(f"    Abas: {len(wbf.sheetnames)}")

    # --------------------------------------------------------- 11. NAO DUPLICIDADE
    print("\n[11] Verificacoes de nao duplicidade")
    wl = wbv["LANCAMENTOS"]
    n = 0
    soma_rec_transf = soma_desp_transf = soma_desp_aporte = soma_rec_resgate = 0
    soma_desp_pagcart = 0
    ids = []
    for i in range(1200):
        rr = DAT + i
        if not wl.cell(rr, 2).value:
            continue
        n += 1
        ids.append(wl.cell(rr, 1).value)
        tipo = wl.cell(rr, 4).value
        vrec = wl.cell(rr, 30).value or 0
        vdesp = wl.cell(rr, 31).value or 0
        if tipo == "TRANSFERENCIA":
            soma_rec_transf += vrec
            soma_desp_transf += vdesp
        if tipo == "APORTE":
            soma_desp_aporte += vdesp
        if tipo == "RESGATE":
            soma_rec_resgate += vrec
        if tipo == "PAGAMENTO_CARTAO":
            soma_desp_pagcart += vdesp
    cmp("Lancamentos lidos = lancamentos gerados", len(L), n, fmt=str)
    cmp("IDs de lancamento unicos", len(ids), len(set(ids)), fmt=str)
    cmp("Transferencia nao vira receita", 0, soma_rec_transf)
    cmp("Transferencia nao vira despesa", 0, soma_desp_transf)
    cmp("Aporte nao vira despesa", 0, soma_desp_aporte)
    cmp("Resgate nao vira receita", 0, soma_rec_resgate)
    cmp("Pagamento de cartao nao vira despesa", 0, soma_desp_pagcart)
    erros_val = sum(1 for i in range(1200)
                    if str(wl.cell(DAT + i, 36).value or "").startswith("ERRO"))
    cmp("Lancamentos com erro de validacao", 0, erros_val, fmt=str)
    wm = wbv["MOV_METAS"]
    erros_mov = sum(1 for i in range(300)
                    if str(wm.cell(DAT + i, 13).value or "").startswith("ERRO"))
    cmp("Movimentos de meta com erro de validacao", 0, erros_mov, fmt=str)
    print(f"    Lancamentos: {n}  IDs unicos: {len(set(ids))}  erros de validacao: {erros_val}")

    # --------------------------------------------------------------- 12. ALERTAS
    print("\n[12] Alertas ativos")
    wa = wbv["ALERTAS"]
    total_al = None
    for r in range(1, 12):
        if wa.cell(r, 1).value == "TOTAL DE ALERTAS ATIVOS":
            total_al = wa.cell(r, 2).value
    print(f"    Total de alertas ativos: {total_al}")
    ativos = []
    for r in range(1, wa.max_row + 1):
        for c in range(4, 8):
            if wa.cell(r, c).value == "ALERTA":
                ativos.append(str(wa.cell(r, 1).value))
                break
    for a in ativos:
        print(f"    ALERTA -> {a}")
    cmp("Contador de alertas confere com as linhas em alerta", len(ativos), total_al or 0, fmt=str)

    # ------------------------------------------------- 13. RELATORIOS PREENCHIDOS
    print("\n[13] Abas de relatorio, projecao e indicadores")
    dash = wbv["DASHBOARD"]
    kpis = [(r, c) for r in range(1, 45) for c in (1, 3, 5, 7)
            if isinstance(dash.cell(r, c).value, str)
            and dash.cell(r, c).value.isupper() and len(dash.cell(r, c).value) > 8]
    vazios = [dash.cell(r, c).value for r, c in kpis
              if dash.cell(r + 1, c).value is None]
    cmp("KPIs do dashboard preenchidos", 0, len(vazios), fmt=str)
    if vazios:
        FALHAS.append(f"KPIs vazios: {vazios[:6]}")
    print(f"    KPIs do dashboard: {len(kpis)} cartoes, {len(vazios)} vazios")
    cmp("Dashboard: saldo total", sum(M.saldos_contas(L, ("REALIZADO",)).values()),
        dash.cell(kpis[0][0] + 1, 1).value)

    wp = wbv["PROJECAO"]
    r_pj = next(r for r in range(1, 20) if wp.cell(r, 1).value == "Competencia") + 1
    saldos_pj = [wp.cell(r_pj + i, 15).value for i in range(24)]
    cmp("Projecao com 24 meses preenchidos", 24, sum(1 for x in saldos_pj if x is not None), fmt=str)
    d_pj = dt.date(HOJE.year, HOJE.month, 1)
    for i in range(24):
        fimm = (dt.date(d_pj.year + (d_pj.month == 12), d_pj.month % 12 + 1, 1)
                - dt.timedelta(days=1))
        esp_saldo = sum(M.saldos_contas(L, ("REALIZADO", "PROJETADO"), fimm).values())
        cmp(f"Projecao saldo {d_pj.year}-{d_pj.month:02d}", esp_saldo, saldos_pj[i])
        d_pj = dt.date(d_pj.year + (d_pj.month == 12), d_pj.month % 12 + 1, 1)
    print(f"    Projecao: {saldos_pj[0] and gs(saldos_pj[0])} -> {gs(saldos_pj[-1])} "
          f"(menor {gs(min(x for x in saldos_pj if x is not None))})")

    wi = wbv["INDICADORES"]
    ind = {}
    for r in range(1, 60):
        k = wi.cell(r, 1).value
        if isinstance(k, str) and k and wi.cell(r, 2).value is not None:
            ind[k] = wi.cell(r, 2).value
    cmp("Indicadores preenchidos", True, len(ind) >= 25, fmt=str)
    for k in ["Categoria de maior gasto no ano (REL_Anual)", "Taxa de poupanca dos ultimos 6 meses",
              "Percentual da renda comprometida (12 meses)", "Tendencia do patrimonio",
              "Menor saldo projetado no horizonte"]:
        v = ind.get(k)
        print(f"    {k[:52]:54s} {v if isinstance(v, str) else (f'{v*100:.1f}%' if isinstance(v, float) and abs(v) < 10 else gs(v))}")
        if v is None:
            FALHAS.append(f"Indicador ausente: {k}")

    wa2 = wbv["REL_Anual"]
    r_ra = next(r for r in range(1, 20) if wa2.cell(r, 1).value == "Mes") + 1
    tot_rec = sum(wa2.cell(r_ra + i, 4).value or 0 for i in range(12))
    tot_desp = sum(wa2.cell(r_ra + i, 6).value or 0 for i in range(12))
    er2 = M.resultado_periodo(L, dt.date(ANO_BASE, 1, 1), dt.date(ANO_BASE, 12, 31), "REALIZADO")
    ep2 = M.resultado_periodo(L, dt.date(ANO_BASE, 1, 1), dt.date(ANO_BASE, 12, 31), "PROJETADO")
    cmp("REL_Anual receitas do ano", er2["receita"] + ep2["receita"], tot_rec)
    cmp("REL_Anual despesas do ano", er2["despesa"] + ep2["despesa"], tot_desp)
    print(f"    REL_Anual {ANO_BASE}: receitas {gs(tot_rec)}  despesas {gs(tot_desp)}  "
          f"resultado {gs(tot_rec - tot_desp)}")

    wf = wbv["FLUXO_CAIXA"]
    blocos_fx = {}
    for r in range(1, wf.max_row + 1):
        v0 = wf.cell(r, 1).value
        if v0 in ("Data", "Semana_Inicio", "Ano"):
            blocos_fx[v0] = r + 1
    for nome, col in [("Data", 7), ("Semana_Inicio", 6), ("Ano", 7)]:
        r0 = blocos_fx.get(nome)
        if r0 is None:
            FALHAS.append(f"Bloco de fluxo '{nome}' nao encontrado")
            continue
        preenchidos = sum(1 for i in range(6) if wf.cell(r0 + i, col).value is not None)
        cmp(f"Bloco de fluxo {nome} preenchido", 6, preenchidos, fmt=str)
    print(f"    Fluxo de caixa: blocos mensal, diario, semanal e anual preenchidos")

    wpt = wbv["PATRIMONIO"]
    r_pt = next(r for r in range(1, 40) if wpt.cell(r, 1).value == "Competencia") + 1
    pl = [wpt.cell(r_pt + i, 9).value for i in range(30)]
    cmp("Evolucao patrimonial com 30 meses", 30, sum(1 for x in pl if x is not None), fmt=str)
    espev = M.evolucao_patrimonial(L, D.PARCELAS, dt.date(2026, 1, 1), 30)
    for i, e in enumerate(espev):
        rr = r_pt + i
        rot = f"{e['ano']}-{e['mes']:02d}"
        cmp(f"Evolucao {rot} contas", e["contas"], wpt.cell(rr, 3).value)
        cmp(f"Evolucao {rot} investimentos", e["investimentos"], wpt.cell(rr, 4).value)
        cmp(f"Evolucao {rot} cartoes", e["cartoes"], wpt.cell(rr, 6).value)
        cmp(f"Evolucao {rot} compromissos", e["compromissos"], wpt.cell(rr, 7).value)
        cmp(f"Evolucao {rot} PL financeiro", e["pl_fin"], wpt.cell(rr, 8).value)
        cmp(f"Evolucao {rot} PL total", e["pl_total"], wpt.cell(rr, 9).value)
        cmp(f"Identidade do PL em {rot}",
            (wpt.cell(rr, 3).value or 0) + (wpt.cell(rr, 4).value or 0)
            + (wpt.cell(rr, 5).value or 0) - (wpt.cell(rr, 6).value or 0)
            - (wpt.cell(rr, 7).value or 0), wpt.cell(rr, 9).value)
    print(f"    Patrimonio: {gs(pl[0])} (jan/2026) -> {gs(pl[-1])} (fim da grade)")

    # ------------------------------------------------------------- 14. TELA LANCAR
    print("\n[14] Tela de lancamento rapido")
    wln = wbv["LANCAR"]
    sit = None
    for r in range(1, 40):
        if wln.cell(r, 1).value == "Situacao do preenchimento":
            sit = wln.cell(r, 2).value
    cmp_txt("LANCAR: exemplo pronto para colar", "PRONTO PARA COLAR", sit)
    r_pronta = None
    for r in range(1, 45):
        if wln.cell(r, 1).value == "Data" and wln.cell(r, 2).value == "Descricao":
            r_pronta = r + 1
            break
    linha = [wln.cell(r_pronta, c).value for c in range(1, 19)] if r_pronta else []
    print(f"    Linha pronta: {linha[1]} | {linha[2]} | {gs(linha[4])} | origem {linha[6]} | "
          f"cat {linha[9]} | sub {linha[10]}")
    cmp("LANCAR resolveu a conta de origem", True, str(linha[6]).startswith("CON-"), fmt=str)
    cmp("LANCAR resolveu a categoria", True, str(linha[9]).startswith("CAT-"), fmt=str)
    cmp("LANCAR resolveu a subcategoria", True, str(linha[10]).startswith("SUB-"), fmt=str)

    # -------------------------------------- 14b. COBERTURA DAS LISTAS SUSPENSAS
    print("\n[14b] Listas suspensas nas colunas de dominio")
    from openpyxl.utils import get_column_letter as _CL
    esperadas = {
        "CAD_Instituicoes": ["Tipo"],
        "CAD_Contas": ["Instituicao", "Tipo", "Status"],
        "CAD_Cartoes": ["Instituicao", "Dia_Fechamento", "Dia_Vencimento", "Status"],
        "CAD_Investimentos": ["Instituicao", "Tipo", "Base_Rendimento", "Status"],
        "CAD_Categorias": ["Tipo_Padrao", "Status"],
        "CAD_Subcategorias": ["Categoria", "Status"],
        "CAD_Compromissos": ["Categoria", "Subcategoria", "Dia_Vencimento",
                             "Periodicidade", "Entidade_Tipo", "Entidade_ID", "Status"],
        "CAD_Metas": ["Conta_Vinculada_ID", "Investimento_Vinculado_ID",
                      "Considerar_No_Fluxo", "Status"],
        "CAD_Recorrencias": ["Tipo_Operacao", "Periodicidade", "Dia", "Categoria_ID",
                             "Subcategoria_ID", "Origem_Tipo", "Origem_ID",
                             "Destino_Tipo", "Destino_ID", "Forma_Pagamento", "Status"],
        "CAD_Bens": ["Tipo", "ID_Compromisso_Vinculado", "Status"],
        "LANCAMENTOS": ["Tipo_Operacao", "Status", "Origem_Tipo", "Origem_ID",
                        "Destino_Tipo", "Destino_ID", "Categoria_ID", "Subcategoria_ID",
                        "Forma_Pagamento", "ID_Compromisso", "ID_Parcela", "ID_Meta",
                        "ID_Recorrencia"],
        "PARCELAS": ["ID_Compromisso"],
        "MOV_METAS": ["ID_Meta", "Tipo_Movimento", "Destino_Retirada", "ID_Meta_Destino"],
    }
    sem_lista = []
    total_dv = 0
    for aba, colunas in esperadas.items():
        w = wbf[aba]
        cobertas = set()
        for dvv in w.data_validations.dataValidation:
            total_dv += 1
            for rng in dvv.sqref.ranges:
                for cc in range(rng.min_col, rng.max_col + 1):
                    cobertas.add(cc)
        for h in colunas:
            col = next((c for c in range(1, w.max_column + 1)
                        if w.cell(HDR, c).value == h), None)
            if col is None:
                sem_lista.append(f"{aba}!{h} (coluna inexistente)")
            elif col not in cobertas:
                sem_lista.append(f"{aba}!{_CL(col)} {h}")
    cmp("Toda coluna de dominio tem lista suspensa", 0, len(sem_lista), fmt=str)
    if sem_lista:
        FALHAS.append(f"Colunas sem lista: {sem_lista}")
    print(f"    {sum(len(v) for v in esperadas.values())} colunas de dominio verificadas, "
          f"{len(sem_lista)} sem lista, {total_dv} regras de validacao")

    # ------------------------------------------------ 15. INTEGRIDADE ESTRUTURAL
    print("\n[15] Integridade estrutural")
    abas = set(wbf.sheetnames)
    quebrados = []
    for nome_dn, dn in wbf.defined_names.items():
        ref = str(dn.value)
        if "#REF" in ref:
            quebrados.append(nome_dn)
            continue
        aba_ref = ref.split("!")[0].strip("'")
        if aba_ref not in abas:
            quebrados.append(f"{nome_dn} -> {aba_ref}")
    cmp("Intervalos nomeados apontam para abas existentes", 0, len(quebrados), fmt=str)
    if quebrados:
        FALHAS.append(f"Nomes quebrados: {quebrados[:8]}")
    ref_quebradas = 0
    for aba in wbf.sheetnames:
        for row in wbf[aba].iter_rows():
            for c in row:
                if isinstance(c.value, str) and "#REF!" in c.value:
                    ref_quebradas += 1
    cmp("Nenhuma referencia #REF! nas formulas", 0, ref_quebradas, fmt=str)
    print(f"    {len(wbf.defined_names)} nomes definidos, {len(abas)} abas, "
          f"{ref_quebradas} referencias quebradas")

    # --------------------------------------------------------------- RESULTADO
    print("\n" + "=" * 78)
    print(f"RESULTADO: {len(OKS)} verificacoes OK, {len(FALHAS)} falhas")
    print("=" * 78)
    for f in FALHAS[:60]:
        print("  FALHA:", f)
    return len(FALHAS)


if __name__ == "__main__":
    arq = sys.argv[1] if len(sys.argv) > 1 else "SISTEMA_FINANCEIRO_PESSOAL_V1_EXEMPLO.xlsx"
    sys.exit(1 if auditar(arq) else 0)
