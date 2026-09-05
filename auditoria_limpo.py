#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Confere o arquivo LIMPO: zerado, sem erro e pronto para receber dados reais."""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "gerador"))

import openpyxl

from comum import LINHA_DADOS as DAT

ARQ = sys.argv[1] if len(sys.argv) > 1 else "SISTEMA_FINANCEIRO_PESSOAL_V1_LIMPO.xlsx"
FALHAS, OKS = [], []


def ok(nome, cond, detalhe=""):
    (OKS if cond else FALHAS).append(f"{nome} {detalhe}".strip())
    print(f"  [{'OK   ' if cond else 'FALHA'}] {nome} {detalhe}".rstrip())


def main():
    v = openpyxl.load_workbook(ARQ, data_only=True)
    print("=" * 74)
    print("CONFERENCIA DO ARQUIVO LIMPO")
    print("=" * 74)

    vazias = ["LANCAMENTOS", "PARCELAS", "MOV_METAS", "CAD_Contas", "CAD_Cartoes",
              "CAD_Investimentos", "CAD_Compromissos", "CAD_Metas", "CAD_Recorrencias",
              "CAD_Instituicoes", "CAD_Bens"]
    for aba in vazias:
        ws = v[aba]
        n = sum(1 for i in range(200) if ws.cell(DAT + i, 2).value not in (None, ""))
        ok(f"{aba} sem registros", n == 0, f"({n} linhas preenchidas)")

    cat = v["CAD_Categorias"]
    n_cat = sum(1 for i in range(40) if cat.cell(DAT + i, 2).value)
    ok("Categorias padrao presentes", n_cat == 13, f"({n_cat})")
    sub = v["CAD_Subcategorias"]
    n_sub = sum(1 for i in range(120) if sub.cell(DAT + i, 4).value)
    ok("Subcategorias padrao presentes", n_sub == 54, f"({n_sub})")
    ok("ID da primeira categoria gerado", cat.cell(DAT, 1).value == "CAT-000001",
       f"({cat.cell(DAT, 1).value})")
    aux = v["AUX"]
    lista = [aux.cell(DAT + i, 11).value for i in range(6)]
    ok("Lista dependente de Alimentacao funcionando", "Supermercado" in lista, str(lista[:5]))

    ts = v["TESTES"]
    falhas = exec_ = None
    for r in range(1, 12):
        if ts.cell(r, 1).value == "TESTES COM FALHA":
            falhas, exec_ = ts.cell(r, 2).value, ts.cell(r, 5).value
    ok("Bateria interna sem falhas", falhas == 0, f"({exec_} testes executados)")

    d = v["DASHBOARD"]
    kpis = {}
    for r in range(1, 45):
        for c in (1, 3, 5, 7):
            t = d.cell(r, c).value
            if isinstance(t, str) and t.isupper() and len(t) > 8:
                kpis[t] = d.cell(r + 1, c).value
    zerar = ["SALDO TOTAL DAS CONTAS", "RESERVADO PARA METAS", "TOTAL INVESTIDO",
             "DIVIDA DE CARTOES", "SALDO DEVEDOR DE COMPROMISSOS", "METAS ATIVAS",
             "LANCAMENTOS REGISTRADOS", "LANCAMENTOS COM ERRO", "TESTES COM FALHA",
             "PATRIMONIO LIQUIDO TOTAL"]
    for k in zerar:
        ok(f"Dashboard: {k} zerado", (kpis.get(k) or 0) == 0, f"({kpis.get(k)})")
    ok("Dashboard sem KPI vazio", all(x is not None for x in kpis.values()))

    erros = 0
    for aba in v.sheetnames:
        for row in v[aba].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(("#REF", "#VALUE", "#DIV",
                                                                    "#NAME", "#N/A", "#NUM")):
                    erros += 1
    ok("Nenhuma celula de erro", erros == 0, f"({erros})")

    al = v["ALERTAS"]
    total_al = None
    for r in range(1, 12):
        if al.cell(r, 1).value == "TOTAL DE ALERTAS ATIVOS":
            total_al = al.cell(r, 2).value
    ok("Nenhum alerta ativo no arquivo vazio", (total_al or 0) == 0, f"({total_al})")

    print("=" * 74)
    print(f"RESULTADO: {len(OKS)} verificacoes OK, {len(FALHAS)} falhas")
    print("=" * 74)
    for f in FALHAS:
        print("  FALHA:", f)
    return len(FALHAS)


if __name__ == "__main__":
    sys.exit(1 if main() else 0)
