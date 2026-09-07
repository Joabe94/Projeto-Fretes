#!/usr/bin/env python3
"""Preenche uma copia da planilha com dados conhecidos, recalcula no LibreOffice
e confere cada numero contra um calculo independente feito em Python."""
import datetime as dt, shutil, sys, random
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "gerador"))
from recalc_lo import recalc

LIMPO = RAIZ / "CONTROLE_FINANCEIRO_SIMPLES.xlsx"
TESTE = RAIZ / "simples" / "_teste.xlsx"
P1, F_LANC, F_TRF, F_CONTA, F_COMPRA, N_PARC = 6, 1005, 305, 25, 65, 36

rnd = random.Random(20260907)
CONTAS = [
    dict(nome="Conta Ueno", banco="Banco Ueno", tipo="Conta corrente", inicial=12000000),
    dict(nome="Dinheiro", banco="—", tipo="Dinheiro / Caixa", inicial=800000),
    dict(nome="Poupança", banco="Banco Ueno", tipo="Poupança", inicial=25000000),
    dict(nome="Conta Itaú", banco="Itaú", tipo="Conta corrente", inicial=-1500000),
]
NOMES = [c["nome"] for c in CONTAS]
CARTOES_N = ["Cartão 1", "Cartão 2", "Cartão 3", "Cartão 4"]
LIMITES = [20000000, 15000000, 5000000, None]
CAT_E = ["Salário", "Venda", "Aluguel recebido", "Rendimento"]
CAT_S = ["Moradia", "Alimentação", "Transporte", "Saúde", "Lazer", "Impostos e taxas"]

LANC = []
for i in range(240):
    d = dt.date(2026, rnd.randint(1, 12), rnd.randint(1, 28))
    ent = rnd.random() < 0.35
    LANC.append(dict(data=d, tipo="ENTRADA" if ent else "SAÍDA",
                     status="REALIZADO" if rnd.random() < 0.6 else "PROJETADO",
                     cat=rnd.choice(CAT_E if ent else CAT_S), desc=f"Teste {i+1}",
                     conta=rnd.choice(NOMES), valor=rnd.randrange(50, 9000) * 1000))
for i in range(20):   # outro ano, para provar que o filtro de ano nao vaza
    LANC.append(dict(data=dt.date(2027, rnd.randint(1, 12), 10), tipo="SAÍDA",
                     status="REALIZADO", cat="Lazer", desc=f"Outro ano {i+1}",
                     conta="Conta Ueno", valor=rnd.randrange(50, 900) * 1000))

TRANSF = []
for i in range(40):
    d = dt.date(2026, rnd.randint(1, 12), rnd.randint(1, 28))
    de = rnd.choice(NOMES)
    para = rnd.choice([n for n in NOMES if n != de] + CARTOES_N[:3])
    TRANSF.append(dict(data=d, status="REALIZADO" if rnd.random() < 0.7 else "PROJETADO",
                       de=de, para=para, desc=f"Transferência {i+1}",
                       valor=rnd.randrange(100, 3000) * 1000))
# casos de borda garantidos
TRANSF += [
    dict(data=dt.date(2026, 2, 10), status="REALIZADO", de="Conta Ueno", para="Cartão 1",
         desc="Fatura de fevereiro", valor=500000),
    dict(data=dt.date(2026, 1, 20), status="REALIZADO", de="Conta Ueno", para="Dinheiro",
         desc="Saque no caixa", valor=1000000),
    dict(data=dt.date(2026, 3, 10), status="PROJETADO", de="Conta Ueno", para="Cartão 1",
         desc="Fatura de março", valor=500000),
    dict(data=dt.date(2027, 1, 5), status="REALIZADO", de="Poupança", para="Conta Ueno",
         desc="Resgate no ano seguinte", valor=3000000),
]

COMPRAS = [
    dict(data=dt.date(2026, 1, 18), cartao="Cartão 1", desc="Geladeira", cat="Manutenção",
         total=6000000, n=12, primeira=None),
    dict(data=dt.date(2026, 3, 2), cartao="Cartão 1", desc="Passagem", cat="Lazer",
         total=4500000, n=6, primeira=dt.date(2026, 4, 1)),
    dict(data=dt.date(2026, 2, 8), cartao="Cartão 2", desc="Farmácia", cat="Saúde",
         total=350000, n=1, primeira=None),
    dict(data=dt.date(2026, 5, 4), cartao="Cartão 2", desc="Notebook", cat="Educação",
         total=10000000, n=7, primeira=None),
    dict(data=dt.date(2026, 12, 20), cartao="Cartão 1", desc="Presente", cat="Família",
         total=3333333, n=5, primeira=None),
    dict(data=dt.date(2025, 11, 10), cartao="Cartão 3", desc="Compra antiga", cat="Moradia",
         total=2400000, n=24, primeira=None),
]


def idx(d): return d.year * 12 + d.month

def inicio(c):
    if c["primeira"]: return c["primeira"]
    a, m = c["data"].year, c["data"].month
    return dt.date(a + 1, 1, 1) if m == 12 else dt.date(a, m + 1, 1)

def parcelas(c):
    base = round(c["total"] / c["n"]); i0 = idx(inicio(c)); out = []
    for k in range(1, c["n"] + 1):
        out.append((i0 + k - 1, c["total"] - (c["n"] - 1) * base if k == c["n"] else base))
    return out

# ----------------------------------------------------------------- preencher
shutil.copy(LIMPO, TESTE)
wb = openpyxl.load_workbook(TESTE)
con, lan, trf, car, lst = (wb["CONTAS"], wb["LANÇAMENTOS"], wb["TRANSFERÊNCIAS"],
                           wb["CARTÕES"], wb["LISTAS"])
for i, c in enumerate(CONTAS):
    r = P1 + i
    con.cell(r, 1).value = c["nome"]; con.cell(r, 2).value = c["banco"]
    con.cell(r, 3).value = c["tipo"]; con.cell(r, 4).value = c["inicial"]
    con.cell(r, 5).value = dt.date(2025, 12, 31)
for i, v in enumerate(LIMITES):
    lst.cell(P1 + i, 5).value = v
for i, l in enumerate(LANC):
    r = P1 + i
    for c, k in enumerate(["data", "tipo", "status", "cat", "desc", "conta", "valor"], 1):
        lan.cell(r, c).value = l[k]
for i, t in enumerate(TRANSF):
    r = P1 + i
    for c, k in enumerate(["data", "status", "de", "para", "desc"], 1):
        trf.cell(r, c).value = t[k]
    trf.cell(r, 6).value = t["valor"]
for i, c in enumerate(COMPRAS):
    r = P1 + i
    for c2, k in enumerate(["data", "cartao", "desc", "cat", "total", "n", "primeira"], 1):
        car.cell(r, c2).value = c[k]
wb["RESUMO MENSAL"]["A4"].value = 2026
wb["RESUMO MENSAL"]["B4"].value = 3
wb["RESUMO ANUAL"]["A4"].value = 2026
wb.save(TESTE)

print("recalculando no LibreOffice...")
res = recalc(TESTE)
if res.get("error"): print("ERRO:", res["error"]); sys.exit(1)
print(f"  {res['total_formulas']} formulas, {res['seconds']}s, {res['total_errors']} com erro")
for k, v in list(res.get("error_summary", {}).items())[:10]:
    print("   ", k, v[:12])

# ------------------------------------------------------------------ conferir
wv = openpyxl.load_workbook(TESTE, data_only=True)
ok = falhas = 0
def conf(nome, esperado, obtido, tol=0.5):
    global ok, falhas
    o = obtido if isinstance(obtido, (int, float)) else 0
    bom = abs(float(esperado) - float(o)) <= tol
    print(("  [OK   ] " if bom else "  [FALHA] ") + nome.ljust(50)
          + f"esperado {esperado:>16,.0f}  obtido {o:>16,.0f}".replace(",", "."))
    if bom: ok += 1
    else: falhas += 1

def soma(tipo=None, status=None, cat=None, ano=None, mes=None, conta=None):
    return sum(l["valor"] for l in LANC
               if (not tipo or l["tipo"] == tipo) and (not status or l["status"] == status)
               and (not cat or l["cat"] == cat) and (not ano or l["data"].year == ano)
               and (not mes or l["data"].month == mes) and (not conta or l["conta"] == conta))

def tr(status=None, de=None, para=None, ano=None, mes=None):
    return sum(t["valor"] for t in TRANSF
               if (not status or t["status"] == status) and (not de or t["de"] == de)
               and (not para or t["para"] == para) and (not ano or t["data"].year == ano)
               and (not mes or t["data"].month == mes))

def cartao_mes(i): return sum(v for c in COMPRAS for (m, v) in parcelas(c) if m == i)

cv, rm, ra, cr = wv["CONTAS"], wv["RESUMO MENSAL"], wv["RESUMO ANUAL"], wv["CARTÕES"]
hoje_i = dt.date.today().year * 12 + dt.date.today().month

print("\n[1] CONTAS — saldo de cada conta")
tot_atual = tot_prev = 0
for i, c in enumerate(CONTAS):
    r, n = P1 + i, c["nome"]
    ent = soma("ENTRADA", "REALIZADO", conta=n); sai = soma("SAÍDA", "REALIZADO", conta=n)
    rec = tr("REALIZADO", para=n); env = tr("REALIZADO", de=n)
    saldo = c["inicial"] + ent - sai + rec - env
    entp = soma("ENTRADA", "PROJETADO", conta=n); saip = soma("SAÍDA", "PROJETADO", conta=n)
    trp = tr("PROJETADO", para=n) - tr("PROJETADO", de=n)
    prev = saldo + entp - saip + trp
    tot_atual += saldo; tot_prev += prev
    conf(f"{n} · entradas realizadas", ent, cv.cell(r, 6).value)
    conf(f"{n} · saídas realizadas", sai, cv.cell(r, 7).value)
    conf(f"{n} · transferências recebidas", rec, cv.cell(r, 8).value)
    conf(f"{n} · transferências enviadas", env, cv.cell(r, 9).value)
    conf(f"{n} · SALDO ATUAL", saldo, cv.cell(r, 10).value)
    conf(f"{n} · SALDO PREVISTO", prev, cv.cell(r, 14).value)
conf("Total saldo atual (todas as contas)", tot_atual, cv["A4"].value)
conf("Total saldo previsto", tot_prev, cv["C4"].value)
pago_cartao = sum(t["valor"] for t in TRANSF if t["status"] == "REALIZADO" and t["para"] in CARTOES_N)
divida = sum(c["total"] for c in COMPRAS) - pago_cartao
conf("Dívida nos cartões (espelho de CARTÕES)", divida, cv["E4"].value)
conf("Posição líquida (contas − cartões)", tot_atual - divida, cv["G4"].value)

print("\n[2] TRANSFERÊNCIAS")
conf("Total transferido (realizado)", tr("REALIZADO"), wv["TRANSFERÊNCIAS"]["A4"].value)
conf("Pagamentos de fatura de cartão", pago_cartao, wv["TRANSFERÊNCIAS"]["C4"].value)
conf("Nenhuma transferência marcada com erro", 0, wv["TRANSFERÊNCIAS"]["F4"].value)
tv = wv["TRANSFERÊNCIAS"]
marcadas = sum(1 for r in range(P1, F_TRF + 1)
               if str(tv.cell(r, 9).value or "").startswith("Pagamento"))
conf("Linhas reconhecidas como pagamento de fatura",
     sum(1 for t in TRANSF if t["para"] in CARTOES_N), marcadas)

print("\n[3] RESUMO MENSAL — março/2026")
I3 = 2026 * 12 + 3
conf("Entradas realizadas", soma("ENTRADA", "REALIZADO", ano=2026, mes=3), rm["A7"].value)
conf("Entradas projetadas", soma("ENTRADA", "PROJETADO", ano=2026, mes=3), rm["C7"].value)
conf("Saídas realizadas", soma("SAÍDA", "REALIZADO", ano=2026, mes=3), rm["E7"].value)
conf("Saídas projetadas", soma("SAÍDA", "PROJETADO", ano=2026, mes=3), rm["G7"].value)
conf("Entradas do mês (total)", soma("ENTRADA", ano=2026, mes=3), rm["B9"].value)
conf("Saídas do mês (total)", soma("SAÍDA", ano=2026, mes=3), rm["B10"].value)
conf("Parcelas de cartão do mês", cartao_mes(I3), rm["B11"].value)
conf("Resultado do mês", soma("ENTRADA", ano=2026, mes=3) - soma("SAÍDA", ano=2026, mes=3)
     - cartao_mes(I3), rm["B12"].value)
conf("Saldo em contas hoje (espelho)", tot_atual, rm["G9"].value)
conf("Dívida nos cartões (espelho)", divida, rm["G10"].value)
conf("Transferências do mês", tr("REALIZADO", ano=2026, mes=3), rm["G11"].value)
conf("Posição líquida (espelho)", tot_atual - divida, rm["G12"].value)
for cat in ["Salário", "Venda"]:
    lin = next(r for r in range(16, 60) if rm.cell(r, 1).value == cat)
    conf(f"Entrada por categoria · {cat}", soma("ENTRADA", cat=cat, ano=2026, mes=3),
         rm.cell(lin, 4).value)
for cat in ["Moradia", "Alimentação", "Lazer"]:
    lin = next(r for r in range(16, 60) if rm.cell(r, 6).value == cat)
    conf(f"Saída por categoria · {cat}", soma("SAÍDA", cat=cat, ano=2026, mes=3),
         rm.cell(lin, 9).value)
print("  conferência do mês (tudo tem que dar zero):")
for r in range(16 + 40 + 3, 16 + 40 + 13):
    rot = rm.cell(r, 1).value
    if isinstance(rot, str) and rot and not rot.startswith("CONFER") and rm.cell(r, 4).value is not None:
        conf("  " + rot[:46], 0, rm.cell(r, 4).value)

print("\n[4] RESUMO ANUAL — 2026")
acum = 0
for m in range(1, 13):
    r = 5 + m
    e = soma("ENTRADA", ano=2026, mes=m); s = soma("SAÍDA", ano=2026, mes=m)
    c = cartao_mes(2026 * 12 + m); acum += e - s - c
    conf(f"{m:02d}/2026 entradas", e, ra.cell(r, 4).value)
    conf(f"{m:02d}/2026 saídas", s, ra.cell(r, 7).value)
    conf(f"{m:02d}/2026 cartão", c, ra.cell(r, 8).value)
    conf(f"{m:02d}/2026 resultado", e - s - c, ra.cell(r, 9).value)
    conf(f"{m:02d}/2026 acumulado", acum, ra.cell(r, 10).value)
    conf(f"{m:02d}/2026 transferências", tr(ano=2026, mes=m), ra.cell(r, 11).value)
conf("Total do ano · entradas", soma("ENTRADA", ano=2026), ra.cell(18, 4).value)
conf("Total do ano · saídas", soma("SAÍDA", ano=2026), ra.cell(18, 7).value)
conf("Total do ano · cartão", sum(cartao_mes(2026 * 12 + m) for m in range(1, 13)),
     ra.cell(18, 8).value)
conf("Saldo acumulado final", acum, ra.cell(18, 10).value)
conf("Total do ano · transferências", tr(ano=2026), ra.cell(18, 11).value)
print("  conferência do ano:")
for r in range(ra.max_row - 4, ra.max_row + 1):
    rot = ra.cell(r, 1).value
    if isinstance(rot, str) and ra.cell(r, 4).value is not None and not rot.startswith("CONFER"):
        conf("  " + rot[:46], 0, ra.cell(r, 4).value)

print("\n[5] CARTÕES")
conf("Total comprado", sum(c["total"] for c in COMPRAS), cr["A4"].value)
vencidas = sum(v for c in COMPRAS for m, v in parcelas(c) if m <= hoje_i)
conf("Parcelas a vencer", sum(c["total"] for c in COMPRAS) - vencidas, cr["C4"].value)
conf("Pago via transferências", pago_cartao, cr["E4"].value)
conf("DÍVIDA ATUAL", divida, cr["G4"].value)
for i, c in enumerate(COMPRAS):
    r, p = P1 + i, parcelas(c)
    conf(f"Compra {i+1} · soma das parcelas = total", c["total"], sum(v for _, v in p))
    conf(f"Compra {i+1} · valor da parcela", round(c["total"] / c["n"]), cr.cell(r, 9).value)
    conf(f"Compra {i+1} · parcelas vencidas", sum(v for m, v in p if m <= hoje_i), cr.cell(r, 12).value)
    conf(f"Compra {i+1} · parcelas a vencer", sum(v for m, v in p if m > hoje_i), cr.cell(r, 13).value)
r0 = F_COMPRA + 3
for i, nome in enumerate(CARTOES_N):
    r = r0 + 2 + i
    tot = sum(c["total"] for c in COMPRAS if c["cartao"] == nome)
    pg = sum(t["valor"] for t in TRANSF if t["status"] == "REALIZADO" and t["para"] == nome)
    conf(f"{nome} · total comprado", tot, cr.cell(r, 3).value)
    conf(f"{nome} · parcelas a vencer",
         sum(v for c in COMPRAS if c["cartao"] == nome for m, v in parcelas(c) if m > hoje_i),
         cr.cell(r, 5).value)
    conf(f"{nome} · pago (transferências)", pg, cr.cell(r, 6).value)
    conf(f"{nome} · dívida atual", tot - pg, cr.cell(r, 7).value)
    if LIMITES[i] is not None:
        conf(f"{nome} · limite disponível", LIMITES[i] - (tot - pg), cr.cell(r, 8).value)

print("\n[6] PARCELAS — grade automática")
pa = wv["PARCELAS"]
for i, c in enumerate(COMPRAS):
    r, esp = P1 + i, parcelas(c)
    got = [(pa.cell(r, 2 + k).value, pa.cell(r, 38 + k).value) for k in range(1, N_PARC + 1)]
    got = [(m, v) for m, v in got if m]
    conf(f"Compra {i+1} · nº de parcelas geradas", len(esp), len(got))
    conf(f"Compra {i+1} · 1º mês", esp[0][0], got[0][0] if got else 0)
    conf(f"Compra {i+1} · último mês", esp[-1][0], got[-1][0] if got else 0)
    conf(f"Compra {i+1} · soma na grade", c["total"], sum(v for _, v in got))
conf("Linhas sem compra não geram parcela", 0,
     sum(1 for r in range(P1 + len(COMPRAS), F_COMPRA + 1)
         for k in range(1, N_PARC + 1) if pa.cell(r, 2 + k).value))

print("\n[7] Regras econômicas — o dinheiro não pode ser contado duas vezes")
conf("Transferência não entra em entradas do ano", soma("ENTRADA", ano=2026), ra.cell(18, 4).value)
conf("Transferência não entra em saídas do ano", soma("SAÍDA", ano=2026), ra.cell(18, 7).value)
conf("Pagamento de fatura não vira despesa",
     soma("SAÍDA", ano=2026), ra.cell(18, 7).value)
soma_saldos = sum(cv.cell(P1 + i, 10).value or 0 for i in range(len(CONTAS)))
conf("Soma dos saldos das contas = total do topo", soma_saldos, cv["A4"].value)
mov = (sum(c["inicial"] for c in CONTAS) + soma("ENTRADA", "REALIZADO")
       - soma("SAÍDA", "REALIZADO") + tr("REALIZADO")
       - tr("REALIZADO") - pago_cartao)
conf("Caixa total = inicial + entradas − saídas − faturas pagas", mov, cv["A4"].value)

print("\n" + "=" * 96)
print(f"RESULTADO: {ok} conferências OK, {falhas} falhas, {res['total_errors']} células com erro")
print("=" * 96)
sys.exit(1 if falhas or res["total_errors"] else 0)
