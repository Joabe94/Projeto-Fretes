#!/usr/bin/env python3
"""Testes de mutacao: introduz um erro de cada vez e exige que a conferencia
correspondente acuse. Uma conferencia que so sabe dar zero nao serve de nada."""
import datetime as dt, shutil, sys
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "gerador"))
from recalc_lo import recalc

LIMPO = RAIZ / "CONTROLE_FINANCEIRO_SIMPLES.xlsx"
TMP = RAIZ / "simples" / "_mut.xlsx"
P1 = 6

BASE_CONTAS = [("Conta Ueno", "Conta corrente", 10000000), ("Dinheiro", "Dinheiro / Caixa", 500000)]
BASE_LANC = [(dt.date(2026, 1, 5), "ENTRADA", "REALIZADO", "Salário", "Salário", "Conta Ueno", 8000000),
             (dt.date(2026, 1, 9), "SAÍDA", "REALIZADO", "Moradia", "Aluguel", "Conta Ueno", 2000000)]
BASE_TRF = [(dt.date(2026, 1, 15), "REALIZADO", "Conta Ueno", "Dinheiro", "Saque", 1000000)]
BASE_CAR = [(dt.date(2026, 1, 18), "Cartão 1", "Geladeira", "Manutenção", 6000000, 12, None)]


def montar(mutacao=None):
    shutil.copy(LIMPO, TMP)
    wb = openpyxl.load_workbook(TMP)
    con, lan, trf, car = wb["CONTAS"], wb["LANÇAMENTOS"], wb["TRANSFERÊNCIAS"], wb["CARTÕES"]
    for i, (n, t, s) in enumerate(BASE_CONTAS):
        r = P1 + i
        con.cell(r, 1).value = n; con.cell(r, 3).value = t; con.cell(r, 4).value = s
    for i, l in enumerate(BASE_LANC):
        for c, v in enumerate(l, 1): lan.cell(P1 + i, c).value = v
    for i, t in enumerate(BASE_TRF):
        for c, v in enumerate(t[:5], 1): trf.cell(P1 + i, c).value = v
        trf.cell(P1 + i, 6).value = t[5]
    for i, c0 in enumerate(BASE_CAR):
        for c, v in enumerate(c0, 1): car.cell(P1 + i, c).value = v
    wb["RESUMO MENSAL"]["A4"].value = 2026
    wb["RESUMO MENSAL"]["B4"].value = 1
    wb["RESUMO ANUAL"]["A4"].value = 2026
    if mutacao: mutacao(wb)
    wb.save(TMP)
    r = recalc(TMP)
    return openpyxl.load_workbook(TMP, data_only=True), r


# linhas da conferencia do RESUMO MENSAL (a secao comeca em 16+40+2 = 58)
CONF_LIN = {
    "entradas_categoria": 59, "saidas_categoria": 60, "campos_vazios": 61,
    "conta_inexistente": 62, "transf_erro": 63, "transf_destino": 64,
    "cartao_sem_parcela": 65, "parcelas_total": 66,
}

ok = falhas = 0
def conf(nome, esperado, obtido):
    global ok, falhas
    bom = bool(esperado) == bool(obtido) if isinstance(esperado, bool) else esperado == obtido
    print(("  [OK   ] " if bom else "  [FALHA] ") + nome.ljust(58)
          + f"esperado {esperado!r:>12}  obtido {obtido!r:>12}")
    if bom: ok += 1
    else: falhas += 1


def valor_conf(wb, chave):
    return wb["RESUMO MENSAL"].cell(CONF_LIN[chave], 4).value


print("[0] Sem nenhum erro, todas as conferências dão zero")
wb, res = montar()
rm = wb["RESUMO MENSAL"]
for chave, lin in CONF_LIN.items():
    rot = str(rm.cell(lin, 1).value or "")[:44]
    conf(f"{chave} · {rot}", 0, rm.cell(lin, 4).value)
conf("nenhuma célula de erro", 0, res["total_errors"])
base_saldo = wb["CONTAS"]["A4"].value
conf("saldo inicial do cenário (10.500.000 + 8.000.000 − 2.000.000)", 16500000, base_saldo)

print("\n[1] Lançamento sem conta preenchida")
wb, _ = montar(lambda w: setattr(w["LANÇAMENTOS"].cell(P1, 6), "value", None))
conf("conferência acusa campo vazio", True, valor_conf(wb, "campos_vazios") > 0)

print("\n[2] Lançamento com conta que não existe")
wb, _ = montar(lambda w: setattr(w["LANÇAMENTOS"].cell(P1, 6), "value", "Conta Fantasma"))
conf("conferência acusa conta inexistente", True, valor_conf(wb, "conta_inexistente") > 0)
conf("o dinheiro some do saldo das contas", True, wb["CONTAS"]["A4"].value != 16500000)

print("\n[3] Transferência com origem igual ao destino")
def mut3(w):
    w["TRANSFERÊNCIAS"].cell(P1, 4).value = "Conta Ueno"
wb, _ = montar(mut3)
conf("coluna Conferência marca ERRO", True,
     str(wb["TRANSFERÊNCIAS"].cell(P1, 9).value or "").startswith("ERRO"))
conf("conferência do mês acusa", True, valor_conf(wb, "transf_erro") > 0)

print("\n[4] Transferência sem valor")
wb, _ = montar(lambda w: setattr(w["TRANSFERÊNCIAS"].cell(P1, 6), "value", None))
conf("coluna Conferência marca ERRO", True,
     str(wb["TRANSFERÊNCIAS"].cell(P1, 9).value or "").startswith("ERRO"))
conf("conferência do mês acusa", True, valor_conf(wb, "transf_erro") > 0)

print("\n[5] Transferência para um destino que não existe")
wb, _ = montar(lambda w: setattr(w["TRANSFERÊNCIAS"].cell(P1, 4), "value", "Banco Inventado"))
conf("conferência acusa destino perdido", True, valor_conf(wb, "transf_destino") != 0)

print("\n[6] Compra de cartão sem nº de parcelas")
wb, _ = montar(lambda w: setattr(w["CARTÕES"].cell(P1, 6), "value", None))
conf("conferência acusa compra sem parcelas", True, valor_conf(wb, "cartao_sem_parcela") > 0)
conf("a compra não gera nenhuma parcela", 0,
     sum(1 for k in range(1, 37) if wb["PARCELAS"].cell(P1, 2 + k).value))

print("\n[7] Pagamento de fatura lançado errado, como SAÍDA")
def mut7(w):
    lan = w["LANÇAMENTOS"]
    lan.cell(P1 + 2, 1).value = dt.date(2026, 2, 10)
    lan.cell(P1 + 2, 2).value = "SAÍDA"; lan.cell(P1 + 2, 3).value = "REALIZADO"
    lan.cell(P1 + 2, 4).value = "Moradia"; lan.cell(P1 + 2, 5).value = "Fatura do cartão"
    lan.cell(P1 + 2, 6).value = "Conta Ueno"; lan.cell(P1 + 2, 7).value = 500000
    w["RESUMO MENSAL"]["B4"].value = 2
wb, _ = montar(mut7)
fev_saidas = wb["RESUMO MENSAL"]["B10"].value
fev_cartao = wb["RESUMO MENSAL"]["B11"].value
conf("a saída errada aparece nas saídas de fevereiro", 500000, fev_saidas)
conf("a parcela do cartão continua contando em fevereiro", 500000, fev_cartao)
print("         ^ é isto que as INSTRUÇÕES proíbem: o mesmo gasto conta duas vezes")
print("           (Gs. 1.000.000 no mês em vez de Gs. 500.000). Nenhuma fórmula")
print("           consegue adivinhar isso — por isso a regra está escrita em 3 abas.")

print("\n[8] Pagamento de fatura lançado do jeito certo")
def mut8(w):
    t = w["TRANSFERÊNCIAS"]
    t.cell(P1 + 1, 1).value = dt.date(2026, 2, 10)
    t.cell(P1 + 1, 2).value = "REALIZADO"; t.cell(P1 + 1, 3).value = "Conta Ueno"
    t.cell(P1 + 1, 4).value = "Cartão 1"; t.cell(P1 + 1, 5).value = "Fatura"
    t.cell(P1 + 1, 6).value = 500000
    w["RESUMO MENSAL"]["B4"].value = 2
wb, _ = montar(mut8)
conf("não vira saída de fevereiro", 0, wb["RESUMO MENSAL"]["B10"].value)
conf("a parcela do cartão continua contando", 500000, wb["RESUMO MENSAL"]["B11"].value)
conf("sai do saldo da conta", 16000000, wb["CONTAS"]["A4"].value)
conf("abate a dívida do cartão", 5500000, wb["CARTÕES"]["G4"].value)
conf("é marcada como pagamento de fatura", "Pagamento de fatura de cartão",
     wb["TRANSFERÊNCIAS"].cell(P1 + 1, 9).value)
conf("todas as conferências continuam em zero", True,
     all(wb["RESUMO MENSAL"].cell(l, 4).value == 0 for l in CONF_LIN.values()))

print("\n[9] Transferência entre contas não muda o total, só o lugar do dinheiro")
def mut9(w):
    t = w["TRANSFERÊNCIAS"]
    t.cell(P1 + 1, 1).value = dt.date(2026, 1, 20)
    t.cell(P1 + 1, 2).value = "REALIZADO"; t.cell(P1 + 1, 3).value = "Conta Ueno"
    t.cell(P1 + 1, 4).value = "Dinheiro"; t.cell(P1 + 1, 5).value = "Saque 2"
    t.cell(P1 + 1, 6).value = 3000000
wb, _ = montar(mut9)
conf("saldo total não muda", 16500000, wb["CONTAS"]["A4"].value)
# 10.000.000 inicial + 8.000.000 entrada - 2.000.000 saida - 1.000.000 saque - 3.000.000
conf("Conta Ueno perde o valor", 12000000, wb["CONTAS"].cell(P1, 10).value)
conf("Dinheiro ganha o valor", 500000 + 4000000, wb["CONTAS"].cell(P1 + 1, 10).value)
conf("não aparece como entrada do mês", 8000000, wb["RESUMO MENSAL"]["B9"].value)
conf("não aparece como saída do mês", 2000000, wb["RESUMO MENSAL"]["B10"].value)

TMP.unlink(missing_ok=True)
print("\n" + "=" * 96)
print(f"RESULTADO: {ok} conferências OK, {falhas} falhas")
print("=" * 96)
sys.exit(1 if falhas else 0)
