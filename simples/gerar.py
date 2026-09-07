#!/usr/bin/env python3
"""Gera CONTROLE_FINANCEIRO_SIMPLES.xlsx - entrada/saida, projetado,
cartao com parcelas automaticas, resumo mensal e resumo anual.
Moeda: Guarani. Todas as formulas usam virgula como separador de argumentos."""
import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment

SAIDA = Path(__file__).resolve().parent.parent / "CONTROLE_FINANCEIRO_SIMPLES.xlsx"
ANO = 2026
HOJE = dt.date(2026, 9, 7)
IDX_HOJE = HOJE.year * 12 + HOJE.month

LIN_HDR = 5          # linha do cabecalho das tabelas de digitacao
LIN_1 = 6            # primeira linha de dados
N_LANC = 1000        # linhas de lancamentos
N_COMPRA = 60        # linhas de compras no cartao
N_PARC = 36          # parcelas maximas por compra
FIM_LANC = LIN_1 + N_LANC - 1        # 1005
FIM_COMPRA = LIN_1 + N_COMPRA - 1    # 65

GS = '"Gs. "#,##0;-"Gs. "#,##0;"Gs. "0'
GS0 = '"Gs. "#,##0;-"Gs. "#,##0;""'   # esconde zero
NUM0 = '#,##0;-#,##0;""'
DATA = "DD/MM/YYYY"

AZUL = "1F3864"
AZUL_CLARO = "D9E2F3"
CINZA = "F2F2F2"
VERDE = "E2EFDA"
VERMELHO = "FCE4E4"
AMARELO = "FFF2CC"
BRANCO = "FFFFFF"

f_titulo = Font(name="Calibri", size=16, bold=True, color=AZUL)
f_sub = Font(name="Calibri", size=10, italic=True, color="595959")
f_hdr = Font(name="Calibri", size=11, bold=True, color=BRANCO)
f_b = Font(name="Calibri", size=11, bold=True)
f_n = Font(name="Calibri", size=11)
f_auto = Font(name="Calibri", size=11, color="7F7F7F", italic=True)
f_kpi = Font(name="Calibri", size=14, bold=True, color=AZUL)

fill_hdr = PatternFill("solid", fgColor=AZUL)
fill_az = PatternFill("solid", fgColor=AZUL_CLARO)
fill_cz = PatternFill("solid", fgColor=CINZA)
fill_vd = PatternFill("solid", fgColor=VERDE)
fill_vm = PatternFill("solid", fgColor=VERMELHO)
fill_am = PatternFill("solid", fgColor=AMARELO)

fina = Side(style="thin", color="BFBFBF")
bd = Border(left=fina, right=fina, top=fina, bottom=fina)
ctr = Alignment(horizontal="center", vertical="center")
esq = Alignment(horizontal="left", vertical="center")
wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

CAT_ENTRADA = ["Salário", "Pró-labore", "Aluguel recebido", "Venda", "Rendimento",
               "Empréstimo recebido", "Reembolso", "Outras entradas"]
CAT_SAIDA = ["Moradia", "Alimentação", "Transporte", "Saúde", "Educação", "Lazer",
             "Vestuário", "Serviços (luz, água, internet)", "Impostos e taxas",
             "Empréstimo / Financiamento", "Despesas do cartão", "Família",
             "Manutenção", "Outras saídas"]
MEIOS = ["Dinheiro", "Conta corrente", "Poupança", "Transferência", "Débito", "Boleto"]
CARTOES = ["Cartão 1", "Cartão 2", "Cartão 3", "Cartão 4"]
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
         "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]

wb = openpyxl.Workbook()


def cel(ws, r, c, v=None, *, fonte=None, fill=None, fmt=None, al=None, borda=True):
    x = ws.cell(row=r, column=c)
    x.value = v
    if fonte: x.font = fonte
    if fill: x.fill = fill
    if fmt: x.number_format = fmt
    if al: x.alignment = al
    if borda: x.border = bd
    return x


def cabecalho(ws, titulo, sub, largura=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=largura)
    c = cel(ws, 1, 1, titulo, fonte=f_titulo, al=esq, borda=False)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=largura)
    cel(ws, 2, 1, sub, fonte=f_sub, al=esq, borda=False)
    ws.row_dimensions[1].height = 24
    return c


def linha_hdr(ws, linha, rotulos, col0=1):
    for i, t in enumerate(rotulos):
        cel(ws, linha, col0 + i, t, fonte=f_hdr, fill=fill_hdr, al=ctr)
    ws.row_dimensions[linha].height = 30


def larguras(ws, mapa):
    for col, w in mapa.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- INSTRUÇÕES
ws = wb.active
ws.title = "INSTRUÇÕES"
ws.sheet_properties.tabColor = AZUL
larguras(ws, {"A": 3, "B": 100})
cabecalho(ws, "CONTROLE FINANCEIRO — ENTRADA, SAÍDA E CARTÃO", "Moeda: Guaraní (Gs.)  ·  Ano-base: %d" % ANO, 2)

TEXTO = [
    ("t", "AS 3 REGRAS QUE FAZEM A PLANILHA FECHAR"),
    ("n", "1) Dinheiro que entra ou sai da sua mão / conta → lança na aba LANÇAMENTOS."),
    ("n", "2) Compra no cartão de crédito → lança SÓ na aba CARTÕES. Não repita em LANÇAMENTOS."),
    ("n", "   A planilha divide a compra nas parcelas e joga cada parcela no mês certo, sozinha."),
    ("n", "3) Quando você paga a fatura do cartão, NÃO lance nada. A parcela já foi contada no mês dela."),
    ("a", "Seguindo essas 3 regras é impossível contar o mesmo gasto duas vezes."),
    ("", ""),
    ("t", "COMO USAR, PASSO A PASSO"),
    ("n", "1º  Abra a aba LISTAS e escreva o nome dos seus cartões e, se quiser, ajuste as categorias."),
    ("n", "2º  Abra a aba LANÇAMENTOS e comece a digitar. Data, Tipo, Status, Categoria, Descrição, Meio, Valor."),
    ("n", "3º  Compras parceladas no cartão vão na aba CARTÕES: data, cartão, valor total e nº de parcelas."),
    ("n", "4º  Veja o mês na aba RESUMO MENSAL (escolha ano e mês no alto) e o ano na aba RESUMO ANUAL."),
    ("", ""),
    ("t", "REALIZADO x PROJETADO"),
    ("n", "REALIZADO = já aconteceu, o dinheiro já entrou ou já saiu."),
    ("n", "PROJETADO = ainda vai acontecer. Serve para lançar salário futuro, aluguel, conta de luz prevista."),
    ("n", "Os resumos mostram as duas colunas separadas e também o total, para você comparar."),
    ("", ""),
    ("t", "O QUE VOCÊ NUNCA DEVE FAZER"),
    ("x", "Não EXCLUA linhas das tabelas. Para apagar um lançamento, selecione a linha e aperte DELETE"),
    ("x", "   (limpar o conteúdo). Excluir a linha apaga as fórmulas das colunas automáticas."),
    ("x", "Não digite nada nas colunas cinza escritas (automático). Elas se calculam sozinhas."),
    ("x", "Não mexa na aba PARCELAS. Ela é o motor que espalha as parcelas pelos meses."),
    ("", ""),
    ("t", "AS ABAS"),
    ("n", "LANÇAMENTOS  ·  onde você digita tudo que é dinheiro em espécie, conta, débito, boleto."),
    ("n", "CARTÕES      ·  compras no cartão de crédito, com parcelamento automático."),
    ("n", "RESUMO MENSAL·  um mês por vez, com o detalhe por categoria e o resultado do mês."),
    ("n", "RESUMO ANUAL ·  os 12 meses do ano lado a lado e o total do ano."),
    ("n", "EXEMPLO      ·  modelos preenchidos de cada tipo de lançamento. Pode apagar a aba quando quiser."),
    ("n", "LISTAS       ·  suas categorias, meios de pagamento e nomes dos cartões."),
    ("n", "PARCELAS     ·  aba de apoio, automática. Não mexa."),
    ("", ""),
    ("t", "COMO O RESULTADO DO MÊS É CALCULADO"),
    ("a", "Resultado = Entradas − Saídas − Parcelas de cartão que vencem no mês"),
    ("n", "O saldo acumulado no RESUMO ANUAL soma o resultado de cada mês, de janeiro em diante."),
]
r = 4
for tipo, txt in TEXTO:
    c = cel(ws, r, 2, txt, borda=False, al=wrap)
    if tipo == "t":
        c.font = Font(size=12, bold=True, color=AZUL); c.fill = fill_az
    elif tipo == "a":
        c.font = Font(size=11, bold=True); c.fill = fill_am
    elif tipo == "x":
        c.font = Font(size=11, color="9C0006"); c.fill = fill_vm
    else:
        c.font = f_n
    ws.row_dimensions[r].height = 17
    r += 1
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------------- LISTAS
lst = wb.create_sheet("LISTAS")
lst.sheet_properties.tabColor = "808080"
cabecalho(lst, "LISTAS", "Ajuste aqui suas categorias, meios de pagamento e o nome dos seus cartões.", 6)
larguras(lst, {"A": 32, "B": 14, "C": 3, "D": 26, "E": 3, "F": 26})
linha_hdr(lst, LIN_HDR, ["Categoria", "Tipo"], 1)
linha_hdr(lst, LIN_HDR, ["Meio de pagamento"], 4)
linha_hdr(lst, LIN_HDR, ["Cartões de crédito"], 6)

CATS = [(c, "ENTRADA") for c in CAT_ENTRADA] + [(c, "SAÍDA") for c in CAT_SAIDA]
N_CAT = 40
for i in range(N_CAT):
    r = LIN_1 + i
    v = CATS[i] if i < len(CATS) else ("", "")
    cel(lst, r, 1, v[0] or None, fonte=f_n, al=esq)
    cel(lst, r, 2, v[1] or None, fonte=f_n, al=ctr,
        fill=fill_vd if v[1] == "ENTRADA" else (fill_vm if v[1] == "SAÍDA" else None))
FIM_CAT = LIN_1 + N_CAT - 1

for i in range(12):
    cel(lst, LIN_1 + i, 4, MEIOS[i] if i < len(MEIOS) else None, fonte=f_n, al=esq)
FIM_MEIO = LIN_1 + 11
for i in range(8):
    cel(lst, LIN_1 + i, 6, CARTOES[i] if i < len(CARTOES) else None, fonte=f_n, al=esq)
FIM_CART = LIN_1 + 7
cel(lst, FIM_CART + 2, 6, "Troque \"Cartão 1\" pelo nome real do seu cartão.", fonte=f_sub, borda=False)

lst2 = "TIPOS"
for i, v in enumerate(["ENTRADA", "SAÍDA"]):
    cel(lst, LIN_1 + i, 8, v, fonte=f_n, al=ctr)
for i, v in enumerate(["REALIZADO", "PROJETADO"]):
    cel(lst, LIN_1 + i, 9, v, fonte=f_n, al=ctr)
cel(lst, LIN_HDR, 8, "Tipo", fonte=f_hdr, fill=fill_hdr, al=ctr)
cel(lst, LIN_HDR, 9, "Status", fonte=f_hdr, fill=fill_hdr, al=ctr)
larguras(lst, {"H": 14, "I": 14})
lst.sheet_view.showGridLines = False

R_CAT = f"LISTAS!$A${LIN_1}:$A${FIM_CAT}"
R_CATTIPO = f"LISTAS!$B${LIN_1}:$B${FIM_CAT}"
R_MEIO = f"LISTAS!$D${LIN_1}:$D${FIM_MEIO}"
R_CARTAO = f"LISTAS!$F${LIN_1}:$F${FIM_CART}"
R_TIPO = f"LISTAS!$H${LIN_1}:$H${LIN_1 + 1}"
R_STATUS = f"LISTAS!$I${LIN_1}:$I${LIN_1 + 1}"

# ---------------------------------------------------------------- LANÇAMENTOS
lan = wb.create_sheet("LANÇAMENTOS")
lan.sheet_properties.tabColor = "2E75B6"
cabecalho(lan, "LANÇAMENTOS", "Dinheiro que entra e sai da sua conta ou do seu bolso. Compra no cartão NÃO vem aqui — vai na aba CARTÕES.", 9)
larguras(lan, {"A": 12, "B": 12, "C": 13, "D": 26, "E": 34, "F": 17, "G": 16, "H": 11, "I": 11})

# faixa de totais (linha 3 e 4)
cel(lan, 3, 1, "Total ENTRADAS", fonte=f_b, fill=fill_vd, al=ctr)
cel(lan, 4, 1, f"=SUMIF($B${LIN_1}:$B${FIM_LANC},\"ENTRADA\",$G${LIN_1}:$G${FIM_LANC})",
    fonte=f_kpi, fill=fill_vd, fmt=GS, al=ctr)
cel(lan, 3, 2, "Total SAÍDAS", fonte=f_b, fill=fill_vm, al=ctr)
cel(lan, 4, 2, f"=SUMIF($B${LIN_1}:$B${FIM_LANC},\"SAÍDA\",$G${LIN_1}:$G${FIM_LANC})",
    fonte=f_kpi, fill=fill_vm, fmt=GS, al=ctr)
cel(lan, 3, 3, "Saldo (sem cartão)", fonte=f_b, fill=fill_az, al=ctr)
cel(lan, 4, 3, "=A4-B4", fonte=f_kpi, fill=fill_az, fmt=GS, al=ctr)
cel(lan, 3, 4, "Lançamentos digitados", fonte=f_b, fill=fill_cz, al=ctr)
cel(lan, 4, 4, f"=COUNTA($A${LIN_1}:$A${FIM_LANC})", fonte=f_kpi, fill=fill_cz, fmt="#,##0", al=ctr)
cel(lan, 3, 5, "Espaço livre para digitar", fonte=f_b, fill=fill_cz, al=ctr)
cel(lan, 4, 5, f"={N_LANC}-D4", fonte=f_kpi, fill=fill_cz, fmt="#,##0", al=ctr)

linha_hdr(lan, LIN_HDR, ["Data", "Tipo", "Status", "Categoria", "Descrição",
                         "Meio de pagamento", "Valor (Gs.)", "Mês (auto)", "Índice (auto)"])
lan.freeze_panes = "A6"

for r in range(LIN_1, FIM_LANC + 1):
    cel(lan, r, 1, None, fonte=f_n, fmt=DATA, al=ctr)
    cel(lan, r, 2, None, fonte=f_n, al=ctr)
    cel(lan, r, 3, None, fonte=f_n, al=ctr)
    cel(lan, r, 4, None, fonte=f_n, al=esq)
    cel(lan, r, 5, None, fonte=f_n, al=esq)
    cel(lan, r, 6, None, fonte=f_n, al=ctr)
    cel(lan, r, 7, None, fonte=f_n, fmt=GS, al=ctr)
    cel(lan, r, 8, f'=IF($A{r}="","",MONTH($A{r})&"/"&YEAR($A{r}))',
        fonte=f_auto, fill=fill_cz, al=ctr)
    cel(lan, r, 9, f'=IF($A{r}="",0,YEAR($A{r})*12+MONTH($A{r}))',
        fonte=f_auto, fill=fill_cz, fmt=NUM0, al=ctr)

dv = DataValidation(type="list", formula1=f"={R_TIPO}", allow_blank=True, showDropDown=False)
dv.error = "Escolha ENTRADA ou SAÍDA na lista."; dv.errorTitle = "Tipo inválido"
lan.add_data_validation(dv); dv.add(f"B{LIN_1}:B{FIM_LANC}")

dv = DataValidation(type="list", formula1=f"={R_STATUS}", allow_blank=True, showDropDown=False)
dv.error = "Escolha REALIZADO (já aconteceu) ou PROJETADO (ainda vai acontecer)."
dv.errorTitle = "Status inválido"
lan.add_data_validation(dv); dv.add(f"C{LIN_1}:C{FIM_LANC}")

dv = DataValidation(type="list", formula1=f"={R_CAT}", allow_blank=True, showDropDown=False)
dv.error = "Categoria não está na lista. Cadastre na aba LISTAS."; dv.errorTitle = "Categoria"
lan.add_data_validation(dv); dv.add(f"D{LIN_1}:D{FIM_LANC}")

dv = DataValidation(type="list", formula1=f"={R_MEIO}", allow_blank=True, showDropDown=False)
dv.error = "Meio de pagamento não está na lista. Cadastre na aba LISTAS."
dv.errorTitle = "Meio de pagamento"
lan.add_data_validation(dv); dv.add(f"F{LIN_1}:F{FIM_LANC}")

dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
dv.error = "O valor tem que ser maior que zero. Se é saída, escolha Tipo = SAÍDA."
dv.errorTitle = "Valor inválido"
lan.add_data_validation(dv); dv.add(f"G{LIN_1}:G{FIM_LANC}")

lan.conditional_formatting.add(
    f"A{LIN_1}:I{FIM_LANC}",
    CellIsRule(operator="equal", formula=[f'"PROJETADO"'], fill=fill_am, stopIfTrue=False))
lan["C6"].comment = Comment("REALIZADO = já entrou ou já saiu.\nPROJETADO = ainda vai acontecer.", "Planilha")
lan.sheet_view.showGridLines = False

# -------------------------------------------------------------------- CARTÕES
car = wb.create_sheet("CARTÕES")
car.sheet_properties.tabColor = "C00000"
cabecalho(car, "CARTÕES DE CRÉDITO — COMPRAS E PARCELAS",
          "Digite a compra uma vez. A planilha divide nas parcelas e joga cada uma no mês certo. Não lance a fatura em LANÇAMENTOS.", 13)
larguras(car, {"A": 13, "B": 16, "C": 30, "D": 24, "E": 16, "F": 10, "G": 14,
               "H": 13, "I": 15, "J": 13, "K": 11, "L": 15, "M": 16})

cel(car, 3, 1, "Total comprado", fonte=f_b, fill=fill_cz, al=ctr)
cel(car, 4, 1, f"=SUM($E${LIN_1}:$E${FIM_COMPRA})", fonte=f_kpi, fill=fill_cz, fmt=GS, al=ctr)
cel(car, 3, 2, "Já pago (até hoje)", fonte=f_b, fill=fill_vd, al=ctr)
cel(car, 4, 2, f"=SUM($L${LIN_1}:$L${FIM_COMPRA})", fonte=f_kpi, fill=fill_vd, fmt=GS, al=ctr)
cel(car, 3, 3, "Falta pagar (saldo devedor)", fonte=f_b, fill=fill_vm, al=ctr)
cel(car, 4, 3, f"=SUM($M${LIN_1}:$M${FIM_COMPRA})", fonte=f_kpi, fill=fill_vm, fmt=GS, al=ctr)
cel(car, 3, 4, "Mês de referência de hoje", fonte=f_b, fill=fill_az, al=ctr)
cel(car, 4, 4, "=MONTH(TODAY())&\"/\"&YEAR(TODAY())", fonte=f_kpi, fill=fill_az, al=ctr)

linha_hdr(car, LIN_HDR, ["Data da compra", "Cartão", "Descrição", "Categoria",
                         "Valor total (Gs.)", "Nº de parcelas", "1ª parcela (opcional)",
                         "Início (auto)", "Valor da parcela (auto)", "Última parcela (auto)",
                         "Índice (auto)", "Já pago (auto)", "Falta pagar (auto)"])
car.freeze_panes = "A6"

for r in range(LIN_1, FIM_COMPRA + 1):
    cel(car, r, 1, None, fonte=f_n, fmt=DATA, al=ctr)
    cel(car, r, 2, None, fonte=f_n, al=ctr)
    cel(car, r, 3, None, fonte=f_n, al=esq)
    cel(car, r, 4, None, fonte=f_n, al=esq)
    cel(car, r, 5, None, fonte=f_n, fmt=GS, al=ctr)
    cel(car, r, 6, None, fonte=f_n, fmt="0", al=ctr)
    cel(car, r, 7, None, fonte=f_n, fmt=DATA, al=ctr)
    # H = inicio automatico: se 1a parcela em branco, mes seguinte ao da compra
    cel(car, r, 8,
        f'=IF($A{r}="","",IF($G{r}<>"",$G{r},'
        f'DATE(YEAR($A{r})+IF(MONTH($A{r})=12,1,0),IF(MONTH($A{r})=12,1,MONTH($A{r})+1),1)))',
        fonte=f_auto, fill=fill_cz, fmt=DATA, al=ctr)
    # I = valor da parcela
    cel(car, r, 9, f'=IF(OR($E{r}="",$F{r}=""),"",ROUND($E{r}/$F{r},0))',
        fonte=f_auto, fill=fill_cz, fmt=GS, al=ctr)
    # K = indice do mes da 1a parcela
    cel(car, r, 11, f'=IF($H{r}="",0,YEAR($H{r})*12+MONTH($H{r}))',
        fonte=f_auto, fill=fill_cz, fmt=NUM0, al=ctr)
    # J = ultima parcela em mes/ano
    cel(car, r, 10,
        f'=IF($K{r}=0,"",'
        f'($K{r}+$F{r}-1-INT(($K{r}+$F{r}-2)/12)*12)&"/"&INT(($K{r}+$F{r}-2)/12))',
        fonte=f_auto, fill=fill_cz, al=ctr)
    # L = ja pago ate hoje  (soma das parcelas com mes <= mes de hoje)
    cel(car, r, 12,
        f'=IF($K{r}=0,"",SUMPRODUCT((PARCELAS!$C{r}:$AL{r}>0)*'
        f'(PARCELAS!$C{r}:$AL{r}<=YEAR(TODAY())*12+MONTH(TODAY()))*PARCELAS!$AM{r}:$BV{r}))',
        fonte=f_auto, fill=fill_cz, fmt=GS, al=ctr)
    # M = falta pagar
    cel(car, r, 13, f'=IF($K{r}=0,"",$E{r}-$L{r})', fonte=f_auto, fill=fill_cz, fmt=GS, al=ctr)

dv = DataValidation(type="list", formula1=f"={R_CARTAO}", allow_blank=True, showDropDown=False)
dv.error = "Cartão não está na lista. Cadastre o nome dele na aba LISTAS."
dv.errorTitle = "Cartão"
car.add_data_validation(dv); dv.add(f"B{LIN_1}:B{FIM_COMPRA}")

dv = DataValidation(type="list", formula1=f"={R_CAT}", allow_blank=True, showDropDown=False)
dv.error = "Categoria não está na lista. Cadastre na aba LISTAS."; dv.errorTitle = "Categoria"
car.add_data_validation(dv); dv.add(f"D{LIN_1}:D{FIM_COMPRA}")

dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(N_PARC),
                    allow_blank=True)
dv.error = f"O número de parcelas tem que ser de 1 a {N_PARC}."; dv.errorTitle = "Parcelas"
car.add_data_validation(dv); dv.add(f"F{LIN_1}:F{FIM_COMPRA}")

dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
dv.error = "O valor total da compra tem que ser maior que zero."; dv.errorTitle = "Valor"
car.add_data_validation(dv); dv.add(f"E{LIN_1}:E{FIM_COMPRA}")

# resumo por cartao
r0 = FIM_COMPRA + 3
cel(car, r0, 1, "SALDO POR CARTÃO", fonte=Font(size=12, bold=True, color=AZUL), fill=fill_az, al=esq)
car.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=5)
linha_hdr(car, r0 + 1, ["Cartão", "Total comprado", "Já pago", "Falta pagar", "Parcelas em aberto"])
for i in range(8):
    r = r0 + 2 + i
    lr = LIN_1 + i
    cel(car, r, 1, f'=IF(LISTAS!$F{lr}="","",LISTAS!$F{lr})', fonte=f_b, al=esq)
    cel(car, r, 2, f'=IF($A{r}="","",SUMIF($B${LIN_1}:$B${FIM_COMPRA},$A{r},$E${LIN_1}:$E${FIM_COMPRA}))',
        fonte=f_n, fmt=GS0, al=ctr)
    cel(car, r, 3, f'=IF($A{r}="","",SUMIF($B${LIN_1}:$B${FIM_COMPRA},$A{r},$L${LIN_1}:$L${FIM_COMPRA}))',
        fonte=f_n, fmt=GS0, al=ctr)
    cel(car, r, 4, f'=IF($A{r}="","",SUMIF($B${LIN_1}:$B${FIM_COMPRA},$A{r},$M${LIN_1}:$M${FIM_COMPRA}))',
        fonte=f_b, fill=fill_vm, fmt=GS0, al=ctr)
    cel(car, r, 5,
        f'=IF($A{r}="","",SUMPRODUCT((PARCELAS!$C${LIN_1}:$AL${FIM_COMPRA}>'
        f'YEAR(TODAY())*12+MONTH(TODAY()))*(PARCELAS!$AM${LIN_1}:$BV${FIM_COMPRA}>0)*1*'
        f'(CARTÕES!$B${LIN_1}:$B${FIM_COMPRA}=$A{r})))',
        fonte=f_n, fmt=NUM0, al=ctr)
r = r0 + 10
cel(car, r, 1, "TOTAL", fonte=f_b, fill=fill_az, al=esq)
for c in (2, 3, 4):
    cel(car, r, c, f"=SUM({L(c)}{r0+2}:{L(c)}{r0+9})", fonte=f_b, fill=fill_az, fmt=GS0, al=ctr)
cel(car, r, 5, f"=SUM(E{r0+2}:E{r0+9})", fonte=f_b, fill=fill_az, fmt=NUM0, al=ctr)
car.sheet_view.showGridLines = False

# ------------------------------------------------------------------- PARCELAS
par = wb.create_sheet("PARCELAS")
par.sheet_properties.tabColor = "A6A6A6"
cabecalho(par, "PARCELAS (automático) — NÃO MEXA NESTA ABA",
          "Cada linha é uma compra da aba CARTÕES. As colunas espalham as parcelas pelos meses.", 10)
larguras(par, {"A": 30, "B": 20})
linha_hdr(par, LIN_HDR, ["Compra", "Valor no mês do RESUMO MENSAL"])
for k in range(1, N_PARC + 1):
    cel(par, LIN_HDR, 2 + k, f"Mês {k}", fonte=f_hdr, fill=fill_hdr, al=ctr)
    cel(par, LIN_HDR, 38 + k, f"Valor {k}", fonte=f_hdr, fill=fill_hdr, al=ctr)
    par.column_dimensions[L(2 + k)].width = 8
    par.column_dimensions[L(38 + k)].width = 12
par.freeze_panes = "C6"

for r in range(LIN_1, FIM_COMPRA + 1):
    cel(par, r, 1, f'=IF(CARTÕES!$C{r}="","",CARTÕES!$C{r})', fonte=f_auto, al=esq)
    cel(par, r, 2,
        f"=SUMPRODUCT((C{r}:AL{r}='RESUMO MENSAL'!$C$4)*AM{r}:BV{r})",
        fonte=f_auto, fill=fill_cz, fmt=GS0, al=ctr)
    for k in range(1, N_PARC + 1):
        cm, cv = 2 + k, 38 + k
        cel(par, r, cm,
            f'=IF(OR(CARTÕES!$K{r}=0,CARTÕES!$F{r}=""),0,IF({k}>CARTÕES!$F{r},0,CARTÕES!$K{r}+{k}-1))',
            fonte=f_auto, fmt=NUM0, al=ctr, borda=False)
        cel(par, r, cv,
            f'=IF({L(cm)}{r}=0,0,IF({k}=CARTÕES!$F{r},'
            f'CARTÕES!$E{r}-(CARTÕES!$F{r}-1)*ROUND(CARTÕES!$E{r}/CARTÕES!$F{r},0),'
            f'ROUND(CARTÕES!$E{r}/CARTÕES!$F{r},0)))',
            fonte=f_auto, fmt=NUM0, al=ctr, borda=False)
par.sheet_view.showGridLines = False

MESGRID = f"PARCELAS!$C${LIN_1}:$AL${FIM_COMPRA}"
VALGRID = f"PARCELAS!$AM${LIN_1}:$BV${FIM_COMPRA}"

# -------------------------------------------------------------- RESUMO MENSAL
rm = wb.create_sheet("RESUMO MENSAL")
rm.sheet_properties.tabColor = "548235"
cabecalho(rm, "RESUMO MENSAL", "Escolha o ano e o mês nas células amarelas. Tudo abaixo se recalcula sozinho.", 8)
larguras(rm, {"A": 30, "B": 18, "C": 18, "D": 18, "E": 4, "F": 30, "G": 18, "H": 18, "I": 18})

cel(rm, 3, 1, "Ano", fonte=f_b, fill=fill_az, al=ctr)
cel(rm, 3, 2, "Mês", fonte=f_b, fill=fill_az, al=ctr)
cel(rm, 3, 3, "Índice (auto)", fonte=f_b, fill=fill_cz, al=ctr)
cel(rm, 4, 1, ANO, fonte=f_kpi, fill=fill_am, fmt="0", al=ctr)
cel(rm, 4, 2, 1, fonte=f_kpi, fill=fill_am, fmt="0", al=ctr)
cel(rm, 4, 3, "=$A$4*12+$B$4", fonte=f_auto, fill=fill_cz, fmt=NUM0, al=ctr)
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="12")
dv.error = "Digite o mês de 1 a 12."; dv.errorTitle = "Mês"
rm.add_data_validation(dv); dv.add("B4")
dv = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100")
dv.error = "Digite o ano com 4 dígitos."; dv.errorTitle = "Ano"
rm.add_data_validation(dv); dv.add("A4")
cel(rm, 4, 4, '=INDEX(LISTAS!$K$6:$K$17,$B$4)&" de "&$A$4', fonte=f_kpi, fill=fill_az, al=ctr)
for i, m in enumerate(MESES):
    cel(lst, LIN_1 + i, 11, m, fonte=f_n, al=esq)
cel(lst, LIN_HDR, 11, "Mês", fonte=f_hdr, fill=fill_hdr, al=ctr)
larguras(lst, {"K": 14})

LR = f"'LANÇAMENTOS'!$I${LIN_1}:$I${FIM_LANC}"     # indice do mes
LT = f"'LANÇAMENTOS'!$B${LIN_1}:$B${FIM_LANC}"     # tipo
LS = f"'LANÇAMENTOS'!$C${LIN_1}:$C${FIM_LANC}"     # status
LC = f"'LANÇAMENTOS'!$D${LIN_1}:$D${FIM_LANC}"     # categoria
LV = f"'LANÇAMENTOS'!$G${LIN_1}:$G${FIM_LANC}"     # valor


def sm(tipo, status=None, cat=None, idx="$C$4"):
    f = f"SUMIFS({LV},{LR},{idx},{LT},\"{tipo}\""
    if status: f += f",{LS},\"{status}\""
    if cat: f += f",{LC},{cat}"
    return f + ")"


CARTAO_MES = f"SUMPRODUCT(({MESGRID}=$C$4)*{VALGRID})"

KPIS = [
    ("Entradas REALIZADAS", sm("ENTRADA", "REALIZADO"), fill_vd),
    ("Entradas PROJETADAS", sm("ENTRADA", "PROJETADO"), fill_am),
    ("Saídas REALIZADAS", sm("SAÍDA", "REALIZADO"), fill_vm),
    ("Saídas PROJETADAS", sm("SAÍDA", "PROJETADO"), fill_am),
]
for i, (rot, f, fl) in enumerate(KPIS):
    c = 1 + i * 2
    cel(rm, 6, c, rot, fonte=f_b, fill=fl, al=ctr)
    cel(rm, 6, c + 1, None, fill=fl)
    cel(rm, 7, c, "=" + f, fonte=f_kpi, fill=fl, fmt=GS, al=ctr)
    cel(rm, 7, c + 1, None, fill=fl)
    rm.merge_cells(start_row=6, start_column=c, end_row=6, end_column=c + 1)
    rm.merge_cells(start_row=7, start_column=c, end_row=7, end_column=c + 1)
rm.row_dimensions[7].height = 24

BLOCO = [
    ("Entradas do mês (realizado + projetado)", "=" + sm("ENTRADA"), fill_vd, f_kpi),
    ("Saídas do mês (realizado + projetado)", "=" + sm("SAÍDA"), fill_vm, f_kpi),
    ("Parcelas de cartão que vencem no mês", "=" + CARTAO_MES, fill_vm, f_kpi),
    ("RESULTADO DO MÊS", "=$B$9-$B$10-$B$11", fill_az, Font(size=16, bold=True, color=AZUL)),
]
for i, (rot, f, fl, fo) in enumerate(BLOCO):
    r = 9 + i
    cel(rm, r, 1, rot, fonte=f_b, fill=fl, al=esq)
    cel(rm, r, 2, f, fonte=fo, fill=fl, fmt=GS, al=ctr)
    rm.row_dimensions[r].height = 22
rm.conditional_formatting.add("B12", CellIsRule(operator="lessThan", formula=["0"],
                                                font=Font(size=16, bold=True, color="9C0006")))

cel(rm, 14, 1, "ENTRADAS POR CATEGORIA", fonte=Font(size=12, bold=True, color=AZUL), fill=fill_az, al=esq)
rm.merge_cells(start_row=14, start_column=1, end_row=14, end_column=4)
linha_hdr(rm, 15, ["Categoria", "Realizado", "Projetado", "Total"])
cel(rm, 14, 6, "SAÍDAS POR CATEGORIA", fonte=Font(size=12, bold=True, color=AZUL), fill=fill_az, al=esq)
rm.merge_cells(start_row=14, start_column=6, end_row=14, end_column=9)
linha_hdr(rm, 15, ["Categoria", "Realizado", "Projetado", "Total"], 6)

for i in range(N_CAT):
    r = 16 + i
    lr = LIN_1 + i
    for base, tipo in ((1, "ENTRADA"), (6, "SAÍDA")):
        cel(rm, r, base,
            f'=IF(LISTAS!$B{lr}="{tipo}",LISTAS!$A{lr},"")', fonte=f_n, al=esq)
        ref = f"${L(base)}{r}"
        cel(rm, r, base + 1, f'=IF({ref}="","",{sm(tipo, "REALIZADO", ref)})',
            fonte=f_n, fmt=GS0, al=ctr)
        cel(rm, r, base + 2, f'=IF({ref}="","",{sm(tipo, "PROJETADO", ref)})',
            fonte=f_n, fmt=GS0, al=ctr)
        cel(rm, r, base + 3, f'=IF({ref}="","",{sm(tipo, None, ref)})',
            fonte=f_b, fmt=GS0, al=ctr)
rtot = 16 + N_CAT
for base, fl in ((1, fill_vd), (6, fill_vm)):
    cel(rm, rtot, base, "TOTAL", fonte=f_b, fill=fl, al=esq)
    for j in (1, 2, 3):
        cel(rm, rtot, base + j, f"=SUM({L(base+j)}16:{L(base+j)}{rtot-1})",
            fonte=f_b, fill=fl, fmt=GS, al=ctr)

r = rtot + 2
cel(rm, r, 1, "CONFERÊNCIA", fonte=Font(size=12, bold=True, color=AZUL), fill=fill_az, al=esq)
rm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
CONF = [
    ("Entradas: KPI menos soma das categorias", f"=$B$9-$D${rtot}"),
    ("Saídas: KPI menos soma das categorias", f"=$B$10-$I${rtot}"),
    ("Lançamentos sem tipo ou sem status",
     f'=SUMPRODUCT((\'LANÇAMENTOS\'!$A${LIN_1}:$A${FIM_LANC}<>"")*'
     f'((\'LANÇAMENTOS\'!$B${LIN_1}:$B${FIM_LANC}="")+'
     f'(\'LANÇAMENTOS\'!$C${LIN_1}:$C${FIM_LANC}="")>0))'),
    ("Compras de cartão sem nº de parcelas",
     f'=SUMPRODUCT((CARTÕES!$A${LIN_1}:$A${FIM_COMPRA}<>"")*'
     f'(CARTÕES!$F${LIN_1}:$F${FIM_COMPRA}=""))'),
    ("Parcelas: soma das parcelas menos total comprado",
     f"=SUMPRODUCT(({MESGRID}>0)*{VALGRID})-SUM(CARTÕES!$E${LIN_1}:$E${FIM_COMPRA})"),
]
for i, (rot, f) in enumerate(CONF):
    rr = r + 1 + i
    cel(rm, rr, 1, rot, fonte=f_n, al=esq)
    rm.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
    cel(rm, rr, 4, f, fonte=f_b, fmt="#,##0", al=ctr)
    rm.conditional_formatting.add(f"D{rr}", CellIsRule(operator="notEqual", formula=["0"],
                                                       fill=fill_vm, font=Font(bold=True, color="9C0006")))
    rm.conditional_formatting.add(f"D{rr}", CellIsRule(operator="equal", formula=["0"],
                                                       fill=fill_vd, font=Font(bold=True, color="375623")))
cel(rm, r + 6, 1, "Tudo em zero = a planilha está fechando. Qualquer número diferente de zero fica vermelho.",
    fonte=f_sub, borda=False, al=esq)
rm.freeze_panes = "A6"
rm.sheet_view.showGridLines = False

# --------------------------------------------------------------- RESUMO ANUAL
ra = wb.create_sheet("RESUMO ANUAL")
ra.sheet_properties.tabColor = "833C00"
cabecalho(ra, "RESUMO ANUAL", "Os 12 meses do ano lado a lado. Troque o ano na célula amarela.", 8)
larguras(ra, {"A": 30})
for c in range(2, 15):
    ra.column_dimensions[L(c)].width = 15

cel(ra, 3, 1, "Ano", fonte=f_b, fill=fill_az, al=ctr)
cel(ra, 4, 1, ANO, fonte=f_kpi, fill=fill_am, fmt="0", al=ctr)
dv = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100")
dv.error = "Digite o ano com 4 dígitos."; dv.errorTitle = "Ano"
ra.add_data_validation(dv); dv.add("A4")

linha_hdr(ra, LIN_HDR, ["Mês", "Entradas realizadas", "Entradas projetadas", "Entradas total",
                        "Saídas realizadas", "Saídas projetadas", "Saídas total",
                        "Cartão (parcelas)", "Resultado do mês", "Saldo acumulado"])

for i in range(12):
    r = LIN_1 + i
    idx = f"($A$4*12+{i+1})"
    cel(ra, r, 1, f'=INDEX(LISTAS!$K$6:$K$17,{i+1})', fonte=f_b, fill=fill_cz, al=esq)
    cel(ra, r, 2, "=" + sm("ENTRADA", "REALIZADO", idx=idx), fonte=f_n, fmt=GS0, al=ctr)
    cel(ra, r, 3, "=" + sm("ENTRADA", "PROJETADO", idx=idx), fonte=f_n, fmt=GS0, al=ctr)
    cel(ra, r, 4, f"=B{r}+C{r}", fonte=f_b, fill=fill_vd, fmt=GS0, al=ctr)
    cel(ra, r, 5, "=" + sm("SAÍDA", "REALIZADO", idx=idx), fonte=f_n, fmt=GS0, al=ctr)
    cel(ra, r, 6, "=" + sm("SAÍDA", "PROJETADO", idx=idx), fonte=f_n, fmt=GS0, al=ctr)
    cel(ra, r, 7, f"=E{r}+F{r}", fonte=f_b, fill=fill_vm, fmt=GS0, al=ctr)
    cel(ra, r, 8, f"=SUMPRODUCT(({MESGRID}={idx})*{VALGRID})", fonte=f_n, fill=fill_vm, fmt=GS0, al=ctr)
    cel(ra, r, 9, f"=D{r}-G{r}-H{r}", fonte=f_b, fill=fill_az, fmt=GS, al=ctr)
    cel(ra, r, 10, (f"=I{r}" if i == 0 else f"=J{r-1}+I{r}"), fonte=f_b, fmt=GS, al=ctr)

rt = LIN_1 + 12
cel(ra, rt, 1, "TOTAL DO ANO", fonte=f_b, fill=fill_az, al=esq)
for c in range(2, 10):
    cel(ra, rt, c, f"=SUM({L(c)}{LIN_1}:{L(c)}{rt-1})", fonte=f_b, fill=fill_az, fmt=GS, al=ctr)
cel(ra, rt, 10, f"=J{rt-1}", fonte=f_b, fill=fill_az, fmt=GS, al=ctr)
for col in ("I", "J"):
    ra.conditional_formatting.add(f"{col}{LIN_1}:{col}{rt}",
                                  CellIsRule(operator="lessThan", formula=["0"],
                                             font=Font(bold=True, color="9C0006"), fill=fill_vm))

r0 = rt + 3
cel(ra, r0, 1, "SAÍDAS POR CATEGORIA, MÊS A MÊS", fonte=Font(size=12, bold=True, color=AZUL),
    fill=fill_az, al=esq)
ra.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=14)
linha_hdr(ra, r0 + 1, ["Categoria"] + [m[:3] for m in MESES] + ["Total do ano"])
for i in range(N_CAT):
    r = r0 + 2 + i
    lr = LIN_1 + i
    cel(ra, r, 1, f'=IF(LISTAS!$B{lr}="SAÍDA",LISTAS!$A{lr},"")', fonte=f_n, al=esq)
    for m in range(12):
        cel(ra, r, 2 + m,
            f'=IF($A{r}="","",{sm("SAÍDA", None, f"$A{r}", idx=f"($A$4*12+{m+1})")})',
            fonte=f_n, fmt=GS0, al=ctr)
    cel(ra, r, 14, f'=IF($A{r}="","",SUM(B{r}:M{r}))', fonte=f_b, fmt=GS0, al=ctr)
rsc = r0 + 2 + N_CAT
cel(ra, rsc, 1, "TOTAL SAÍDAS", fonte=f_b, fill=fill_vm, al=esq)
for c in range(2, 15):
    cel(ra, rsc, c, f"=SUM({L(c)}{r0+2}:{L(c)}{rsc-1})", fonte=f_b, fill=fill_vm, fmt=GS0, al=ctr)

r1 = rsc + 3
cel(ra, r1, 1, "ENTRADAS POR CATEGORIA, MÊS A MÊS", fonte=Font(size=12, bold=True, color=AZUL),
    fill=fill_az, al=esq)
ra.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=14)
linha_hdr(ra, r1 + 1, ["Categoria"] + [m[:3] for m in MESES] + ["Total do ano"])
for i in range(N_CAT):
    r = r1 + 2 + i
    lr = LIN_1 + i
    cel(ra, r, 1, f'=IF(LISTAS!$B{lr}="ENTRADA",LISTAS!$A{lr},"")', fonte=f_n, al=esq)
    for m in range(12):
        cel(ra, r, 2 + m,
            f'=IF($A{r}="","",{sm("ENTRADA", None, f"$A{r}", idx=f"($A$4*12+{m+1})")})',
            fonte=f_n, fmt=GS0, al=ctr)
    cel(ra, r, 14, f'=IF($A{r}="","",SUM(B{r}:M{r}))', fonte=f_b, fmt=GS0, al=ctr)
rec = r1 + 2 + N_CAT
cel(ra, rec, 1, "TOTAL ENTRADAS", fonte=f_b, fill=fill_vd, al=esq)
for c in range(2, 15):
    cel(ra, rec, c, f"=SUM({L(c)}{r1+2}:{L(c)}{rec-1})", fonte=f_b, fill=fill_vd, fmt=GS0, al=ctr)

r2 = rec + 3
cel(ra, r2, 1, "CONFERÊNCIA DO ANO", fonte=Font(size=12, bold=True, color=AZUL), fill=fill_az, al=esq)
ra.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=4)
CONF2 = [
    ("Entradas: total do ano menos soma das categorias", f"=D{rt}-N{rec}"),
    ("Saídas: total do ano menos soma das categorias", f"=G{rt}-N{rsc}"),
    ("Saldo acumulado bate com a soma dos resultados", f"=J{rt}-I{rt}"),
]
for i, (rot, f) in enumerate(CONF2):
    rr = r2 + 1 + i
    cel(ra, rr, 1, rot, fonte=f_n, al=esq)
    ra.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
    cel(ra, rr, 4, f, fonte=f_b, fmt="#,##0", al=ctr)
    ra.conditional_formatting.add(f"D{rr}", CellIsRule(operator="notEqual", formula=["0"],
                                                       fill=fill_vm, font=Font(bold=True, color="9C0006")))
    ra.conditional_formatting.add(f"D{rr}", CellIsRule(operator="equal", formula=["0"],
                                                       fill=fill_vd, font=Font(bold=True, color="375623")))
ra.freeze_panes = "B6"
ra.sheet_view.showGridLines = False

# -------------------------------------------------------------------- EXEMPLO
ex = wb.create_sheet("EXEMPLO")
ex.sheet_properties.tabColor = "FFC000"
cabecalho(ex, "EXEMPLO — COMO PREENCHER CADA CASO",
          "Esta aba é só modelo. Nada aqui entra nas contas. Pode apagar a aba quando não precisar mais.", 8)
larguras(ex, {"A": 12, "B": 12, "C": 13, "D": 26, "E": 34, "F": 17, "G": 16, "H": 44})

cel(ex, 4, 1, "MODELOS PARA A ABA LANÇAMENTOS", fonte=Font(size=12, bold=True, color=AZUL),
    fill=fill_az, al=esq)
ex.merge_cells(start_row=4, start_column=1, end_row=4, end_column=8)
linha_hdr(ex, 5, ["Data", "Tipo", "Status", "Categoria", "Descrição",
                  "Meio de pagamento", "Valor (Gs.)", "Por que assim"])
EX_LANC = [
    (dt.date(2026, 1, 5), "ENTRADA", "REALIZADO", "Salário", "Salário de janeiro", "Conta corrente", 8500000,
     "Dinheiro que já caiu na conta. Status REALIZADO."),
    (dt.date(2026, 1, 10), "SAÍDA", "REALIZADO", "Alimentação", "Supermercado da semana", "Débito", 780000,
     "Saída paga no débito. Sai da conta na hora, então é lançamento normal."),
    (dt.date(2026, 1, 12), "SAÍDA", "REALIZADO", "Serviços (luz, água, internet)", "Conta de luz", "Boleto", 320000,
     "Boleto pago. Também é lançamento normal."),
    (dt.date(2026, 2, 5), "ENTRADA", "PROJETADO", "Salário", "Salário de fevereiro", "Conta corrente", 8500000,
     "Ainda não caiu. Status PROJETADO para já aparecer na projeção."),
    (dt.date(2026, 2, 15), "SAÍDA", "PROJETADO", "Moradia", "Aluguel de fevereiro", "Transferência", 2500000,
     "Gasto certo que ainda vai acontecer. PROJETADO."),
    (dt.date(2026, 1, 20), "ENTRADA", "REALIZADO", "Venda", "Venda de móvel usado", "Dinheiro", 1200000,
     "Entrada em espécie. Meio = Dinheiro."),
]
for i, linha in enumerate(EX_LANC):
    r = 6 + i
    for j, v in enumerate(linha):
        fmt = DATA if j == 0 else (GS if j == 6 else None)
        cel(ex, r, 1 + j, v, fonte=f_n if j < 7 else f_sub,
            fmt=fmt, al=ctr if j in (0, 1, 2, 5, 6) else esq)
    ex.row_dimensions[r].height = 16

r0 = 6 + len(EX_LANC) + 2
cel(ex, r0, 1, "MODELOS PARA A ABA CARTÕES", fonte=Font(size=12, bold=True, color=AZUL),
    fill=fill_az, al=esq)
ex.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=8)
linha_hdr(ex, r0 + 1, ["Data da compra", "Cartão", "Descrição", "Categoria",
                       "Valor total (Gs.)", "Nº parcelas", "1ª parcela", "O que a planilha faz"])
EX_CART = [
    (dt.date(2026, 1, 18), "Cartão 1", "Geladeira", "Manutenção", 6000000, 12, None,
     "Deixe a 1ª parcela em branco: a planilha começa em fev/2026 e cria 12 parcelas de Gs. 500.000."),
    (dt.date(2026, 3, 2), "Cartão 1", "Passagem aérea", "Lazer", 4500000, 6, dt.date(2026, 4, 1),
     "Aqui a 1ª parcela foi digitada. Vão 6 parcelas de Gs. 750.000, de abr a set/2026."),
    (dt.date(2026, 2, 8), "Cartão 2", "Farmácia", "Saúde", 350000, 1, None,
     "Compra à vista no crédito: 1 parcela. Cai inteira no mês seguinte."),
    (dt.date(2026, 5, 4), "Cartão 2", "Notebook", "Educação", 10000000, 7, None,
     "Gs. 10.000.000 em 7x não dá número redondo. A planilha ajusta a última parcela para fechar exato."),
]
for i, linha in enumerate(EX_CART):
    r = r0 + 2 + i
    for j, v in enumerate(linha):
        fmt = DATA if j in (0, 6) else (GS if j == 4 else ("0" if j == 5 else None))
        cel(ex, r, 1 + j, v, fonte=f_n if j < 7 else f_sub,
            fmt=fmt, al=ctr if j in (0, 1, 4, 5, 6) else esq)
    ex.row_dimensions[r].height = 16

r1 = r0 + 2 + len(EX_CART) + 2
cel(ex, r1, 1, "CASOS QUE CONFUNDEM — O QUE FAZER", fonte=Font(size=12, bold=True, color=AZUL),
    fill=fill_az, al=esq)
ex.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=8)
DUVIDAS = [
    ("Paguei a fatura do cartão", "NÃO lance nada. A parcela já foi contada no mês em que venceu. "
     "Lançar de novo contaria o mesmo gasto duas vezes."),
    ("Comprei no cartão e paguei à vista (1x)", "Lance na aba CARTÕES com Nº de parcelas = 1."),
    ("Comprei no débito", "Lance na aba LANÇAMENTOS como SAÍDA, Meio = Débito. Débito não é cartão de crédito."),
    ("Passei dinheiro da conta para a poupança", "Não lance. Transferência entre suas contas não é entrada nem saída."),
    ("Recebi um empréstimo", "ENTRADA, categoria Empréstimo recebido. As parcelas que você vai pagar "
     "são SAÍDA em Empréstimo / Financiamento, uma por mês, com status PROJETADO."),
    ("Parcelei uma compra e depois quitei tudo", "Mude o Nº de parcelas para o número que você realmente pagou "
     "e ajuste o Valor total. A planilha recalcula sozinha."),
    ("Errei um lançamento", "Selecione a linha inteira e aperte DELETE para limpar. "
     "Nunca use Excluir linha, senão as fórmulas automáticas somem."),
]
linha_hdr(ex, r1 + 1, ["Situação", "O que fazer"])
ex.merge_cells(start_row=r1 + 1, start_column=2, end_row=r1 + 1, end_column=8)
for i, (a, b) in enumerate(DUVIDAS):
    r = r1 + 2 + i
    cel(ex, r, 1, a, fonte=f_b, al=wrap)
    cel(ex, r, 2, b, fonte=f_n, al=wrap)
    ex.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    ex.row_dimensions[r].height = 30
ex.sheet_view.showGridLines = False

wb.move_sheet("LISTAS", offset=6)
wb.move_sheet("PARCELAS", offset=3)
wb.active = 0
wb.save(SAIDA)
print("gerado:", SAIDA, f"{SAIDA.stat().st_size/1024:.0f} KB")
print("abas:", wb.sheetnames)
