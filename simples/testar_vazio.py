#!/usr/bin/env python3
"""Confere o arquivo VAZIO: recalcula, exige zero erro de formula, zero numero
nos resumos e confirma que todas as listas suspensas e colunas automaticas existem."""
import shutil, sys
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "gerador"))
from recalc_lo import recalc

LIMPO = RAIZ / "CONTROLE_FINANCEIRO_SIMPLES.xlsx"
COPIA = RAIZ / "simples" / "_vazio.xlsx"
P1, F_LANC, F_TRF, F_CONTA, F_COMPRA, F_CAT, F_CARTAO = 6, 1005, 305, 25, 65, 45, 13
shutil.copy(LIMPO, COPIA)

print("recalculando o arquivo vazio...")
res = recalc(COPIA)
print(f"  {res['total_formulas']} formulas, {res['seconds']}s, {res['total_errors']} com erro")
for k, v in list(res.get("error_summary", {}).items())[:10]:
    print("   ", k, v[:12])

wb = openpyxl.load_workbook(COPIA, data_only=True)
wf = openpyxl.load_workbook(COPIA)
ok = falhas = 0
def conf(nome, esperado, obtido):
    global ok, falhas
    bom = esperado == obtido
    print(("  [OK   ] " if bom else "  [FALHA] ") + nome.ljust(56)
          + f"esperado {esperado!r:>12}  obtido {obtido!r:>12}")
    if bom: ok += 1
    else: falhas += 1

print("\n[1] Abas presentes, na ordem")
conf("10 abas na ordem certa",
     ["INSTRUÇÕES", "CONTAS", "LANÇAMENTOS", "TRANSFERÊNCIAS", "CARTÕES",
      "RESUMO MENSAL", "RESUMO ANUAL", "EXEMPLO", "PARCELAS", "LISTAS"], wb.sheetnames)

print("\n[2] Tudo zerado no arquivo em branco")
cv, rm, ra, cr, tv = (wb["CONTAS"], wb["RESUMO MENSAL"], wb["RESUMO ANUAL"],
                      wb["CARTÕES"], wb["TRANSFERÊNCIAS"])
for aba, cels in [("CONTAS", [("A4", "saldo atual"), ("C4", "saldo previsto"),
                              ("E4", "dívida cartões"), ("G4", "posição líquida")]),
                  ("CARTÕES", [("A4", "total comprado"), ("C4", "parcelas a vencer"),
                               ("E4", "pago"), ("G4", "dívida atual")]),
                  ("TRANSFERÊNCIAS", [("A4", "total transferido"), ("C4", "pagamentos de fatura"),
                                      ("E4", "digitadas"), ("F4", "erros")]),
                  ("RESUMO MENSAL", [("A7", "entradas realizadas"), ("C7", "entradas projetadas"),
                                     ("E7", "saídas realizadas"), ("G7", "saídas projetadas"),
                                     ("B9", "entradas do mês"), ("B10", "saídas do mês"),
                                     ("B11", "cartão do mês"), ("B12", "resultado"),
                                     ("G9", "saldo em contas"), ("G10", "dívida"),
                                     ("G11", "transferências"), ("G12", "posição líquida")])]:
    for c, rot in cels:
        conf(f"{aba} {c} · {rot}", 0, wb[aba][c].value)
conf("RESUMO ANUAL · 12 meses zerados", True,
     all(ra.cell(P1 + m, c).value == 0 for m in range(12) for c in (4, 7, 8, 9, 10, 11)))
conf("CONTAS · 20 linhas de saldo em branco", True,
     all(cv.cell(P1 + i, 10).value in (None, "") for i in range(20)))

print("\n[3] Conferências internas fecham em zero")
for lin in range(58, 67):
    rot = rm.cell(lin, 1).value
    if isinstance(rot, str) and rm.cell(lin, 4).value is not None:
        conf("MENSAL · " + rot[:42], 0, rm.cell(lin, 4).value)
for r in range(ra.max_row - 4, ra.max_row + 1):
    rot = ra.cell(r, 1).value
    if isinstance(rot, str) and ra.cell(r, 4).value is not None:
        conf("ANUAL · " + rot[:42], 0, ra.cell(r, 4).value)

print("\n[4] Listas suspensas (dropdowns)")
ESPERADO = {
    "CONTAS": {f"C{P1}:C{F_CONTA}": "Tipo de conta", f"D{P1}:D{F_CONTA}": "Saldo inicial"},
    "LANÇAMENTOS": {f"B{P1}:B{F_LANC}": "Tipo", f"C{P1}:C{F_LANC}": "Status",
                    f"D{P1}:D{F_LANC}": "Categoria", f"F{P1}:F{F_LANC}": "Conta",
                    f"G{P1}:G{F_LANC}": "Valor"},
    "TRANSFERÊNCIAS": {f"B{P1}:B{F_TRF}": "Status", f"C{P1}:C{F_TRF}": "De",
                       f"D{P1}:D{F_TRF}": "Para", f"F{P1}:F{F_TRF}": "Valor"},
    "CARTÕES": {f"B{P1}:B{F_COMPRA}": "Cartão", f"D{P1}:D{F_COMPRA}": "Categoria",
                f"E{P1}:E{F_COMPRA}": "Valor total", f"F{P1}:F{F_COMPRA}": "Nº de parcelas"},
    "LISTAS": {f"E{P1}:E{F_CARTAO}": "Limite do cartão"},
    "RESUMO MENSAL": {"A4": "Ano", "B4": "Mês"},
    "RESUMO ANUAL": {"A4": "Ano"},
}
for aba, faixas in ESPERADO.items():
    tem = set()
    for dv in wf[aba].data_validations.dataValidation:
        tem.update(str(x) for x in dv.sqref.ranges)
    for faixa, rot in faixas.items():
        conf(f"{aba} · {faixa} · {rot}", True, faixa in tem)

print("\n[5] Colunas automáticas com fórmula em todas as linhas")
con, lan, trf, car, par, lst = (wf["CONTAS"], wf["LANÇAMENTOS"], wf["TRANSFERÊNCIAS"],
                                wf["CARTÕES"], wf["PARCELAS"], wf["LISTAS"])
def qtd(ws, col, r0, r1):
    return sum(1 for r in range(r0, r1 + 1) if str(ws.cell(r, col).value or "").startswith("="))
for c, rot in [(6, "Entradas"), (7, "Saídas"), (8, "Transf. recebidas"), (9, "Transf. enviadas"),
               (10, "SALDO ATUAL"), (11, "Entradas prev."), (12, "Saídas prev."),
               (13, "Transf. prev."), (14, "SALDO PREVISTO")]:
    conf(f"CONTAS · {rot}", 20, qtd(con, c, P1, F_CONTA))
for c, rot in [(8, "Mês"), (9, "Índice")]:
    conf(f"LANÇAMENTOS · {rot}", 1000, qtd(lan, c, P1, F_LANC))
for c, rot in [(7, "Mês"), (8, "Índice"), (9, "Conferência")]:
    conf(f"TRANSFERÊNCIAS · {rot}", 300, qtd(trf, c, P1, F_TRF))
for c, rot in [(8, "Início"), (9, "Valor da parcela"), (10, "Última parcela"),
               (11, "Índice"), (12, "Parcelas vencidas"), (13, "Parcelas a vencer")]:
    conf(f"CARTÕES · {rot}", 60, qtd(car, c, P1, F_COMPRA))
conf("LISTAS · coluna contas+cartões", 28, qtd(lst, 14, P1, P1 + 27))
conf("PARCELAS · grade completa (60 x 36 meses + 60 x 36 valores)", 4320,
     sum(1 for r in range(P1, F_COMPRA + 1) for k in range(1, 37)
         if str(par.cell(r, 2 + k).value or "").startswith("=")
         and str(par.cell(r, 38 + k).value or "").startswith("=")) * 2)

print("\n[6] Nenhum dado de exemplo dentro das abas de trabalho")
for ws, col, r1, rot in [(con, 1, F_CONTA, "CONTAS"), (lan, 1, F_LANC, "LANÇAMENTOS"),
                         (trf, 1, F_TRF, "TRANSFERÊNCIAS"), (car, 1, F_COMPRA, "CARTÕES")]:
    conf(f"{rot} sem nenhuma linha preenchida", 0,
         sum(1 for r in range(P1, r1 + 1) if ws.cell(r, col).value is not None))
conf("LISTAS traz as categorias prontas", 22,
     sum(1 for r in range(P1, F_CAT + 1) if lst.cell(r, 1).value))
conf("LISTAS traz 4 cartões para renomear", 4,
     sum(1 for r in range(P1, F_CARTAO + 1) if lst.cell(r, 4).value))
conf("LISTAS sem limite preenchido", 0,
     sum(1 for r in range(P1, F_CARTAO + 1) if lst.cell(r, 5).value is not None))

COPIA.unlink(missing_ok=True)
print("\n" + "=" * 96)
print(f"RESULTADO: {ok} conferências OK, {falhas} falhas, {res['total_errors']} células com erro")
print("=" * 96)
sys.exit(1 if falhas or res["total_errors"] else 0)
